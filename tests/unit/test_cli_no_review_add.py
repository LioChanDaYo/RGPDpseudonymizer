"""Story 10.18: the command line asks whether to add a name the detector
missed when nothing is left to review (T1-T15, no T11), and what must not
change (G1-G3).

Written by the BMAD dev agent (Claude). The "child" tests run the real
``gdpr-pseudo`` command (``process``, and ``batch --workers 1`` on a folder)
in a child process (``cli_no_review_child.py``), so that its standard input
can be closed, piped or run out as for a real user; only detection, the
review's keys and a Ctrl+C at the Nth prompt are replaced there. G2, G3 and
T15 run in the test process. Each test gets its own database. Preconditions
call ``pytest.fail("PRE: ...")``, never ``assert``; the first ``assert``
after them is the expected failure on the unfixed code. The tests assert
Lionel's approved texts exactly, the written files and the stored rows.
Invented names only; every other word is a main-corpus word. No spaCy
model, no Qt.
"""

from __future__ import annotations

import csv
import json
import os
import re
import subprocess
import sys
from collections.abc import Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from unittest.mock import patch

import openpyxl
import pytest
import typer.main
from rich.prompt import Confirm, Prompt

from gdpr_pseudonymizer.cli.commands.batch import _process_single_document_worker
from gdpr_pseudonymizer.cli.main import app
from gdpr_pseudonymizer.core.document_processor import DocumentProcessor
from gdpr_pseudonymizer.data.database import init_database, open_database
from gdpr_pseudonymizer.data.repositories.mapping_repository import (
    SQLiteMappingRepository,
)
from gdpr_pseudonymizer.nlp.entity_detector import DetectedEntity

PASSPHRASE = "zorbalia_pass_A_123!"
SHEET = "Zorbalia"
CHILD = Path(__file__).with_name("cli_no_review_child.py")
REPO = Path(__file__).resolve().parents[2]
TIMEOUT = 120

ZQ = "Zorbal Quentrix"
VM = "Vorlan Mestrix"
VZ = "Vorlan Zorbal"
ZV = "Zorbaville"

ENTER = "\r"
SPACE = " "

# Lionel's approved texts (Story 10.18, Q1, Q4 and Q5), as the screen shows them.
NO_NAMES_FOUND = "ℹ️  No names were found in this document."
NO_NEW_NAMES = (
    "ℹ️  No new names to review: every name found is already in your mapping"
    " database."
)
SELECTED_NO_NAMES_FOUND = (
    "ℹ️  No names of the selected types were found in this document."
)
SELECTED_NO_NEW_NAMES = (
    "ℹ️  No new names of the selected types to review: every one found is"
    " already in your mapping database."
)
QUESTION = "Add a name the detector missed before the file is written? [y/n]"
NO_ANSWER = (
    "Error: NoAnswerError: No answer could be read, so this document was not"
    " processed: the command line must ask whether to add a name the detector"
    " missed. Run it in a terminal to answer. No output was written for this"
    " file."
)
FALSE_LINE = "All entities have existing mappings"
FAILED = "✗ Processing Failed"
SUCCESS = "Processing Successful"
CANCELLED = "Processing cancelled by user"
BATCH_CANCELLED = "Batch processing cancelled by user"
AUTO_ACCEPTED = "Auto-accepted"

NONE_TXT = "Le contrat avec Zorbal Quentrix est signé. Zorbal Quentrix valide."
NONE_SHEET = [["Nom", "Note"], [ZQ, "valide"], ["ok", "contrat avec Zorbal Quentrix"]]
ADD = b"y\nZorbal Quentrix\nPERSON\nn\n"

_ANSI = re.compile(r"\x1b\[[0-9;?]*[A-Za-z]")

Names = Sequence[tuple[str, str]]
Rows = list[list[str]]


def _pre(condition: bool, what: str) -> None:
    """A precondition: a red failure on every leg, never an AssertionError."""
    if not condition:
        pytest.fail(f"PRE: {what}")


@pytest.fixture(autouse=True)
def _settings(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    """No user or project settings file is read (the in-process tests)."""
    monkeypatch.setenv("HOME", str(tmp_path))
    monkeypatch.setenv("USERPROFILE", str(tmp_path))
    monkeypatch.chdir(tmp_path)


@pytest.fixture
def db(tmp_path: Path) -> Path:
    path = tmp_path / "m.db"
    init_database(str(path), PASSPHRASE)
    return path


@dataclass
class Run:
    code: int
    shown: str


def _child(
    tmp_path: Path,
    args: Sequence[str],
    names: Names = (),
    stdin: bytes | None = None,
    keys: Sequence[str] = (),
    ctrl_c_at: int = 0,
) -> Run:
    """One real ``gdpr-pseudo`` run in a child process. ``stdin=None`` is a
    closed input; bytes are piped in and then run out."""
    env = dict(os.environ)
    env.update(
        HOME=str(tmp_path),
        USERPROFILE=str(tmp_path),
        GDPR_PSEUDO_PASSPHRASE=PASSPHRASE,
        GDPR_PSEUDO_LANG="en",
        PYTHONIOENCODING="utf-8",
        COLUMNS="400",
        NO_REVIEW_NAMES=json.dumps([list(pair) for pair in names]),
        NO_REVIEW_KEYS=json.dumps(list(keys)),
        NO_REVIEW_CTRL_C_AT=str(ctrl_c_at),
    )
    command = [sys.executable, str(CHILD), *args]
    try:
        if stdin is None:
            result = subprocess.run(
                command,
                cwd=tmp_path,
                env=env,
                stdin=subprocess.DEVNULL,
                capture_output=True,
                timeout=TIMEOUT,
            )
        else:
            result = subprocess.run(
                command,
                cwd=tmp_path,
                env=env,
                input=stdin,
                capture_output=True,
                timeout=TIMEOUT,
            )
    except subprocess.TimeoutExpired:
        pytest.fail("PRE: the child process ended")
    shown = (result.stdout + result.stderr).decode("utf-8", errors="replace")
    return Run(result.returncode, _ANSI.sub("", shown))


def _process(
    tmp_path: Path,
    db: Path,
    source: Path,
    output: Path,
    names: Names = (),
    stdin: bytes | None = None,
    keys: Sequence[str] = (),
    ctrl_c_at: int = 0,
    extra: Sequence[str] = (),
) -> Run:
    args = ["process", str(source), "--db", str(db), "-o", str(output), *extra]
    return _child(tmp_path, args, names, stdin, keys, ctrl_c_at)


def _batch(
    tmp_path: Path,
    db: Path,
    folder: Path,
    out_folder: Path,
    names: Names = (),
    stdin: bytes | None = None,
    ctrl_c_at: int = 0,
) -> Run:
    args = [
        "batch",
        str(folder),
        "--db",
        str(db),
        "--output",
        str(out_folder),
        "--workers",
        "1",
    ]
    return _child(tmp_path, args, names, stdin, (), ctrl_c_at)


def _store(tmp_path: Path, db: Path, name: str) -> str:
    """A first run that stores ``name`` (PERSON) through a review."""
    source = tmp_path / f"store_{len(_rows(db))}.txt"
    source.write_text(f"{name} valide.", encoding="utf-8")
    run = _process(
        tmp_path,
        db,
        source,
        tmp_path / f"out_{source.stem}.txt",
        [(name, "PERSON")],
        stdin=b"y\n",
        keys=[ENTER, SPACE],
    )
    _pre(run.code == 0, f"first run: {name} stored (exit 0)")
    return _pseudonym(db, name, "PERSON")


def _write(path: Path, content: str | Rows) -> None:
    if isinstance(content, str):
        path.write_text(content, encoding="utf-8")
        return
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = SHEET
    for row in content:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


def _read(path: Path) -> str | Rows:
    _pre(path.is_file(), f"{path.name} written")
    if path.suffix == ".txt":
        return path.read_text(encoding="utf-8")
    if path.suffix == ".csv":
        with path.open(encoding="utf-8", newline="") as handle:
            return [list(row) for row in csv.reader(handle)]
    workbook = openpyxl.load_workbook(path)
    try:
        return [
            ["" if value is None else str(value) for value in row]
            for row in workbook.worksheets[0].iter_rows(values_only=True)
        ]
    finally:
        workbook.close()


def _flat(content: str | Rows) -> str:
    if isinstance(content, str):
        return content
    return "\n".join("\t".join(row) for row in content)


def _rows(db: Path) -> list[tuple[str, str, str]]:
    """Every stored row: (full_name, entity_type, pseudonym_full)."""
    with open_database(str(db), PASSPHRASE) as session:
        return [
            (row.full_name, row.entity_type, row.pseudonym_full)
            for row in SQLiteMappingRepository(session).find_all()
        ]


def _rows_of(db: Path, name: str, entity_type: str) -> list[tuple[str, str, str]]:
    return [row for row in _rows(db) if row[:2] == (name, entity_type)]


def _pseudonym(db: Path, name: str, entity_type: str) -> str:
    found = _rows_of(db, name, entity_type)
    _pre(len(found) == 1, f"one row stored for ({name}, {entity_type})")
    return found[0][2]


def _with(content: str | Rows, name: str, pseudonym: str) -> str | Rows:
    """The source with every occurrence of ``name`` replaced."""
    if isinstance(content, str):
        return content.replace(name, pseudonym)
    return [[cell.replace(name, pseudonym) for cell in row] for row in content]


# --- AC1, AC2: the question, and "y" adds the names missed ------------------


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason="Story 10.18: reproduces a document with nothing detected written with no question, a missed name in clear on the unfixed code; fixed in Task 2",
)
@pytest.mark.parametrize("ext", [".txt", ".xlsx"])
def test_nothing_found_yes_every_place(tmp_path: Path, db: Path, ext: str) -> None:
    """T1 (AC1, AC2): nothing detected; "y" and a name at two places: both
    places replaced, its row stored, the approved line shown."""
    source, output = tmp_path / f"in{ext}", tmp_path / f"out_in{ext}"
    content: str | Rows = NONE_TXT if ext == ".txt" else NONE_SHEET
    _write(source, content)
    run = _process(tmp_path, db, source, output, stdin=ADD)
    _pre(run.code == 0, "exit code 0")
    written = _read(output)
    assert ZQ not in _flat(written)  # expected failure
    assert len(_rows_of(db, ZQ, "PERSON")) == 1
    pseudonym = _pseudonym(db, ZQ, "PERSON")
    assert written == _with(content, ZQ, pseudonym)
    assert NO_NAMES_FOUND in run.shown
    assert FALSE_LINE not in run.shown


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason="Story 10.18: reproduces a document whose names are all stored written with no question, a missed name in clear on the unfixed code; fixed in Task 2",
)
@pytest.mark.parametrize("ext", [".txt", ".xlsx"])
def test_all_stored_yes_adds_missed(tmp_path: Path, db: Path, ext: str) -> None:
    """T2 (AC1, AC2): every name found is stored; "y" adds the missed one."""
    vm = _store(tmp_path, db, VM)
    source, output = tmp_path / f"in{ext}", tmp_path / f"out_in{ext}"
    content: str | Rows = (
        "Vorlan Mestrix valide. Le contrat avec Zorbal Quentrix."
        if ext == ".txt"
        else [["Nom", "Note"], [VM, "contrat avec Zorbal Quentrix"]]
    )
    _write(source, content)
    run = _process(tmp_path, db, source, output, [(VM, "PERSON")], stdin=ADD)
    _pre(run.code == 0, "exit code 0")
    _pre(AUTO_ACCEPTED in run.shown, "the stored name is accepted without review")
    written = _read(output)
    assert ZQ not in _flat(written)  # expected failure
    assert len(_rows_of(db, ZQ, "PERSON")) == 1
    pseudonym = _pseudonym(db, ZQ, "PERSON")
    assert written == _with(_with(content, VM, vm), ZQ, pseudonym)
    assert NO_NEW_NAMES in run.shown
    assert FALSE_LINE not in run.shown


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason="Story 10.18: reproduces no question, so a missed name next to a stored name left in clear on the unfixed code; fixed in Task 2",
)
@pytest.mark.parametrize("ext", [".txt", ".xlsx"])
def test_stored_name_overlap(tmp_path: Path, db: Path, ext: str) -> None:
    """T3 (AC1, AC9; 10.15's F1 and Q4): the name added partly overlaps a
    stored name accepted without review; nothing of either stays in clear."""
    _store(tmp_path, db, VZ)
    source, output = tmp_path / f"in{ext}", tmp_path / f"out_in{ext}"
    if ext == ".txt":
        _write(source, "Vorlan Zorbal Quentrix a signé.")
    else:
        _write(source, [["Nom", "Note"], ["ok", "Vorlan Zorbal Quentrix a signé"]])
    run = _process(tmp_path, db, source, output, [(VZ, "PERSON")], stdin=ADD)
    _pre(run.code == 0, "exit code 0")
    written = _flat(_read(output))
    assert "Quentrix" not in written and "Vorlan" not in written  # expected failure
    assert len(_rows_of(db, ZQ, "PERSON")) == 1
    pseudonym = _pseudonym(db, ZQ, "PERSON")
    if ext == ".txt":
        assert written == f"{pseudonym} a signé."
    else:
        assert written == f"Nom\tNote\nok\t{pseudonym} a signé"


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason="Story 10.18: reproduces the false line in place of the approved one, and no question on the unfixed code; fixed in Task 2",
)
def test_answer_n_writes_as_today(tmp_path: Path, db: Path) -> None:
    """T4 (AC1, AC2): "n" writes the file as today; the new line and the
    question are shown, the false line is not."""
    source, output = tmp_path / "in.txt", tmp_path / "out_in.txt"
    _write(source, NONE_TXT)
    run = _process(tmp_path, db, source, output, stdin=b"n\n")
    _pre(run.code == 0, "exit code 0")
    written = _read(output)
    assert NO_NAMES_FOUND in run.shown  # expected failure
    assert FALSE_LINE not in run.shown
    assert QUESTION in run.shown
    assert written == NONE_TXT
    assert _rows(db) == []


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason="Story 10.18: reproduces no question, so a missed name left in clear on the unfixed code; fixed in Task 2",
)
def test_hyphen_compound_not_a_place(tmp_path: Path, db: Path) -> None:
    """T5 (AC1; 10.15 Q3): a name inside a hyphenated word is not a place."""
    source, output = tmp_path / "in.txt", tmp_path / "out_in.txt"
    _write(source, "Le contrat avec Saint-Zorbal et Zorbal est signé.")
    run = _process(tmp_path, db, source, output, stdin=b"y\nZorbal\nPERSON\nn\n")
    _pre(run.code == 0, "exit code 0")
    written = _read(output)
    assert "et Zorbal est" not in written  # expected failure
    assert len(_rows_of(db, "Zorbal", "PERSON")) == 1
    pseudonym = _pseudonym(db, "Zorbal", "PERSON")
    assert written == f"Le contrat avec Saint-Zorbal et {pseudonym} est signé."


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason="Story 10.18: reproduces the false line in place of the selected-types line on the unfixed code; fixed in Task 2",
)
@pytest.mark.parametrize("case", ["nothing_of_the_types", "stored_and_left_out"])
def test_entity_types_lines(tmp_path: Path, db: Path, case: str) -> None:
    """T6 (Q5): with --entity-types, the lines speak of the selected types."""
    source, output = tmp_path / "in.txt", tmp_path / "out_in.txt"
    if case == "nothing_of_the_types":
        content, names, line = (
            "Le contrat avec Zorbaville est signé.",
            [(ZV, "ORG")],
            SELECTED_NO_NAMES_FOUND,
        )
        expected = content
    else:
        vm = _store(tmp_path, db, VM)
        content, names, line = (
            "Vorlan Mestrix valide. Le contrat avec Zorbaville est signé.",
            [(VM, "PERSON"), (ZV, "ORG")],
            SELECTED_NO_NEW_NAMES,
        )
        expected = content.replace(VM, vm)
    _write(source, content)
    run = _process(
        tmp_path,
        db,
        source,
        output,
        names,
        stdin=b"n\n",
        extra=["--entity-types", "PERSON"],
    )
    _pre(run.code == 0, "exit code 0")
    if case == "stored_and_left_out":
        _pre(AUTO_ACCEPTED in run.shown, "the stored name is accepted without review")
    written = _read(output)
    assert line in run.shown  # expected failure
    assert [
        text for text in (NO_NAMES_FOUND, NO_NEW_NAMES, FALSE_LINE) if text in run.shown
    ] == []
    assert written == expected


# --- AC3: batch one file at a time -------------------------------------------


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason="Story 10.18: reproduces no question in a batch, a missed name in clear on the unfixed code; fixed in Task 2",
)
def test_batch_asks_each_file(tmp_path: Path, db: Path) -> None:
    """T7 (AC3, Q3): the question is asked for each file."""
    folder, out_folder = tmp_path / "docs", tmp_path / "out"
    folder.mkdir()
    text = "Le contrat avec Zorbal Quentrix est signé."
    for stem in ("a", "b"):
        _write(folder / f"{stem}.txt", text)
    run = _batch(
        tmp_path, db, folder, out_folder, stdin=b"y\nZorbal Quentrix\nPERSON\nn\nn\n"
    )
    _pre(run.code == 0, "exit code 0")
    written_a = _read(out_folder / "a_pseudonymized.txt")
    written_b = _read(out_folder / "b_pseudonymized.txt")
    assert ZQ not in written_a  # expected failure
    assert written_b == text
    assert run.shown.count(QUESTION) == 2


# --- AC4: no answer can be read ----------------------------------------------


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason="Story 10.18: reproduces a document written with no answer read on the unfixed code; fixed in Task 2",
)
@pytest.mark.parametrize("ext", [".txt", ".xlsx"])
def test_no_answer_closed_stdin(tmp_path: Path, db: Path, ext: str) -> None:
    """T8 (AC4, Q4): a closed input stops the run with the approved line."""
    source, output = tmp_path / f"in{ext}", tmp_path / f"out_in{ext}"
    _write(source, NONE_TXT if ext == ".txt" else NONE_SHEET)
    run = _process(tmp_path, db, source, output)
    assert run.code == 1  # expected failure
    assert not output.exists()
    assert _rows(db) == []
    assert FAILED in run.shown and NO_ANSWER in run.shown
    assert SUCCESS not in run.shown


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason="Story 10.18: reproduces a document written with no answer read on the unfixed code; fixed in Task 2",
)
def test_no_answer_input_runs_out(tmp_path: Path, db: Path) -> None:
    """T9 (AC4): an input that runs out inside the manual flow stops the run."""
    source, output = tmp_path / "in.txt", tmp_path / "out_in.txt"
    _write(source, NONE_TXT)
    run = _process(tmp_path, db, source, output, stdin=b"y\nZorbal Quentrix\n")
    assert run.code == 1  # expected failure
    assert not output.exists()
    assert _rows(db) == []
    assert NO_ANSWER in run.shown


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason="Story 10.18: reproduces a batch writing every file with no answer read on the unfixed code; fixed in Task 2",
)
def test_no_answer_batch_each_file_fails(tmp_path: Path, db: Path) -> None:
    """T10 (AC4): in a batch, each file without an answer fails and the batch
    goes on to the next one."""
    folder, out_folder = tmp_path / "docs", tmp_path / "out"
    folder.mkdir()
    for stem in ("a", "b"):
        _write(folder / f"{stem}.txt", "Le contrat avec Zorbal Quentrix est signé.")
    run = _batch(tmp_path, db, folder, out_folder)
    assert run.code == 1  # expected failure
    assert not (out_folder / "a_pseudonymized.txt").exists()
    assert not (out_folder / "b_pseudonymized.txt").exists()
    reason = NO_ANSWER.removeprefix("Error: ")
    assert f"a.txt: {reason}" in run.shown and f"b.txt: {reason}" in run.shown
    assert _rows(db) == []


# --- AC5: Ctrl+C ----------------------------------------------------------------


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason="Story 10.18: reproduces no question to cancel, the document written on the unfixed code; fixed in Task 2",
)
def test_ctrl_c_at_question(tmp_path: Path, db: Path) -> None:
    """T12 (AC5): Ctrl+C at the question cancels at once; nothing written."""
    source, output = tmp_path / "in.txt", tmp_path / "out_in.txt"
    _write(source, NONE_TXT)
    run = _process(tmp_path, db, source, output, ctrl_c_at=1)
    _pre(run.code == 0, "exit code 0")
    assert not output.exists()  # expected failure
    assert CANCELLED in run.shown
    assert _rows(db) == []


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason="Story 10.18: reproduces no question to cancel, the document written on the unfixed code; fixed in Task 2",
)
def test_ctrl_c_while_adding(tmp_path: Path, db: Path) -> None:
    """T13 (AC5): Ctrl+C at 'Entity text' cancels; nothing written or stored."""
    source, output = tmp_path / "in.txt", tmp_path / "out_in.txt"
    _write(source, NONE_TXT)
    run = _process(tmp_path, db, source, output, stdin=b"y\n", ctrl_c_at=2)
    _pre(run.code == 0, "exit code 0")
    assert not output.exists()  # expected failure
    assert CANCELLED in run.shown
    assert _rows(db) == []


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason="Story 10.18: reproduces no question to cancel, the second file written on the unfixed code; fixed in Task 2",
)
def test_ctrl_c_in_batch_keeps_written(tmp_path: Path, db: Path) -> None:
    """T14 (AC5): Ctrl+C at the second file's question cancels the batch; the
    first file's output stays."""
    folder, out_folder = tmp_path / "docs", tmp_path / "out"
    folder.mkdir()
    for stem in ("a", "b"):
        _write(folder / f"{stem}.txt", "Le contrat avec Zorbal Quentrix est signé.")
    run = _batch(tmp_path, db, folder, out_folder, stdin=b"n\n", ctrl_c_at=2)
    _pre(run.code == 0, "exit code 0")
    _pre((out_folder / "a_pseudonymized.txt").is_file(), "a's output written")
    assert not (out_folder / "b_pseudonymized.txt").exists()  # expected failure
    assert BATCH_CANCELLED in run.shown


# --- AC8: the CLI reference -------------------------------------------------------


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason="Story 10.18: reproduces the CLI reference documenting --stop-on-error and a --workers default of 1 on the unfixed code; fixed in Task 2",
)
def test_cli_reference_options() -> None:
    """T15 (AC8, Q6 and the --stop-on-error fix): both CLI references name
    the real batch options and the real --workers default."""
    batch = typer.main.get_command(app).commands["batch"]  # type: ignore[attr-defined]
    options = {
        option
        for parameter in batch.params
        for option in (*parameter.opts, *parameter.secondary_opts)
    }
    _pre(
        "--no-continue-on-error" in options
        and "--workers" in options
        and "--stop-on-error" not in options,
        "the batch command's options",
    )
    for name in ("CLI-REFERENCE.md", "CLI-REFERENCE.fr.md"):
        lines = (REPO / "docs" / name).read_text(encoding="utf-8").splitlines()
        _pre(len(lines) > 192, f"{name} read")
        stale = [
            number for number, line in enumerate(lines, 1) if "--stop-on-error" in line
        ]
        assert stale == []  # expected failure
        assert any(
            line.startswith("| `--no-continue-on-error` | | | ") for line in lines
        )
        assert "gdpr-pseudo batch ./documents/ --no-continue-on-error" in lines
        assert any(line.startswith("| `--workers` | `-w` | 4 | ") for line in lines)


# --- Guards: what must not change (AC7, AC14) ------------------------------------


def test_review_opened_no_question(tmp_path: Path, db: Path) -> None:
    """G1 (AC7, Q2): after a review where every name is decided, no question."""
    source, output = tmp_path / "in.txt", tmp_path / "out_in.txt"
    text = "Vorlan Mestrix valide. Le contrat est signé."
    _write(source, text)
    run = _process(
        tmp_path,
        db,
        source,
        output,
        [(VM, "PERSON")],
        stdin=b"y\n",
        keys=[ENTER, SPACE],
    )
    assert run.code == 0
    assert QUESTION not in run.shown
    assert _read(output) == text.replace(VM, _pseudonym(db, VM, "PERSON"))


class _NoDetection:
    def detect_entities(self, text: str) -> list[DetectedEntity]:
        return []


def _record_prompts(asked: list[str]) -> Any:
    def ask(prompt: Any, *args: Any, **kwargs: Any) -> Any:
        asked.append(str(prompt))
        raise AssertionError("a prompt was called")

    return ask


@pytest.mark.parametrize("ext", [".txt", ".xlsx"])
def test_skip_validation_never_asks(tmp_path: Path, db: Path, ext: str) -> None:
    """G2 (AC7, AC14): the desktop's call (skip_validation=True) never asks."""
    source, output = tmp_path / f"in{ext}", tmp_path / f"out_in{ext}"
    content: str | Rows = NONE_TXT if ext == ".txt" else NONE_SHEET
    _write(source, content)
    asked: list[str] = []
    with (
        patch.object(
            DocumentProcessor, "_get_detector", new=lambda self: _NoDetection()
        ),
        patch.object(Confirm, "ask", new=_record_prompts(asked)),
        patch.object(Prompt, "ask", new=_record_prompts(asked)),
    ):
        result = DocumentProcessor(str(db), PASSPHRASE).process_document(
            str(source), str(output), skip_validation=True
        )
    assert result.success
    assert asked == []
    assert _read(output) == content


def test_parallel_worker_never_asks(tmp_path: Path, db: Path) -> None:
    """G3 (AC7, AC14): the parallel batch's worker never asks."""
    source, output = tmp_path / "in.txt", tmp_path / "out_in.txt"
    _write(source, NONE_TXT)
    asked: list[str] = []
    with (
        patch.object(
            DocumentProcessor, "_get_detector", new=lambda self: _NoDetection()
        ),
        patch.object(Confirm, "ask", new=_record_prompts(asked)),
        patch.object(Prompt, "ask", new=_record_prompts(asked)),
    ):
        outcome = _process_single_document_worker(
            (str(source), str(output), str(db), PASSPHRASE, "neutral", "spacy", None)
        )
    assert outcome["success"]
    assert asked == []
    assert _read(output) == NONE_TXT
