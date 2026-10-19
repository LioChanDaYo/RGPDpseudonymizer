"""Story 10.15: a name corrected ([E]), retyped ([T]) or added ([A]) in the
command-line review is replaced at every place it belongs, in a text file and
in Excel and CSV cells (T1-T15), and what must not change (G1-G5). G6-G8
guard the gaps QA found (TEST-001: the cells through both restarts of the
review; TEST-002: the same span in two cells).

Written by the BMAD dev agent (Claude). Each test drives the real
``gdpr-pseudo process`` command through ``CliRunner``, on a database made by
the real ``init`` command. Detection is replaced by a fake that finds the
names given; the keys, the typed texts and the yes/no answers come from lists.
The assertions read the written file (text or cells) and the stored rows.
Preconditions call ``pytest.fail("PRE: ...")``, never ``assert``; the first
``assert`` after them is the expected failure on the unfixed code. No test
asserts the wording of a new message. Invented names only; every other word
is a main-corpus word. No spaCy model, no Qt.
"""

from __future__ import annotations

import csv
import dataclasses
import re
from collections.abc import Iterable, Sequence
from contextlib import AbstractContextManager, ExitStack
from pathlib import Path
from typing import Any
from unittest.mock import patch

import openpyxl
import pytest
import readchar
from rich.prompt import Confirm, Prompt
from typer.testing import CliRunner

import gdpr_pseudonymizer.core.document_processor as document_processor_module
from gdpr_pseudonymizer.cli.main import app
from gdpr_pseudonymizer.core.document_processor import DocumentProcessor
from gdpr_pseudonymizer.data.database import open_database
from gdpr_pseudonymizer.data.repositories.mapping_repository import (
    SQLiteMappingRepository,
)
from gdpr_pseudonymizer.nlp.entity_detector import DetectedEntity
from gdpr_pseudonymizer.validation.models import (
    EntityReview,
    EntityReviewState,
    ValidationSession,
)

PASSPHRASE = "zorbalia_pass_A_123!"
ENV = {"GDPR_PSEUDO_PASSPHRASE": PASSPHRASE}
SHEET = "Zorbalia"
SHEETS = [".xlsx", ".csv"]

ZQ = "Zorbal Quentrix"
ZQS = "Zorbal Quentrix signé"
VM = "Vorlan Mestrix"
VZ = "Vorlan Zorbal"
BT = "Brimel Testar"
ZV = "Zorbaville"
CUSTOM = "Xantor Vellion"

ENTER = "\r"
SPACE = " "
CTRL_C = readchar.key.CTRL_C
MAX_KEYS = 60

SUCCESS = "Processing Successful"
AUTO_ACCEPTED = "Auto-accepted"
QUIT_QUESTION = "Quit validation? Progress will be lost."
READY_QUESTION = "Ready to process document with validated entities?"

_ANSI = re.compile(r"\x1b\[[0-9;?]*[A-Za-z]")

Names = Sequence[tuple[str, str]]
Rows = list[list[str]]


def _pre(condition: bool, what: str) -> None:
    """A precondition: a red failure on every leg, never an AssertionError."""
    if not condition:
        pytest.fail(f"PRE: {what}")


@pytest.fixture(autouse=True)
def _settings(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    """The command reads no user or project settings file."""
    monkeypatch.setenv("HOME", str(tmp_path))
    monkeypatch.setenv("USERPROFILE", str(tmp_path))
    monkeypatch.chdir(tmp_path)


@pytest.fixture
def db(tmp_path: Path) -> Path:
    """A new database, made by the real ``init`` command."""
    path = tmp_path / "mappings.db"
    result = CliRunner().invoke(
        app, ["init", "--db", str(path)], env=ENV, catch_exceptions=False
    )
    _pre(result.exit_code == 0 and path.is_file(), "init made the database")
    return path


class FakeDetector:
    """Finds every occurrence of each (name, type) given, and drops an
    occurrence that lies inside a longer one it found. With ``first_only``,
    each name is found once, at its first occurrence over all calls."""

    def __init__(self, names: Names, first_only: bool = False) -> None:
        self.names = list(names)
        self.first_only = first_only
        self.seen: set[tuple[str, str]] = set()

    def detect_entities(self, text: str) -> list[DetectedEntity]:
        found: list[DetectedEntity] = []
        for name, entity_type in self.names:
            if self.first_only and (name, entity_type) in self.seen:
                continue
            start = text.find(name)
            while start != -1:
                found.append(
                    DetectedEntity(
                        text=name,
                        entity_type=entity_type,
                        start_pos=start,
                        end_pos=start + len(name),
                        confidence=0.9,
                        source="spacy",
                    )
                )
                if self.first_only:
                    self.seen.add((name, entity_type))
                    break
                start = text.find(name, start + 1)
        return [
            entity
            for entity in found
            if not any(
                other.start_pos <= entity.start_pos
                and entity.end_pos <= other.end_pos
                and other.end_pos - other.start_pos > entity.end_pos - entity.start_pos
                for other in found
            )
        ]


class Keys:
    """``readchar.readkey``: the keys in order, then Space; Ctrl+C raises
    ``KeyboardInterrupt``, as the real function does."""

    def __init__(self, keys: Sequence[str]) -> None:
        self.keys = list(keys)
        self.read = 0

    def __call__(self) -> str:
        self.read += 1
        if self.read > MAX_KEYS:
            raise RuntimeError("more than 60 keys read")
        key = self.keys.pop(0) if self.keys else SPACE
        if key == CTRL_C:
            raise KeyboardInterrupt
        return key


class Typed:
    """``Prompt.ask``: the typed texts in order."""

    def __init__(self, texts: Sequence[str]) -> None:
        self.texts = list(texts)

    def __call__(self, prompt: Any, *args: Any, **kwargs: Any) -> str:
        if not self.texts:
            raise RuntimeError("no typed text left")
        return self.texts.pop(0)


class Answers:
    """``Confirm.ask``: the answers in order, then yes; records each question."""

    def __init__(self, answers: Sequence[bool] = ()) -> None:
        self.answers = list(answers)
        self.asked: list[str] = []

    def __call__(self, prompt: Any, *args: Any, **kwargs: Any) -> bool:
        self.asked.append(str(prompt))
        return self.answers.pop(0) if self.answers else True


@dataclasses.dataclass
class Run:
    exit_code: int
    shown: str
    asked: list[str]


def _process(
    source: Path,
    output: Path,
    db: Path,
    names: Names,
    keys: Sequence[str],
    typed: Sequence[str] = (),
    *,
    first_only: bool = False,
    answers: Sequence[bool] = (),
    extra: Iterable[AbstractContextManager[Any]] = (),
) -> Run:
    """One real ``gdpr-pseudo process`` run."""
    detector = FakeDetector(names, first_only)
    confirm = Answers(answers)
    with ExitStack() as stack:
        stack.enter_context(
            patch.object(DocumentProcessor, "_get_detector", new=lambda self: detector)
        )
        stack.enter_context(
            patch.object(
                DocumentProcessor,
                "_detect_and_filter_entities",
                new=lambda self, text, entity_type_filter=None: (
                    detector.detect_entities(text)
                ),
            )
        )
        stack.enter_context(patch.object(readchar, "readkey", new=Keys(keys)))
        stack.enter_context(patch.object(Prompt, "ask", new=Typed(typed)))
        stack.enter_context(patch.object(Confirm, "ask", new=confirm))
        for context in extra:
            stack.enter_context(context)
        result = CliRunner().invoke(
            app,
            ["process", str(source), "--db", str(db), "-o", str(output)],
            env=ENV,
            catch_exceptions=False,
        )
    return Run(result.exit_code, _ANSI.sub("", result.output), confirm.asked)


def _write_sheet(path: Path, rows: Rows) -> None:
    if path.suffix == ".xlsx":
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = SHEET
        for row in rows:
            sheet.append(row)
        workbook.save(path)
        workbook.close()
    else:
        with path.open("w", encoding="utf-8", newline="") as handle:
            csv.writer(handle).writerows(rows)


def _cells(path: Path) -> Rows:
    """The written cells, as rows of strings."""
    _pre(path.is_file(), "the output file is written")
    if path.suffix == ".xlsx":
        workbook = openpyxl.load_workbook(path)
        try:
            return [
                ["" if value is None else str(value) for value in row]
                for row in workbook.worksheets[0].iter_rows(values_only=True)
            ]
        finally:
            workbook.close()
    with path.open(encoding="utf-8", newline="") as handle:
        return [list(row) for row in csv.reader(handle)]


def _cell(rows: Rows, ref: str) -> str:
    """A cell by its Excel name ("B2": row 2, column 2)."""
    column, row = ord(ref[0]) - ord("A"), int(ref[1:]) - 1
    _pre(row < len(rows) and column < len(rows[row]), f"cell {ref} exists")
    return rows[row][column]


def _flat(rows: Rows) -> str:
    return "\n".join("\t".join(row) for row in rows)


def _text(path: Path) -> str:
    _pre(path.is_file(), "the output file is written")
    return path.read_text(encoding="utf-8")


def _rows(db: Path) -> list[tuple[str, str, str, str]]:
    """Every stored row: (full_name, entity_type, pseudonym_full, theme)."""
    with open_database(str(db), PASSPHRASE) as session:
        return [
            (row.full_name, row.entity_type, row.pseudonym_full, row.theme)
            for row in SQLiteMappingRepository(session).find_all()
        ]


def _pseudonym(db: Path, name: str, entity_type: str) -> str:
    """The pseudonym of the one row stored for (name, type)."""
    found = [row for row in _rows(db) if row[:2] == (name, entity_type)]
    _pre(len(found) == 1, f"one row stored for ({name}, {entity_type})")
    return found[0][2]


def _files(tmp_path: Path, ext: str, tag: str = "") -> tuple[Path, Path]:
    return tmp_path / f"in{tag}{ext}", tmp_path / f"out{tag}{ext}"


# --- AC1 and AC2: a corrected or retyped name keeps its cell ----------------


@pytest.mark.parametrize("ext", SHEETS)
def test_type_change_keeps_cell(tmp_path: Path, db: Path, ext: str) -> None:
    """T1 (AC1, [T]; QA's #2): a name retyped with [T] is written in its cell."""
    source, output = _files(tmp_path, ext)
    _write_sheet(source, [["Nom", "Note"], [VM, "valide"], [ZQ, "signé"]])
    run = _process(
        source,
        output,
        db,
        [(VM, "PERSON"), (ZQ, "ORG")],
        [ENTER, SPACE, "t"],
        ["PERSON"],
    )
    _pre(run.exit_code == 0, "exit code 0")
    cells = _cells(output)
    pseudonym = _pseudonym(db, ZQ, "PERSON")
    _pre(_cell(cells, "A2") == _pseudonym(db, VM, "PERSON"), "A2 holds VM's row")
    assert _cell(cells, "A3") == pseudonym  # expected failure
    assert ZQ not in _flat(cells)


@pytest.mark.parametrize("ext", SHEETS)
def test_correction_keeps_cell(tmp_path: Path, db: Path, ext: str) -> None:
    """T2 (AC1, [E]; QA's #2): a name corrected with [E] is written in its cell."""
    source, output = _files(tmp_path, ext)
    _write_sheet(
        source,
        [
            ["Nom", "Note"],
            ["ok", "Zorbal Quentrix signé le contrat"],
            ["ok", "chez Zorbaville"],
        ],
    )
    run = _process(
        source,
        output,
        db,
        [(ZQS, "PERSON"), (ZV, "ORG")],
        [ENTER, "e", SPACE],
        [ZQ],
    )
    _pre(run.exit_code == 0, "exit code 0")
    cells = _cells(output)
    pseudonym = _pseudonym(db, ZQ, "PERSON")
    _pre(_cell(cells, "B3") == f"chez {_pseudonym(db, ZV, 'ORG')}", "B3 control")
    assert _cell(cells, "B2") == f"{pseudonym} signé le contrat"  # expected failure


@pytest.mark.parametrize("ext", SHEETS)
def test_correction_then_type_change_keeps_cell(
    tmp_path: Path, db: Path, ext: str
) -> None:
    """T3 (AC1, [E] then [T]): the corrected and retyped name keeps its cell."""
    source, output = _files(tmp_path, ext)
    _write_sheet(
        source,
        [
            ["Nom", "Note"],
            [VM, "valide"],
            ["ok", "Zorbal Quentrix signé le contrat"],
            ["ok", "chez Zorbaville"],
        ],
    )
    run = _process(
        source,
        output,
        db,
        [(VM, "PERSON"), (ZQS, "ORG"), (ZV, "ORG")],
        [ENTER, SPACE, "e", "p", "t", SPACE],
        [ZQ, "PERSON"],
    )
    _pre(run.exit_code == 0, "exit code 0")
    cells = _cells(output)
    pseudonym = _pseudonym(db, ZQ, "PERSON")
    _pre(_cell(cells, "B4") == f"chez {_pseudonym(db, ZV, 'ORG')}", "B4 control")
    assert _cell(cells, "B3") == f"{pseudonym} signé le contrat"  # expected failure


@pytest.mark.parametrize("ext", SHEETS)
def test_custom_after_correction_keeps_cell(tmp_path: Path, db: Path, ext: str) -> None:
    """T4 (AC1, [C] after [E]): the corrected name with its own pseudonym keeps
    its cell."""
    source, output = _files(tmp_path, ext)
    _write_sheet(
        source,
        [
            ["Nom", "Note"],
            ["ok", "Zorbal Quentrix signé le contrat"],
            ["ok", "avec Brimel Testar"],
        ],
    )
    run = _process(
        source,
        output,
        db,
        [(ZQS, "PERSON"), (BT, "PERSON")],
        [ENTER, "e", "p", "c", SPACE],
        [ZQ, CUSTOM],
    )
    _pre(run.exit_code == 0, "exit code 0")
    cells = _cells(output)
    _pre(_pseudonym(db, ZQ, "PERSON") == CUSTOM, "the custom row is stored")
    _pre(_cell(cells, "B3") == f"avec {_pseudonym(db, BT, 'PERSON')}", "B3 control")
    assert _cell(cells, "B2") == f"{CUSTOM} signé le contrat"  # expected failure


@pytest.mark.parametrize("ext", SHEETS)
def test_correction_located_in_its_cell(tmp_path: Path, db: Path, ext: str) -> None:
    """T9 (AC2): a correction is located in its own cell, not in the joined
    review text, where the same text stands earlier."""
    source, output = _files(tmp_path, ext)
    _write_sheet(
        source, [[ZQ, "Note"], ["ok", "chez Zorbal Quentrix signé le contrat"]]
    )
    run = _process(
        source,
        output,
        db,
        [(ZQ, "ORG"), (ZQS, "PERSON")],
        [ENTER, "e", SPACE],
        [ZQ],
    )
    _pre(run.exit_code == 0, "exit code 0")
    cells = _cells(output)
    pseudonym = _pseudonym(db, ZQ, "PERSON")
    _pre(_cell(cells, "A1") == _pseudonym(db, ZQ, "ORG"), "A1 holds the ORG row")
    assert (
        _cell(cells, "B2") == f"chez {pseudonym} signé le contrat"
    )  # expected failure


def test_sent_entity_fallback_keeps_cell() -> None:
    """T12 (AC1): the corrected-name fallback of ``_sent_entity`` keeps the cell."""
    label = f"{SHEET}!B2"
    review = EntityReview(
        entity=DetectedEntity(
            text=ZQS,
            entity_type="PERSON",
            start_pos=0,
            end_pos=len(ZQS),
            confidence=0.9,
            source="spacy",
            context_label=label,
        ),
        state=EntityReviewState.MODIFIED,
        user_modification=ZQ,
    )
    _pre(review.modified_entity is None, "no modified_entity")
    sent = ValidationSession._sent_entity(review)
    _pre(sent.text == ZQ, "the fallback sends the corrected text")
    assert sent.context_label == label  # expected failure


# --- AC3 and AC14: a name added with [A] is replaced at every place ---------


@pytest.mark.parametrize("ext", SHEETS)
def test_add_on_spreadsheet(tmp_path: Path, db: Path, ext: str) -> None:
    """T5 (AC3; the PO's probe): a name added with [A] is written in its cell."""
    source, output = _files(tmp_path, ext)
    _write_sheet(
        source,
        [["Nom", "Note"], [VM, "valide"], ["ok", "contrat avec Zorbal Quentrix"]],
    )
    run = _process(
        source, output, db, [(VM, "PERSON")], [ENTER, "a", SPACE], [ZQ, "PERSON"]
    )
    _pre(run.exit_code == 0, "exit code 0")
    cells = _cells(output)
    pseudonym = _pseudonym(db, ZQ, "PERSON")
    _pre(_cell(cells, "A2") == _pseudonym(db, VM, "PERSON"), "A2 control")
    assert _cell(cells, "B3") == f"contrat avec {pseudonym}"  # expected failure


@pytest.mark.parametrize("ext", [".txt", ".xlsx"])
def test_add_name_at_two_places(tmp_path: Path, db: Path, ext: str) -> None:
    """T6 (AC3; the PO's two-place cases): both places of an added name."""
    source, output = _files(tmp_path, ext)
    text = (
        "Vorlan Mestrix valide. Le contrat avec Zorbal Quentrix. "
        "Signé : Zorbal Quentrix."
    )
    if ext == ".txt":
        source.write_text(text, encoding="utf-8")
    else:
        _write_sheet(
            source, [["Nom", "Note"], [VM, "avec Zorbal Quentrix"], [ZQ, "signé"]]
        )
    run = _process(
        source, output, db, [(VM, "PERSON")], [ENTER, "a", SPACE], [ZQ, "PERSON"]
    )
    _pre(run.exit_code == 0, "exit code 0")
    pseudonym = _pseudonym(db, ZQ, "PERSON")
    vm = _pseudonym(db, VM, "PERSON")
    if ext == ".txt":
        written = _text(output)
        _pre(written.startswith(f"{vm} valide."), "VM replaced")
        assert ZQ not in written  # expected failure
        assert written == text.replace(VM, vm).replace(ZQ, pseudonym)
    else:
        cells = _cells(output)
        _pre(_cell(cells, "A2") == vm, "VM replaced")
        assert ZQ not in _flat(cells)  # expected failure
        assert _cell(cells, "B2") == f"avec {pseudonym}"
        assert _cell(cells, "A3") == pseudonym


@pytest.mark.parametrize("ext", [".txt", ".xlsx"])
def test_add_keeps_found_place_decision(tmp_path: Path, db: Path, ext: str) -> None:
    """T7 (AC3, AC14 (b), 10.10): a place the detector found keeps its own
    decision (here rejected); the other place is added."""
    source, output = _files(tmp_path, ext)
    text = "Zorbal Quentrix valide. Le contrat avec Zorbal Quentrix."
    if ext == ".txt":
        source.write_text(text, encoding="utf-8")
    else:
        _write_sheet(
            source,
            [["Nom", "Note"], [ZQ, "valide"], ["ok", "contrat avec Zorbal Quentrix"]],
        )
    run = _process(
        source,
        output,
        db,
        [(ZQ, "PERSON")],
        [ENTER, "a", "r"],
        [ZQ, "PERSON"],
        first_only=True,
    )
    _pre(run.exit_code == 0, "exit code 0")
    pseudonym = _pseudonym(db, ZQ, "PERSON")
    if ext == ".txt":
        written = _text(output)
        assert written.endswith(f"Le contrat avec {pseudonym}.")  # expected failure
        assert written == f"{ZQ} valide. Le contrat avec {pseudonym}."
    else:
        cells = _cells(output)
        assert _cell(cells, "B3") == f"contrat avec {pseudonym}"  # expected failure
        assert _cell(cells, "A2") == ZQ


def test_add_on_found_places_adds_nothing(tmp_path: Path, db: Path) -> None:
    """T13 (AC14 (b), Q2's case): [A] of a name found at every place adds
    nothing; the user's reject of the found places holds."""
    source, output = _files(tmp_path, ".txt")
    text = "Zorbal Quentrix valide. Le contrat avec Zorbal Quentrix."
    source.write_text(text, encoding="utf-8")
    run = _process(
        source, output, db, [(ZQ, "PERSON")], [ENTER, "a", "r"], [ZQ, "PERSON"]
    )
    _pre(run.exit_code == 0, "exit code 0")
    written = _text(output)
    assert _rows(db) == []  # expected failure
    assert written == text


@pytest.mark.parametrize("ext", [".txt", ".xlsx"])
def test_add_overlapping_a_written_name(tmp_path: Path, db: Path, ext: str) -> None:
    """T10 (AC14 (d), Q4): an added place that partly overlaps a found name is
    widened; the whole span takes the added name's pseudonym."""
    source, output = _files(tmp_path, ext)
    if ext == ".txt":
        source.write_text("Vorlan Zorbal Quentrix a signé.", encoding="utf-8")
    else:
        _write_sheet(
            source, [["Nom", "Note"], ["ok", "Vorlan Zorbal Quentrix a signé"]]
        )
    run = _process(
        source, output, db, [(VZ, "PERSON")], [ENTER, "a", SPACE], [ZQ, "PERSON"]
    )
    _pre(run.exit_code == 0, "exit code 0")
    pseudonym = _pseudonym(db, ZQ, "PERSON")
    _pseudonym(db, VZ, "PERSON")
    if ext == ".txt":
        written = _text(output)
        assert "Quentrix" not in written and "Vorlan" not in written  # expected failure
        assert written == f"{pseudonym} a signé."
    else:
        cells = _cells(output)
        flat = _flat(cells)
        assert "Quentrix" not in flat and "Vorlan" not in flat  # expected failure
        assert _cell(cells, "B2") == f"{pseudonym} a signé"


def test_add_whole_words_only(tmp_path: Path, db: Path) -> None:
    """T11 (AC14 (a), Q3): only whole words are places; a hyphen joins words,
    an elided article does not."""
    source, output = _files(tmp_path, ".txt")
    source.write_text(
        "Vorlan Mestrix a signé avec Saint-Zorbal, Zorbalix, l'Zorbal et Zorbal.",
        encoding="utf-8",
    )
    run = _process(
        source, output, db, [(VM, "PERSON")], [ENTER, "a", SPACE], ["Zorbal", "PERSON"]
    )
    _pre(run.exit_code == 0, "exit code 0")
    pseudonym = _pseudonym(db, "Zorbal", "PERSON")
    vm = _pseudonym(db, VM, "PERSON")
    written = _text(output)
    assert written.endswith(f" et {pseudonym}.")  # expected failure
    assert written == (
        f"{vm} a signé avec Saint-Zorbal, Zorbalix, l'{pseudonym} et {pseudonym}."
    )


@pytest.mark.parametrize("ext", [".txt", ".xlsx"])
def test_add_overlapping_a_stored_name(tmp_path: Path, db: Path, ext: str) -> None:
    """T14 (AC14 (d), Q4; the PO's F1): the widening also covers a name stored
    in the mapping database and accepted without review."""
    first, first_out = _files(tmp_path, ext, "1")
    if ext == ".txt":
        first.write_text("Vorlan Zorbal valide.", encoding="utf-8")
    else:
        _write_sheet(first, [["Nom"], [VZ]])
    run = _process(first, first_out, db, [(VZ, "PERSON")], [ENTER, SPACE])
    _pre(run.exit_code == 0, "first run: exit code 0")
    stored = _pseudonym(db, VZ, "PERSON")

    source, output = _files(tmp_path, ext, "2")
    if ext == ".txt":
        source.write_text(
            "Vorlan Zorbal Quentrix a signé avec Brimel Testar.", encoding="utf-8"
        )
    else:
        _write_sheet(
            source,
            [
                ["Nom", "Note"],
                ["ok", "Vorlan Zorbal Quentrix a signé avec Brimel Testar"],
            ],
        )
    run = _process(
        source,
        output,
        db,
        [(VZ, "PERSON"), (BT, "PERSON")],
        [ENTER, "a", SPACE],
        [ZQ, "PERSON"],
    )
    _pre(run.exit_code == 0, "second run: exit code 0")
    _pre(AUTO_ACCEPTED in run.shown, "the stored name is accepted without review")
    _pre(_pseudonym(db, VZ, "PERSON") == stored, "the stored row is unchanged")
    pseudonym = _pseudonym(db, ZQ, "PERSON")
    bt = _pseudonym(db, BT, "PERSON")
    if ext == ".txt":
        written = _text(output)
        assert "Quentrix" not in written and "Vorlan" not in written  # expected failure
        assert written == f"{pseudonym} a signé avec {bt}."
    else:
        cells = _cells(output)
        flat = _flat(cells)
        assert "Quentrix" not in flat and "Vorlan" not in flat  # expected failure
        assert _cell(cells, "B2") == f"{pseudonym} a signé avec {bt}"


def test_add_twice_writes_once(tmp_path: Path, db: Path) -> None:
    """T15 (AC14 (b)): the same [A] twice adds each place once; both places
    are written."""
    source, output = _files(tmp_path, ".txt")
    text = (
        "Vorlan Mestrix valide. Le contrat avec Zorbal Quentrix. "
        "Signé : Zorbal Quentrix."
    )
    source.write_text(text, encoding="utf-8")
    run = _process(
        source,
        output,
        db,
        [(VM, "PERSON")],
        [ENTER, "a", "a", SPACE],
        [ZQ, "PERSON", ZQ, "PERSON"],
    )
    _pre(run.exit_code == 0, "exit code 0")
    pseudonym = _pseudonym(db, ZQ, "PERSON")
    vm = _pseudonym(db, VM, "PERSON")
    written = _text(output)
    assert ZQ not in written  # expected failure
    assert written == text.replace(VM, vm).replace(ZQ, pseudonym)


# --- AC4: a name that matches no cell stops the run before any row ----------


def _lose_cell(how: str, changed: list[str]) -> AbstractContextManager[Any]:
    """Wrap the real review: one validated ZQ loses its cell, gets an unknown
    cell, or gets a span outside its cell's text."""
    real = document_processor_module.run_validation_workflow

    def review(*args: Any, **kwargs: Any) -> list[DetectedEntity]:
        validated = list(real(*args, **kwargs))
        for index, entity in enumerate(validated):
            if entity.text != ZQ:
                continue
            if how == "no_cell":
                validated[index] = dataclasses.replace(entity, context_label=None)
            elif how == "unknown_cell":
                validated[index] = dataclasses.replace(
                    entity, context_label=f"{SHEET}!Z99"
                )
            else:
                validated[index] = dataclasses.replace(
                    entity, start_pos=len(ZQ) + 10, end_pos=len(ZQ) + 25
                )
            changed.append(how)
            break
        return validated

    return patch.object(document_processor_module, "run_validation_workflow", review)


@pytest.mark.parametrize("how", ["no_cell", "unknown_cell", "span_outside"])
def test_unmatched_cell_stops_run(tmp_path: Path, db: Path, how: str) -> None:
    """T8 (AC4): a validated name that matches no cell stops the run: exit
    code 1, no output file, no mapping row, no success line."""
    source, output = _files(tmp_path, ".xlsx")
    _write_sheet(source, [["Nom", "Note"], [VM, "valide"], [ZQ, "signé"]])
    changed: list[str] = []
    run = _process(
        source,
        output,
        db,
        [(VM, "PERSON"), (ZQ, "PERSON")],
        [ENTER, SPACE, SPACE],
        extra=[_lose_cell(how, changed)],
    )
    _pre(changed == [how], "the wrapper changed one name")
    assert run.exit_code == 1  # expected failure
    assert not output.exists()
    assert _rows(db) == []
    assert SUCCESS not in run.shown


# --- Guards: what must not change (AC11) and 10.13's stored names ----------


@pytest.mark.parametrize("ext", SHEETS)
def test_spreadsheet_accept_only_unchanged(tmp_path: Path, db: Path, ext: str) -> None:
    """G1 (AC11): a spreadsheet whose names are only accepted is written as
    today."""
    source, output = _files(tmp_path, ext)
    _write_sheet(
        source,
        [["Nom", "Note"], [VM, "valide"], ["ok", "contrat avec Zorbal Quentrix"]],
    )
    run = _process(
        source,
        output,
        db,
        [(VM, "PERSON"), (ZQ, "PERSON")],
        [ENTER, SPACE, SPACE],
    )
    assert run.exit_code == 0
    assert SUCCESS in run.shown
    vm = _pseudonym(db, VM, "PERSON")
    pseudonym = _pseudonym(db, ZQ, "PERSON")
    assert _cells(output) == [
        ["Nom", "Note"],
        [vm, "valide"],
        ["ok", f"contrat avec {pseudonym}"],
    ]


def test_text_add_one_place_unchanged(tmp_path: Path, db: Path) -> None:
    """G2 (AC11; the PO's control): [A] of a name at one place of a text."""
    source, output = _files(tmp_path, ".txt")
    text = "Vorlan Mestrix valide. Le contrat avec Zorbal Quentrix."
    source.write_text(text, encoding="utf-8")
    run = _process(
        source, output, db, [(VM, "PERSON")], [ENTER, "a", SPACE], [ZQ, "PERSON"]
    )
    assert run.exit_code == 0
    vm = _pseudonym(db, VM, "PERSON")
    pseudonym = _pseudonym(db, ZQ, "PERSON")
    assert _text(output) == text.replace(VM, vm).replace(ZQ, pseudonym)


def test_text_correction_unchanged(tmp_path: Path, db: Path) -> None:
    """G3 (AC11): [E] in a text is written as today."""
    source, output = _files(tmp_path, ".txt")
    source.write_text(
        "Zorbal Quentrix signé le contrat avec Zorbaville.", encoding="utf-8"
    )
    run = _process(
        source,
        output,
        db,
        [(ZQS, "PERSON"), (ZV, "ORG")],
        [ENTER, "e", SPACE],
        [ZQ],
    )
    assert run.exit_code == 0
    pseudonym = _pseudonym(db, ZQ, "PERSON")
    zv = _pseudonym(db, ZV, "ORG")
    assert _text(output) == f"{pseudonym} signé le contrat avec {zv}."


def test_text_type_change_unchanged(tmp_path: Path, db: Path) -> None:
    """G4 (AC11): [T] in a text is written as today."""
    source, output = _files(tmp_path, ".txt")
    source.write_text("Vorlan Mestrix valide. Zorbal Quentrix signé.", encoding="utf-8")
    run = _process(
        source,
        output,
        db,
        [(VM, "PERSON"), (ZQ, "ORG")],
        [ENTER, SPACE, "t"],
        ["PERSON"],
    )
    assert run.exit_code == 0
    vm = _pseudonym(db, VM, "PERSON")
    pseudonym = _pseudonym(db, ZQ, "PERSON")
    assert _text(output) == f"{vm} valide. {pseudonym} signé."


def test_known_name_kept_through_both_restarts(tmp_path: Path, db: Path) -> None:
    """G5 (QA 10.13's TEST-002): a stored name still covers its place after
    the review restarts, once through each restart (Ctrl+C answered No, and
    the final question answered No)."""
    try:
        from gdpr_pseudonymizer.validation.ui import STILL_TO_REVIEW_LINE
    except ImportError:
        pytest.fail("PRE: 10.13 is merged (its line exists)")

    first, first_out = _files(tmp_path, ".txt", "1")
    first.write_text("Zorbaville valide.", encoding="utf-8")
    run = _process(first, first_out, db, [(ZV, "ORG")], [ENTER, SPACE])
    _pre(run.exit_code == 0, "first run: exit code 0")

    source, output = _files(tmp_path, ".txt", "2")
    source.write_text("Zorbaville a signé avec Quentrix et Zorbalix.", encoding="utf-8")
    run = _process(
        source,
        output,
        db,
        [
            (ZV, "ORG"),
            (ZV, "LOCATION"),
            ("Quentrix", "LOCATION"),
            ("Zorbalix", "LOCATION"),
        ],
        [ENTER, CTRL_C, ENTER, "n", "n", SPACE, ENTER, "n", "n", SPACE],
        answers=[False, False, True],
    )
    _pre(run.exit_code == 0, "second run: exit code 0")
    _pre(
        run.asked == [QUIT_QUESTION, READY_QUESTION, READY_QUESTION],
        "the three questions asked in that order",
    )
    _pre(AUTO_ACCEPTED in run.shown, "the stored name is accepted without review")
    one = STILL_TO_REVIEW_LINE.replace("%1", "1")
    two = STILL_TO_REVIEW_LINE.replace("%1", "2")
    assert run.shown.count(one) == 2
    assert two not in run.shown
    zv = _pseudonym(db, ZV, "ORG")
    zorbalix = _pseudonym(db, "Zorbalix", "LOCATION")
    assert _text(output) == f"{zv} a signé avec Quentrix et {zorbalix}."


# --- Guards from QA's review: the cells through both restarts, one span in two
# cells (TEST-001, TEST-002) ----------------------------------------------------


def test_spreadsheet_correction_after_ctrl_c_restart(tmp_path: Path, db: Path) -> None:
    """G6 (QA TEST-001; mutant M4a): after the Ctrl+C restart of the review
    (« Quit validation? » answered No), a correction is still located in its
    own cell."""
    source, output = _files(tmp_path, ".xlsx")
    _write_sheet(
        source, [[ZQ, "Note"], ["ok", "chez Zorbal Quentrix signé le contrat"]]
    )
    run = _process(
        source,
        output,
        db,
        [(ZQ, "ORG"), (ZQS, "PERSON")],
        [ENTER, CTRL_C, ENTER, "e", SPACE],
        [ZQ],
        answers=[False],
    )
    _pre(run.asked[:1] == [QUIT_QUESTION], "the review restarted after Ctrl+C")
    assert run.exit_code == 0
    cells = _cells(output)
    pseudonym = _pseudonym(db, ZQ, "PERSON")
    assert _cell(cells, "B2") == f"chez {pseudonym} signé le contrat"
    assert _cell(cells, "A1") == _pseudonym(db, ZQ, "ORG")


def test_spreadsheet_add_after_ready_restart(tmp_path: Path, db: Path) -> None:
    """G7 (QA TEST-001; mutant M4b): after the restart of the review (« Ready
    to process … » answered No), a name added with [A] is replaced in every
    cell where it stands, and the run does not stop."""
    source, output = _files(tmp_path, ".xlsx")
    _write_sheet(source, [["Nom", "Note"], [VM, "avec Zorbal Quentrix"], [ZQ, "signé"]])
    run = _process(
        source,
        output,
        db,
        [(VM, "PERSON")],
        [ENTER, SPACE, ENTER, "a", SPACE],
        [ZQ, "PERSON"],
        answers=[False],
    )
    _pre(
        run.asked[:2] == [READY_QUESTION, READY_QUESTION],
        "the review restarted after No",
    )
    assert run.exit_code == 0
    assert SUCCESS in run.shown
    cells = _cells(output)
    pseudonym = _pseudonym(db, ZQ, "PERSON")
    assert ZQ not in _flat(cells)
    assert _cell(cells, "B2") == f"avec {pseudonym}"
    assert _cell(cells, "A3") == pseudonym
    assert _cell(cells, "A2") == _pseudonym(db, VM, "PERSON")


def test_add_same_span_in_another_cell(tmp_path: Path, db: Path) -> None:
    """G8 (QA TEST-002; mutant M5): a place in another cell at the same
    position as a found place is a place of its own: [A] adds it, and both
    cells are replaced."""
    source, output = _files(tmp_path, ".xlsx")
    _write_sheet(source, [["Nom", "Note"], [ZQ, "valide"], [ZQ, "signé"]])
    run = _process(
        source,
        output,
        db,
        [(ZQ, "PERSON")],
        [ENTER, "a", SPACE],
        [ZQ, "PERSON"],
        first_only=True,
    )
    assert run.exit_code == 0
    cells = _cells(output)
    pseudonym = _pseudonym(db, ZQ, "PERSON")
    assert ZQ not in _flat(cells)
    assert _cell(cells, "A2") == pseudonym
    assert _cell(cells, "A3") == pseudonym


# --- v1.4: a correction that partly overlaps another written name (the read-only
# Codex review's P0; Lionel's Q4 applied to corrections) ------------------------

OVERLAP = "Vorlan Zorbal Quentrix a signé."
MIRROR = "Zorbal Quentrix Vorlan a signé."
QV = "Quentrix Vorlan"
NAME_PARTS = ("Vorlan", "Zorbal", "Quentrix")


def _batch(
    folder: Path,
    out_folder: Path,
    db: Path,
    names: Names,
    keys: Sequence[str],
    typed: Sequence[str] = (),
) -> Run:
    """One real ``gdpr-pseudo batch --workers 1`` run, with ``_process``'s
    detection, keys, typed texts and yes/no answers."""
    detector = FakeDetector(names)
    confirm = Answers()
    with ExitStack() as stack:
        stack.enter_context(
            patch.object(DocumentProcessor, "_get_detector", new=lambda self: detector)
        )
        stack.enter_context(
            patch.object(
                DocumentProcessor,
                "_detect_and_filter_entities",
                new=lambda self, text, entity_type_filter=None: (
                    detector.detect_entities(text)
                ),
            )
        )
        stack.enter_context(patch.object(readchar, "readkey", new=Keys(keys)))
        stack.enter_context(patch.object(Prompt, "ask", new=Typed(typed)))
        stack.enter_context(patch.object(Confirm, "ask", new=confirm))
        result = CliRunner().invoke(
            app,
            [
                "batch",
                str(folder),
                "--db",
                str(db),
                "--output",
                str(out_folder),
                "--workers",
                "1",
            ],
            env=ENV,
            catch_exceptions=False,
        )
    return Run(result.exit_code, _ANSI.sub("", result.output), confirm.asked)


def _write_overlap(path: Path, text: str) -> None:
    """The text as a text file, or in cell B2 of a sheet."""
    if path.suffix == ".txt":
        path.write_text(text, encoding="utf-8")
    else:
        _write_sheet(path, [["Nom", "Note"], ["ok", text]])


def _overlap_written(path: Path) -> str:
    """The text written, or cell B2 of a sheet whose other cells are unchanged."""
    if path.suffix == ".txt":
        return _text(path)
    cells = _cells(path)
    _pre(
        len(cells) == 2
        and [_cell(cells, "A1"), _cell(cells, "B1"), _cell(cells, "A2")]
        == ["Nom", "Note", "ok"],
        "the sheet's other cells unchanged",
    )
    return _cell(cells, "B2")


def _names(db: Path) -> list[tuple[str, str]]:
    return sorted((row[0], row[1]) for row in _rows(db))


@pytest.mark.parametrize("ext", [".txt", ".xlsx", ".csv"])
def test_correction_overlapping_accepted_name(
    tmp_path: Path, db: Path, ext: str
) -> None:
    """T16 (AC1, AC14 (d) for [E]; the Codex P0): a correction that reaches
    into an accepted name is widened over both and takes its own pseudonym."""
    source, output = _files(tmp_path, ext)
    _write_overlap(source, OVERLAP)
    run = _process(
        source,
        output,
        db,
        [(VZ, "PERSON"), ("Quentrix", "PERSON")],
        [ENTER, SPACE, "e"],
        [ZQ],
    )
    _pre(run.exit_code == 0, "exit code 0")
    _pseudonym(db, VZ, "PERSON")
    pseudonym = _pseudonym(db, ZQ, "PERSON")
    written = _overlap_written(output)
    assert [part for part in NAME_PARTS if part in written] == []  # expected failure
    assert written == f"{pseudonym} a signé."
    assert _names(db) == sorted([(VZ, "PERSON"), (ZQ, "PERSON")])


@pytest.mark.parametrize("ext", [".txt", ".xlsx", ".csv"])
def test_correction_then_accepted_overlap(tmp_path: Path, db: Path, ext: str) -> None:
    """T17 (the mirror order): a correction that reaches into a name accepted
    after it is widened over both and takes its own pseudonym."""
    source, output = _files(tmp_path, ext)
    _write_overlap(source, MIRROR)
    run = _process(
        source,
        output,
        db,
        [("Zorbal", "PERSON"), (QV, "PERSON")],
        [ENTER, "e", SPACE],
        [ZQ],
    )
    _pre(run.exit_code == 0, "exit code 0")
    _pseudonym(db, QV, "PERSON")
    pseudonym = _pseudonym(db, ZQ, "PERSON")
    written = _overlap_written(output)
    assert [part for part in NAME_PARTS if part in written] == []  # expected failure
    assert written == f"{pseudonym} a signé."
    assert _names(db) == sorted([(QV, "PERSON"), (ZQ, "PERSON")])


@pytest.mark.parametrize("ext", [".txt", ".xlsx"])
def test_correction_overlapping_stored_name(tmp_path: Path, db: Path, ext: str) -> None:
    """T18 (AC14 (d)'s stored names, for [E]): a correction that reaches into
    a name stored in the mapping database, accepted without review, is widened
    over both; the stored row is unchanged."""
    first, first_out = _files(tmp_path, ext, "1")
    _write_overlap(first, "Vorlan Zorbal valide.")
    run = _process(first, first_out, db, [(VZ, "PERSON")], [ENTER, SPACE])
    _pre(run.exit_code == 0, "first run: exit code 0")
    stored = _pseudonym(db, VZ, "PERSON")

    source, output = _files(tmp_path, ext, "2")
    _write_overlap(source, OVERLAP)
    run = _process(
        source,
        output,
        db,
        [(VZ, "PERSON"), ("Quentrix", "PERSON")],
        [ENTER, "e"],
        [ZQ],
    )
    _pre(run.exit_code == 0, "second run: exit code 0")
    _pre(AUTO_ACCEPTED in run.shown, "the stored name is accepted without review")
    _pre(_pseudonym(db, VZ, "PERSON") == stored, "the stored row is unchanged")
    pseudonym = _pseudonym(db, ZQ, "PERSON")
    written = _overlap_written(output)
    assert [part for part in NAME_PARTS if part in written] == []  # expected failure
    assert written == f"{pseudonym} a signé."
    assert _names(db) == sorted([(VZ, "PERSON"), (ZQ, "PERSON")])


@pytest.mark.parametrize("ext", [".txt", ".xlsx", ".csv"])
def test_batch_correction_overlap(tmp_path: Path, db: Path, ext: str) -> None:
    """T19: the same correction through ``batch --workers 1``."""
    folder, out_folder = tmp_path / "docs", tmp_path / "out"
    folder.mkdir()
    _write_overlap(folder / f"in{ext}", OVERLAP)
    run = _batch(
        folder,
        out_folder,
        db,
        [(VZ, "PERSON"), ("Quentrix", "PERSON")],
        [ENTER, SPACE, "e"],
        [ZQ],
    )
    _pre(run.exit_code == 0, "exit code 0")
    output = out_folder / f"in_pseudonymized{ext}"
    _pseudonym(db, VZ, "PERSON")
    pseudonym = _pseudonym(db, ZQ, "PERSON")
    written = _overlap_written(output)
    assert [part for part in NAME_PARTS if part in written] == []  # expected failure
    assert written == f"{pseudonym} a signé."
    assert _names(db) == sorted([(VZ, "PERSON"), (ZQ, "PERSON")])


# --- v1.5: every place [E] reaches is typed (the PO's N4) ----------------------

SECOND_PLACE = "Zorbal Quentrix valide. Vorlan Zorbal Quentrix a signé."


@pytest.mark.parametrize("ext", [".txt", ".xlsx", ".csv"])
def test_correction_overlap_at_second_place(tmp_path: Path, db: Path, ext: str) -> None:
    """T20 (AC1, AC14 (d) for [E] at every place of its group; the PO's N4):
    [E] corrects both places of its group; at the second, the correction
    reaches into an accepted name and is widened over both there too."""
    source, output = _files(tmp_path, ext)
    _write_overlap(source, SECOND_PLACE)
    _pre(
        [SECOND_PLACE.find("Quentrix"), SECOND_PLACE.rfind("Quentrix")] == [7, 38]
        and SECOND_PLACE.find(VZ) == 24,
        "the places of the Test Plan",
    )
    run = _process(
        source,
        output,
        db,
        [("Quentrix", "PERSON"), (VZ, "PERSON")],
        [ENTER, "e", SPACE],
        [ZQ],
    )
    _pre(run.exit_code == 0, "exit code 0")
    _pseudonym(db, VZ, "PERSON")
    pseudonym = _pseudonym(db, ZQ, "PERSON")
    written = _overlap_written(output)
    assert [part for part in NAME_PARTS if part in written] == []  # expected failure
    assert written == f"{pseudonym} valide. {pseudonym} a signé."
    assert _names(db) == sorted([(VZ, "PERSON"), (ZQ, "PERSON")])


@pytest.mark.parametrize("ext", [".xlsx", ".csv"])
def test_correction_overlap_in_another_cell(tmp_path: Path, db: Path, ext: str) -> None:
    """T21 (T20 across two cells): the correction's second place, in another
    cell, reaches into an accepted name and is widened over both there."""
    source, output = _files(tmp_path, ext)
    _write_sheet(
        source,
        [
            ["Nom", "Note"],
            ["ok", "Zorbal Quentrix valide."],
            ["ok", "Vorlan Zorbal Quentrix a signé."],
        ],
    )
    run = _process(
        source,
        output,
        db,
        [("Quentrix", "PERSON"), (VZ, "PERSON")],
        [ENTER, SPACE, "e"],
        [ZQ],
    )
    _pre(run.exit_code == 0, "exit code 0")
    _pseudonym(db, VZ, "PERSON")
    pseudonym = _pseudonym(db, ZQ, "PERSON")
    cells = _cells(output)
    _pre(
        len(cells) == 3
        and [_cell(cells, ref) for ref in ("A1", "B1", "A2", "A3")]
        == ["Nom", "Note", "ok", "ok"],
        "the sheet's other cells unchanged",
    )
    b2, b3 = _cell(cells, "B2"), _cell(cells, "B3")
    both = f"{b2}\n{b3}"
    assert [part for part in NAME_PARTS if part in both] == []  # expected failure
    assert (b2, b3) == (f"{pseudonym} valide.", f"{pseudonym} a signé.")
    assert _names(db) == sorted([(VZ, "PERSON"), (ZQ, "PERSON")])


# --- Guards from QA's re-check: the widening's fixed point and the tie rule
# (TEST-003, TEST-004) ----------------------------------------------------------

QB = "Quentrix Brimel"


@pytest.mark.parametrize("ext", [".txt", ".xlsx"])
def test_widening_reaches_fixed_point(tmp_path: Path, db: Path, ext: str) -> None:
    """G9 (QA TEST-003; mutant MC, one pass): a correction widened over an
    accepted name that itself partly overlaps a stored name is widened again,
    until it covers both."""
    first, first_out = _files(tmp_path, ext, "1")
    _write_overlap(first, "Vorlan Mestrix valide.")
    run = _process(first, first_out, db, [(VM, "PERSON")], [ENTER, SPACE])
    _pre(run.exit_code == 0, "first run: exit code 0")
    stored = _pseudonym(db, VM, "PERSON")

    source, output = _files(tmp_path, ext, "2")
    _write_overlap(source, "ok Zorbal Quentrix Vorlan Mestrix a signé.")
    run = _process(
        source,
        output,
        db,
        [("Zorbal", "PERSON"), (QV, "PERSON"), (VM, "PERSON")],
        [ENTER, "e", SPACE],
        [ZQ],
    )
    _pre(AUTO_ACCEPTED in run.shown, "the stored name is accepted without review")
    assert run.exit_code == 0
    written = _overlap_written(output)
    pseudonym = _pseudonym(db, ZQ, "PERSON")
    parts = (*NAME_PARTS, "Mestrix")
    assert [part for part in parts if part in written] == []
    assert written == f"ok {pseudonym} a signé."
    assert _pseudonym(db, VM, "PERSON") == stored


def test_overlapping_typed_places_first_in_list_wins(tmp_path: Path, db: Path) -> None:
    """G10 (QA TEST-004; mutants MD, the later place wins, and MD2, the typed
    places taken last to first): two corrections that partly overlap each
    other and a stored name both end with the same span; the first in the
    list sent (the first group reviewed) gives the pseudonym."""
    first, first_out = _files(tmp_path, ".txt", "1")
    first.write_text("Quentrix Brimel valide.", encoding="utf-8")
    run = _process(first, first_out, db, [(QB, "PERSON")], [ENTER, SPACE])
    _pre(run.exit_code == 0, "first run: exit code 0")
    stored = _pseudonym(db, QB, "PERSON")

    source, output = _files(tmp_path, ".txt", "2")
    source.write_text("Vorlan Zorbal Quentrix Brimel a signé.", encoding="utf-8")
    run = _process(
        source,
        output,
        db,
        [("Vorlan", "PERSON"), ("Zorbal", "PERSON"), (QB, "PERSON")],
        [ENTER, "e", "e"],
        [VZ, ZQ],
    )
    _pre(AUTO_ACCEPTED in run.shown, "the stored name is accepted without review")
    assert run.exit_code == 0
    written = _text(output)
    parts = (*NAME_PARTS, "Brimel")
    assert [part for part in parts if part in written] == []
    assert written == f"{_pseudonym(db, VZ, 'PERSON')} a signé."
    _pseudonym(db, ZQ, "PERSON")
    assert _pseudonym(db, QB, "PERSON") == stored
