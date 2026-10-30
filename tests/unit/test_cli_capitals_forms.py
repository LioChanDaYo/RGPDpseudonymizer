"""Story 10.19: a name added ([A]) or corrected ([E]) in the command-line
review also reaches its forms in capitals, and ``delete-mapping`` by name also
erases the saved pseudonyms of the name in other capitals (T1-T7, T9-T13,
T17, T18, T19 CLI, T20, T21 CLI, T22-T25, T28, T30 CLI, T31, T33-T51 CLI),
and what must not change (G1, G2, G7, G8 CLI, G11-G30, T23's case 2, T44's,
T45's, T47's, T48's and T50's accented cases, T49's case without accents).

Written by the BMAD dev agent (Claude), before the fix. Each test drives the
real ``gdpr-pseudo process`` or ``delete-mapping`` command through
``CliRunner``, on a database made by the real ``init`` command, as Story
10.15's harness (``test_cli_corrections_hand_added.py``). Detection is
replaced by a fake that finds every occurrence of each (name, type) given,
with its exact capitals; the keys, the typed texts and the yes/no answers come
from lists. The assertions read the written file (text or cells), the stored
rows, the ERASURE entries and the lines shown.

Preconditions call ``pytest.fail("PRE: ...")``, never ``assert``; the first
``assert`` after them is the expected failure on the unfixed code
(``# expected failure``). Invented names and main-corpus words only; a
decomposed text is made at run time with ``unicodedata.normalize("NFD", ...)``
and the one combining mark T20 needs with ``unicodedata.lookup``: no combining
mark is written in this file. No spaCy model, no Qt.
"""

from __future__ import annotations

import dataclasses
import json
import re
import sqlite3
import unicodedata
from collections.abc import Sequence
from contextlib import ExitStack, closing
from pathlib import Path
from typing import Any
from unittest.mock import patch

import openpyxl
import pytest
import readchar
from rich.prompt import Confirm, Prompt
from typer.testing import CliRunner

from gdpr_pseudonymizer.cli.main import app
from gdpr_pseudonymizer.core.document_processor import DocumentProcessor
from gdpr_pseudonymizer.core.mapping_key import _key_steps
from gdpr_pseudonymizer.data.database import open_database
from gdpr_pseudonymizer.data.models import Entity
from gdpr_pseudonymizer.data.repositories.mapping_repository import (
    SQLiteMappingRepository,
)
from gdpr_pseudonymizer.nlp.entity_detector import DetectedEntity
from gdpr_pseudonymizer.validation.models import fold
from gdpr_pseudonymizer.validation.workflow import ADD_MISSED_NAME_QUESTION

PASSPHRASE = "zorbalia_pass_A_123!"
ENV = {"GDPR_PSEUDO_PASSPHRASE": PASSPHRASE}
SHEET = "Zorbalia"
PERSON = "PERSON"
LOCATION = "LOCATION"
ORG = "ORG"

ENTER = "\r"
SPACE = " "
MAX_KEYS = 60

VM = "Vorlan Mestrix"
VZ = "Vorlan Zorbal"
ZQ = "Zorbal Quentrix"
VQ = "Vorlène Quentrix"
QZ = "Quovôme Zorbal"

T1_TEXT = (
    "Vorlan Mestrix valide. Le contrat avec Zorbal. Signé : ZORBAL. ZORBAL a signé."
)
T1_ROWS = [
    ["Nom", "Note"],
    [VM, "valide"],
    ["ok", "contrat avec Zorbal"],
    ["ok", "Signé : ZORBAL"],
    ["ZORBAL", "a signé"],
]
T2_TEXT = "Zorbal valide. Le contrat avec Zorbal. Signé : ZORBAL. ZORBAL a signé."
T3_TEXT = "Zorbal valide. Signé : ZORBAL. Le contrat avec ZORBAL."
T4_TEXT = "Vorlan Mestrix valide. Le petit contrat avec Petit. PETIT a signé."
T5_TEXT = "Vorlan Mestrix valide. Le contrat avec Vorlan Zorbal et Vorlan ZORBAL."
T6_TEXT = "Vorlan Mestrix valide. Le contrat avec Vorlan de Zorbal."
T7_TEXT = "Zorbal Quentrix a signé. Signé : ZORBAL QUENTRIX."
T7_ROWS = [
    ["Nom", "Note"],
    ["ok", "Zorbal Quentrix a signé."],
    ["ok", "Signé : ZORBAL QUENTRIX."],
]
T10_TEXT = (
    "Vorlan Mestrix valide. Le contrat avec Vorlène. VORLÈNE a signé. "
    "VORLENE a signé. Vorlene a signé."
)
T17_TEXT = "Vorlène Quentrix a signé. Signé : VORLENE QUENTRIX."
G1_TEXT = "Vorlan Mestrix valide. Le contrat avec Zorbal et SAINT-ZORBAL."
G7_TEXT = "Vorlène Quentrix a signé. Signé : Vorlene Quentrix."
G8_TEXT = (
    "Vorlène Quentrix a signé. Signé : VORLENE QUENTRIX. "
    "Le contrat avec Vorlene Quentrix."
)
T21_TEXT = f"Vorlan Mestrix valide. {G8_TEXT}"
T22_TEXT = (
    "MEZOR a signé. Vorlan Mézor valide. Signé : Mézor. "
    "Le contrat avec Zorbal Mèzor. Signé : Mèzor."
)
T22_SWAPPED = (
    "Vorlan Mézor valide. Signé : Mézor. MEZOR a signé. "
    "Le contrat avec Zorbal Mèzor. Signé : Mèzor."
)
T22_SET = "MEZOR a signé. Mézor valide. Le contrat avec Mèzor."
T22_NAMES = [
    ("MEZOR", PERSON),
    ("Vorlan Mézor", PERSON),
    ("Mézor", PERSON),
    ("Zorbal Mèzor", PERSON),
    ("Mèzor", PERSON),
]
T24_TEXT = (
    "Brimel Testar valide. Zélvine Testar a signé. Signé : ZELVINE TESTAR. "
    "Le contrat avec Zelvine Testar."
)
T25_TEXT = "Signé : ZELVINE TESTAR. Zélvine Testar a signé."
G11_TEXT = "Signé : VORLENE QUENTRIX. Vorlène Quentrix a signé."
G13_TEXT = "Le contrat est validé. Signé : ZORBAL QUENTRIX."
T30_TEXT = "Signé : ZÉLVINE TESTAR. Zélvine Testar a signé."
T31_TEXT = "Zélvine Mézor a signé. Zélvine Mèzor a signé."
T35_TEXT = "Signé : ZÉLVINE TESTAR. Zélvine\nTestar a signé."
T36_TEXT = "Vorlan Mestrix valide. Vorlène a signé. VORLENE a signé."
T36B_TEXT = "Vorlan Mestrix valide. Signé : Bélora. Signé : BÊLORA."
# T40: the punctuation written right after « Testar »
PUNCTUATION = {
    "ellipsis": "\u2026",
    "em_dash": "\u2014",
    "en_dash": "\u2013",
    "slash": "/",
    "asterisk": "*",
    "footnote_digit": "1",
}
# T42: (the text, the two places found, the correction typed): an initial,
# a comma, an initial typed without its full stops
T42_CASES = {
    "initial": (
        "Signé : ZÉLVINE J. TESTAR. Zélvine J. Testar a signé.",
        ("ZÉLVINE", "Zélvine"),
        "ZELVINE J. TESTAR",
    ),
    "double_initial": (
        "Signé : ZÉLVINE J.-P. TESTAR. Zélvine J.-P. Testar a signé.",
        ("ZÉLVINE", "Zélvine"),
        "ZELVINE J.-P. TESTAR",
    ),
    "compound_initial": (
        "Signé : ZÉLVINE J.-M. TESTAR. Zélvine J.-M. Testar a signé.",
        ("ZÉLVINE", "Zélvine"),
        "ZELVINE J.-M. TESTAR",
    ),
    "comma": (
        "Signé : TESTAR, ZÉLVINE. Testar, Zélvine a signé.",
        ("TESTAR", "Testar"),
        "TESTAR, ZELVINE",
    ),
    "typed_without_the_full_stop": (
        "Signé : ZÉLVINE J. TESTAR. Zélvine J. Testar a signé.",
        ("ZÉLVINE", "Zélvine"),
        "ZELVINE J TESTAR",
    ),
    "typed_without_the_full_stops": (
        "Signé : ZÉLVINE J.-M. TESTAR. Zélvine J.-M. Testar a signé.",
        ("ZÉLVINE", "Zélvine"),
        "ZELVINE J-M TESTAR",
    ),
}
# T51: (the name in capitals, the name typed) with a bracket, a square
# bracket or a slash before a word the old step cut
T51_CASES = {
    "brackets_in_the_middle": ("MÉLORA (MÉLIME) TESTAR", "Melora (Melime) Testar"),
    "brackets_at_the_end": ("QUOVANE (MÉLORA)", "Quovane (Mélora)"),
    "square_brackets": ("QUOVANE [MÉLORA]", "Quovane [Melora]"),
    "slash": ("QUOVANE/MÉLORA", "Quovane/Melora"),
}
# T37: the second name between quotation marks (opening, closing)
QUOTES = {
    "straight": ("'", "'"),
    "curly": ("\u2018", "\u2019"),
    "curly_double": ("\u201c", "\u201d"),
}

# Story 10.15's approved line (Lionel's Q2 there), for « Zorbal » as PERSON.
ALREADY_LINE = (
    "'Zorbal' is already in the review as PERSON at every place: nothing added"
)
# Lionel's 10.19 Q2 line (2026-10-19, "Approve (Rec.)"), N = 2, byte for byte.
SEVERAL_LINE = (
    "You are about to delete the following 2 saved pseudonyms "
    "(all of this name, whatever its type or its capitals):"
)
ONE_ROW_LINE = "You are about to delete the following entity mapping:"
# Lionel's 10.19 Q5b lines (2026-10-19, "Say mapping (Rec.)"), byte for byte.
ID_LINE_ONE = (
    "1 other mapping with the same name remains (type: ORG). Delete one with "
    "'gdpr-pseudo delete-mapping --id <uuid>', or every mapping of the name with "
    "'gdpr-pseudo delete-mapping \"<name>\"'."
)
ID_LINE_TWO = (
    "2 other mappings with the same name remain (types: LOCATION, ORG). Delete "
    "one with 'gdpr-pseudo delete-mapping --id <uuid>', or every mapping of the "
    "name with 'gdpr-pseudo delete-mapping \"<name>\"'."
)
DELETED = "deleted successfully"
NAME_LINE = "Name:       "

_ANSI = re.compile(r"\x1b\[[0-9;?]*[A-Za-z]")

Names = Sequence[tuple[str, str]]
Cells = list[list[str]]


def nfd(text: str) -> str:
    return unicodedata.normalize("NFD", text)


def _pre(condition: bool, what: str) -> None:
    """A precondition: a red failure on every leg, never an AssertionError."""
    if not condition:
        pytest.fail(f"PRE: {what}")


def _collapsed(text: str) -> str:
    """``text`` with every whitespace run made one space (Rich wraps lines)."""
    return " ".join(text.split())


@pytest.fixture(autouse=True)
def _settings(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    """No user or project settings file is read or written."""
    monkeypatch.setenv("HOME", str(tmp_path))
    monkeypatch.setenv("USERPROFILE", str(tmp_path))
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        "gdpr_pseudonymizer.gui.config._config_path",
        lambda: tmp_path / ".gdpr-pseudo.yaml",
    )


@pytest.fixture
def db(tmp_path: Path) -> Path:
    """A new database, made by the real ``init`` command."""
    path = tmp_path / "mappings.db"
    result = CliRunner().invoke(
        app, ["init", "--db", str(path)], env=ENV, catch_exceptions=False
    )
    _pre(result.exit_code == 0 and path.is_file(), "init made the database")
    return path


# ---------------------------------------------------------------------------
# Harness
# ---------------------------------------------------------------------------


class FakeDetector:
    """Finds every occurrence of each (name, type) given, with its exact
    capitals (``str.find``), and drops an occurrence that lies inside a longer
    one it found. With ``first_only``, each name is found once, at its first
    occurrence over all calls. Records every text it is given."""

    def __init__(self, names: Names, first_only: bool = False) -> None:
        self.names = list(names)
        self.first_only = first_only
        self.seen: set[tuple[str, str]] = set()
        self.texts: list[str] = []

    def detect_entities(self, text: str) -> list[DetectedEntity]:
        self.texts.append(text)
        found: list[DetectedEntity] = []
        for name, entity_type in self.names:
            if self.first_only and (name, entity_type) in self.seen:
                continue
            start = text.find(name)
            while start != -1:
                found.append(
                    DetectedEntity(
                        text=name,
                        entity_type=entity_type,
                        start_pos=start,
                        end_pos=start + len(name),
                        confidence=0.9,
                        source="spacy",
                    )
                )
                if self.first_only:
                    self.seen.add((name, entity_type))
                    break
                start = text.find(name, start + 1)
        return [
            entity
            for entity in found
            if not any(
                other.start_pos <= entity.start_pos
                and entity.end_pos <= other.end_pos
                and other.end_pos - other.start_pos > entity.end_pos - entity.start_pos
                for other in found
            )
        ]


class Keys:
    """``readchar.readkey``: the keys in order, then Space."""

    def __init__(self, keys: Sequence[str]) -> None:
        self.keys = list(keys)
        self.read = 0

    def __call__(self) -> str:
        self.read += 1
        if self.read > MAX_KEYS:
            raise RuntimeError("more than 60 keys read")
        return self.keys.pop(0) if self.keys else SPACE


class Typed:
    """``Prompt.ask``: the typed texts in order."""

    def __init__(self, texts: Sequence[str]) -> None:
        self.texts = list(texts)

    def __call__(self, prompt: Any, *args: Any, **kwargs: Any) -> str:
        if not self.texts:
            raise RuntimeError("no typed text left")
        return self.texts.pop(0)


class Answers:
    """``Confirm.ask``: no to Story 10.18's question, yes otherwise."""

    def __call__(self, prompt: Any, *args: Any, **kwargs: Any) -> bool:
        return str(prompt) != ADD_MISSED_NAME_QUESTION


@dataclasses.dataclass
class Run:
    exit_code: int
    shown: str

    @property
    def flat(self) -> str:
        return _collapsed(self.shown)

    @property
    def screen(self) -> str:
        """What is shown before the success lines."""
        return self.shown.split(DELETED)[0]

    @property
    def after(self) -> str:
        """What is shown from the success lines on, whitespace collapsed."""
        parts = self.shown.split(DELETED, 1)
        return _collapsed(parts[1]) if len(parts) == 2 else ""


def _invoke(
    args: Sequence[str],
    detector: FakeDetector | None = None,
    keys: Sequence[str] = (),
    typed: Sequence[str] = (),
) -> Run:
    """One real ``gdpr-pseudo`` command run."""
    with ExitStack() as stack:
        if detector is not None:
            fake = detector
            stack.enter_context(
                patch.object(DocumentProcessor, "_get_detector", new=lambda self: fake)
            )
            stack.enter_context(
                patch.object(
                    DocumentProcessor,
                    "_detect_and_filter_entities",
                    new=lambda self, text, entity_type_filter=None: (
                        fake.detect_entities(text)
                    ),
                )
            )
        stack.enter_context(patch.object(readchar, "readkey", new=Keys(keys)))
        stack.enter_context(patch.object(Prompt, "ask", new=Typed(typed)))
        stack.enter_context(patch.object(Confirm, "ask", new=Answers()))
        result = CliRunner().invoke(app, list(args), env=ENV, catch_exceptions=False)
    return Run(result.exit_code, _ANSI.sub("", result.output))


def _process(
    source: Path,
    output: Path,
    db: Path,
    names: Names,
    keys: Sequence[str],
    typed: Sequence[str] = (),
    *,
    first_only: bool = False,
    detector: FakeDetector | None = None,
) -> Run:
    """One real ``gdpr-pseudo process`` run."""
    return _invoke(
        ["process", str(source), "--db", str(db), "-o", str(output)],
        detector if detector is not None else FakeDetector(names, first_only),
        keys,
        typed,
    )


def _delete(db: Path, *args: str, typed: Sequence[str] = ("yes",)) -> Run:
    """One real ``gdpr-pseudo delete-mapping`` run."""
    return _invoke(["delete-mapping", *args, "--db", str(db)], typed=typed)


def _write_sheet(path: Path, rows: Cells) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = SHEET
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


def _write(source: Path, text: str, rows: Cells) -> None:
    """The document: ``text`` in a text file, ``rows`` in a workbook."""
    if source.suffix == ".xlsx":
        _write_sheet(source, rows)
    else:
        source.write_text(text, encoding="utf-8")


def _cells(path: Path) -> Cells:
    """The written cells, as rows of strings."""
    _pre(path.is_file(), "the output file is written")
    workbook = openpyxl.load_workbook(path)
    try:
        return [
            ["" if value is None else str(value) for value in row]
            for row in workbook.worksheets[0].iter_rows(values_only=True)
        ]
    finally:
        workbook.close()


def _flat(rows: Cells) -> str:
    return "\n".join("\t".join(row) for row in rows)


def _written(path: Path) -> str:
    """The written text; a workbook's cells, joined."""
    if path.suffix == ".xlsx":
        return _flat(_cells(path))
    _pre(path.is_file(), "the output file is written")
    return path.read_text(encoding="utf-8")


def _files(tmp_path: Path, ext: str) -> tuple[Path, Path]:
    return tmp_path / f"in{ext}", tmp_path / f"out{ext}"


def _entities(db: Path) -> list[Entity]:
    """Every stored row, decrypted."""
    with open_database(str(db), PASSPHRASE) as session:
        return SQLiteMappingRepository(session).find_all()


def _names(db: Path) -> list[tuple[str, str]]:
    """Every stored (full_name, entity_type), sorted."""
    return sorted((row.full_name, row.entity_type) for row in _entities(db))


def _pseudonym(db: Path, name: str, entity_type: str = PERSON) -> str:
    """The pseudonym of the one row stored for (name, type)."""
    found = [
        row.pseudonym_full
        for row in _entities(db)
        if (row.full_name, row.entity_type) == (name, entity_type)
    ]
    _pre(len(found) == 1, f"one row stored for ({name}, {entity_type})")
    return str(found[0])


def _ids(db: Path) -> dict[tuple[str, str], str]:
    """The id of each stored (full_name, entity_type)."""
    return {(row.full_name, row.entity_type): row.id for row in _entities(db)}


def _store(db: Path, tmp_path: Path, rows: Names) -> None:
    """Each (name, type) stored by a real finalization of « <name> valide. »."""
    processor = DocumentProcessor(db_path=str(db), passphrase=PASSPHRASE)
    for index, (name, entity_type) in enumerate(rows):
        result = processor.finalize_document(
            document_text=f"{name} valide.",
            validated_entities=[
                DetectedEntity(
                    text=name,
                    entity_type=entity_type,
                    start_pos=0,
                    end_pos=len(name),
                    confidence=0.9,
                    source="spacy",
                )
            ],
            output_path=str(tmp_path / f"stored{index}.txt"),
        )
        _pre(result.success, f"row {index} stored")


def _save(db: Path, rows: Sequence[tuple[str, str, str]]) -> None:
    """Rows saved as given, (full_name, type, pseudonym), as an earlier
    version saved them."""
    with open_database(str(db), PASSPHRASE) as session:
        repo = SQLiteMappingRepository(session)
        for full_name, entity_type, pseudonym in rows:
            repo.save(
                Entity(
                    entity_type=entity_type,
                    full_name=full_name,
                    pseudonym_full=pseudonym,
                    theme="neutral",
                )
            )


def _erasures(db: Path) -> list[str]:
    """The ``user_modifications`` JSON of every ERASURE entry, as stored."""
    with closing(sqlite3.connect(str(db))) as connection:
        rows = connection.execute(
            "SELECT user_modifications FROM operations "
            "WHERE operation_type = 'ERASURE'"
        ).fetchall()
    return [str(row[0]) for row in rows]


def _erased_ids(db: Path) -> list[str]:
    """The ``deleted_entity_id`` of every ERASURE entry, sorted."""
    return sorted(str(json.loads(raw)["deleted_entity_id"]) for raw in _erasures(db))


# ---------------------------------------------------------------------------
# AC1, AC2: [A] reaches every form of the text typed
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("ext", [".txt", ".xlsx"])
def test_add_reaches_capital_forms(tmp_path: Path, db: Path, ext: str) -> None:
    """T1 (AC1, AC2): [A] « Zorbal » also writes « ZORBAL » with its pseudonym,
    in a text file and in every cell; one row is stored."""
    source, output = _files(tmp_path, ext)
    _write(source, T1_TEXT, T1_ROWS)
    run = _process(
        source, output, db, [(VM, PERSON)], [ENTER, "a", SPACE], ["Zorbal", PERSON]
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    vm, z = _pseudonym(db, VM), _pseudonym(db, "Zorbal")
    assert "ZORBAL" not in written, written  # expected failure
    if ext == ".xlsx":
        assert _cells(output) == [
            ["Nom", "Note"],
            [vm, "valide"],
            ["ok", f"contrat avec {z}"],
            ["ok", f"Signé : {z}"],
            [z, "a signé"],
        ]
    else:
        assert written == (
            f"{vm} valide. Le contrat avec {z}. Signé : {z}. {z} a signé."
        )
    assert _names(db) == sorted([(VM, PERSON), ("Zorbal", PERSON)])


def test_capitals_left_when_every_title_form_found(tmp_path: Path, db: Path) -> None:
    """T2 (AC2): every « Zorbal » found, « ZORBAL » not: the approved line is
    not shown and the capital places are added."""
    source, output = _files(tmp_path, ".txt")
    source.write_text(T2_TEXT, encoding="utf-8")
    run = _process(
        source,
        output,
        db,
        [("Zorbal", PERSON)],
        [ENTER, "a", SPACE],
        ["Zorbal", PERSON],
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    assert ALREADY_LINE not in run.flat, run.shown  # expected failure
    assert "ZORBAL" not in _written(output)
    assert "Added PERSON entity: Zorbal" in run.flat
    assert _names(db) == [("Zorbal", PERSON)]


def test_found_capital_place_keeps_its_decision(tmp_path: Path, db: Path) -> None:
    """T3 (AC2): « ZORBAL » found at its first place and rejected keeps its
    decision; [A] « Zorbal » writes the other « ZORBAL »."""
    source, output = _files(tmp_path, ".txt")
    source.write_text(T3_TEXT, encoding="utf-8")
    run = _process(
        source,
        output,
        db,
        [("ZORBAL", PERSON)],
        [ENTER, "a", "r"],
        ["Zorbal", PERSON],
        first_only=True,
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    assert written.count("ZORBAL") == 1, written  # expected failure
    z = _pseudonym(db, "Zorbal")
    assert written == f"{z} valide. Signé : ZORBAL. Le contrat avec {z}."
    assert _names(db) == [("Zorbal", PERSON)]


def test_small_letter_word_never_reached(tmp_path: Path, db: Path) -> None:
    """T4 (AC1): [A] « Petit » writes « PETIT », never « petit »."""
    source, output = _files(tmp_path, ".txt")
    source.write_text(T4_TEXT, encoding="utf-8")
    run = _process(
        source, output, db, [(VM, PERSON)], [ENTER, "a", SPACE], ["Petit", PERSON]
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    assert "PETIT" not in written, written  # expected failure
    assert "Le petit contrat" in written
    vm, p = _pseudonym(db, VM), _pseudonym(db, "Petit")
    assert written == f"{vm} valide. Le petit contrat avec {p}. {p} a signé."


def test_mixed_form_reached(tmp_path: Path, db: Path) -> None:
    """T5 (AC1): [A] « Vorlan Zorbal » also writes « Vorlan ZORBAL »."""
    source, output = _files(tmp_path, ".txt")
    source.write_text(T5_TEXT, encoding="utf-8")
    run = _process(
        source, output, db, [(VM, PERSON)], [ENTER, "a", SPACE], [VZ, PERSON]
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    assert "ZORBAL" not in written, written  # expected failure
    vm, vz = _pseudonym(db, VM), _pseudonym(db, VZ)
    assert written == f"{vm} valide. Le contrat avec {vz} et {vz}."


def test_small_particle_form_reached(tmp_path: Path, db: Path) -> None:
    """T6 (AC1): [A] « VORLAN DE ZORBAL » writes « Vorlan de Zorbal »."""
    source, output = _files(tmp_path, ".txt")
    source.write_text(T6_TEXT, encoding="utf-8")
    run = _process(
        source,
        output,
        db,
        [(VM, PERSON)],
        [ENTER, "a", SPACE],
        ["VORLAN DE ZORBAL", PERSON],
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    assert "Zorbal" not in written, written  # expected failure
    vm, p = _pseudonym(db, VM), _pseudonym(db, "VORLAN DE ZORBAL")
    assert written == f"{vm} valide. Le contrat avec {p}."


# ---------------------------------------------------------------------------
# AC3: a correction is found in its forms; the CLI review's groups
# ---------------------------------------------------------------------------


@pytest.mark.parametrize("ext", [".txt", ".xlsx"])
def test_correction_reaches_capital_place(tmp_path: Path, db: Path, ext: str) -> None:
    """T7 (AC3, CLI [E]): « Zorbal » and « ZORBAL » are one name; corrected as
    « Zorbal Quentrix », no part of « ZORBAL QUENTRIX » stays in clear."""
    source, output = _files(tmp_path, ext)
    _write(source, T7_TEXT, T7_ROWS)
    run = _process(
        source, output, db, [("Zorbal", PERSON), ("ZORBAL", PERSON)], [ENTER, "e"], [ZQ]
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    _pre("(2 occurrences, 1 unique)" in run.flat, "one group shown")
    written = _written(output)
    assert "QUENTRIX" not in written, written  # expected failure
    zq = _pseudonym(db, ZQ)
    if ext == ".xlsx":
        assert _cells(output) == [
            ["Nom", "Note"],
            ["ok", f"{zq} a signé."],
            ["ok", f"Signé : {zq}."],
        ]
    else:
        assert written == f"{zq} a signé. Signé : {zq}."
    assert _names(db) == [(ZQ, PERSON)]


def test_cli_review_groups_accentless_capitals(tmp_path: Path, db: Path) -> None:
    """T17 (AC3, the CLI review's groups): « Vorlène » and « VORLENE » are one
    name; corrected as « Vorlène Quentrix », « VORLENE QUENTRIX » is covered."""
    source, output = _files(tmp_path, ".txt")
    source.write_text(T17_TEXT, encoding="utf-8")
    run = _process(
        source,
        output,
        db,
        [("Vorlène", PERSON), ("VORLENE", PERSON)],
        [ENTER, "e"],
        [VQ],
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    assert "QUENTRIX" not in written, written  # expected failure
    vq = _pseudonym(db, VQ)
    assert written == f"{vq} a signé. Signé : {vq}."
    assert "(2 occurrences, 1 unique)" in run.flat
    assert _names(db) == [(VQ, PERSON)]


def test_cli_decomposed_document(tmp_path: Path, db: Path) -> None:
    """T19 CLI (AC1-AC3 with Story 10.20's readers): T17's document written
    with every accent as a separate mark, the correction typed the same way."""
    source, output = _files(tmp_path, ".txt")
    source.write_text(nfd(T17_TEXT), encoding="utf-8")
    stored = source.read_text(encoding="utf-8")
    _pre(
        stored != T17_TEXT and unicodedata.normalize("NFC", stored) == T17_TEXT,
        "the file holds separate marks",
    )
    detector = FakeDetector([("Vorlène", PERSON), ("VORLENE", PERSON)])
    run = _process(source, output, db, (), [ENTER, "e"], [nfd(VQ)], detector=detector)
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    _pre(
        bool(detector.texts) and all(text == T17_TEXT for text in detector.texts),
        "the detector is given the composed text",
    )
    written = _written(output)
    assert "QUENTRIX" not in written, written  # expected failure
    vq = _pseudonym(db, VQ)
    assert written == f"{vq} a signé. Signé : {vq}."
    assert "(2 occurrences, 1 unique)" in run.flat
    assert _names(db) == [(VQ, PERSON)]


def test_hand_group_joined_by_its_written_texts(tmp_path: Path, db: Path) -> None:
    """T21 CLI (AC3, G8's rule for places added with [A]; the orchestrator's
    ruling, 2026-10-25): [A] « VORLENE » also reaches « Vorlene », which is
    not a form of « Vorlène ». The review's join reads the text written at
    each place, so that name stays apart, and the correction of « Vorlène »
    never writes « Vorlene » with the corrected name's pseudonym."""
    source, output = _files(tmp_path, ".txt")
    source.write_text(T21_TEXT, encoding="utf-8")
    run = _process(
        source,
        output,
        db,
        [(VM, PERSON), ("Vorlène", LOCATION)],
        [ENTER, "a", SPACE, "e", SPACE],
        ["VORLENE", LOCATION, VQ],
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    vq = _pseudonym(db, VQ, LOCATION)
    assert f"avec {vq}" not in written, written  # expected failure
    assert "(3 occurrences, 2 unique)" in run.flat
    vm, caps = _pseudonym(db, VM), _pseudonym(db, "VORLENE", LOCATION)
    assert written == (
        f"{vm} valide. {vq} a signé. Signé : {caps} QUENTRIX. "
        f"Le contrat avec {caps} Quentrix."
    )


def test_join_never_chains(tmp_path: Path, db: Path) -> None:
    """T22 (AC3, Codex P0-1): « MEZOR » is a form of « Mézor » and of
    « Mèzor », but the « Vorlan Mézor » and « Zorbal Mèzor » groups never
    meet through it: three names, and « Zorbal Mèzor » keeps its own
    pseudonym."""
    source, output = _files(tmp_path, ".txt")
    source.write_text(T22_TEXT, encoding="utf-8")
    run = _process(source, output, db, T22_NAMES, [ENTER, "e"], ["Vorlan Mézor"])
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    vm = _pseudonym(db, "Vorlan Mézor")
    assert f"avec {vm}" not in written, written  # expected failure
    assert "(5 occurrences, 3 unique)" in run.flat
    assert _pseudonym(db, "Zorbal Mèzor") != vm


def test_join_never_chains_swapped(tmp_path: Path, db: Path) -> None:
    """T22, swapped (AC3, Codex P0-1; the orchestrator's addition): the full
    name first. « MEZOR » does not join the « Vorlan Mézor » group, so its
    place keeps its own decision."""
    source, output = _files(tmp_path, ".txt")
    source.write_text(T22_SWAPPED, encoding="utf-8")
    run = _process(
        source, output, db, T22_NAMES, [ENTER, "e"], ["Vorlan Mézor Quentrix"]
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    vmq = _pseudonym(db, "Vorlan Mézor Quentrix")
    assert f"{vmq} a signé" not in written, written  # expected failure
    assert "(5 occurrences, 3 unique)" in run.flat
    mezor = _pseudonym(db, "MEZOR")
    assert written.startswith(f"{vmq} valide. Signé : {vmq}. {mezor} a signé. ")


def test_join_never_chains_through_the_set(tmp_path: Path, db: Path) -> None:
    """T22, case 3 (AC3; the orchestrator's addition): « Mézor » joins
    « MEZOR »; « Mèzor », a form of « MEZOR » but not of « Mézor », does not.
    The correction reaches no place of the joined group, so « MEZOR » and
    « Mézor » come back to the review as their own names (T25's rule), each
    corrected; « Mèzor » is accepted."""
    source, output = _files(tmp_path, ".txt")
    source.write_text(T22_SET, encoding="utf-8")
    run = _process(
        source,
        output,
        db,
        [("MEZOR", PERSON), ("Mézor", PERSON), ("Mèzor", PERSON)],
        [ENTER, "e", "e", "e", SPACE],
        ["Vorlan Mézor"] * 3,
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    vm = _pseudonym(db, "Vorlan Mézor")
    assert f"avec {vm}" not in written, written  # expected failure
    assert "(3 occurrences, 2 unique)" in run.flat
    other = _pseudonym(db, "Mèzor")
    assert written == f"{vm} a signé. {vm} valide. Le contrat avec {other}."


def test_hand_place_never_joins_another_name(tmp_path: Path, db: Path) -> None:
    """T24 (AC2, AC3, QA's JOIN-001): [A] « ZÉLVINE » (LOCATION) puts its
    « ZELVINE » place in the « Zélvine » group; « Zelvine », another name
    under Q3, never joins it, so the correction of « Zélvine » never writes
    « Zelvine » with its pseudonym."""
    source, output = _files(tmp_path, ".txt")
    source.write_text(T24_TEXT, encoding="utf-8")
    run = _process(
        source,
        output,
        db,
        [("Brimel Testar", PERSON), ("Zélvine", LOCATION), ("Zelvine", LOCATION)],
        [ENTER, "a", SPACE, "e", SPACE],
        ["ZÉLVINE", LOCATION, "Zélvine Testar"],
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    zt = _pseudonym(db, "Zélvine Testar", LOCATION)
    assert f"avec {zt}" not in written, written  # expected failure
    assert "(3 occurrences, 2 unique)" in run.flat
    assert ("Zelvine", LOCATION) in _names(db)


def test_unreached_place_goes_back_to_review(tmp_path: Path, db: Path) -> None:
    """T25 (AC3, QA's CAPS-001; the orchestrator's ruling): a correction typed
    in capitals without accents does not reach « Zélvine Testar »; that place
    goes back to the review as its own name, never written with a partial
    span, and is corrected there."""
    source, output = _files(tmp_path, ".txt")
    source.write_text(T25_TEXT, encoding="utf-8")
    run = _process(
        source,
        output,
        db,
        [("ZELVINE", PERSON), ("Zélvine", PERSON)],
        [ENTER, "e", "e"],
        ["ZELVINE TESTAR", "Zélvine Testar"],
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    assert "TESTAR" not in written.upper(), written  # expected failure
    assert "(2 occurrences, 1 unique)" in run.flat
    caps, title = _pseudonym(db, "ZELVINE TESTAR"), _pseudonym(db, "Zélvine Testar")
    assert written == f"Signé : {caps}. {title} a signé."


def test_cli_added_words_never_left_in_clear_at_own_places(
    tmp_path: Path, db: Path
) -> None:
    """T30 CLI (AC3, Lionel's answer on the leak, "fix it in 10.19"):
    « ZÉLVINE » and « Zélvine » are one name (no join); corrected as
    « ZELVINE TESTAR », the place « Zélvine » is widened over « Testar »,
    the word the correction adds, so nothing of the name stays in clear."""
    source, output = _files(tmp_path, ".txt")
    source.write_text(T30_TEXT, encoding="utf-8")
    run = _process(
        source,
        output,
        db,
        [("ZÉLVINE", PERSON), ("Zélvine", PERSON)],
        [ENTER, "e"],
        ["ZELVINE TESTAR"],
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    assert "TESTAR" not in written.upper(), written  # expected failure
    p = _pseudonym(db, "ZELVINE TESTAR")
    assert written == f"Signé : {p}. {p} a signé."


def test_cli_widening_never_takes_another_name(tmp_path: Path, db: Path) -> None:
    """T31 CLI (AC3; the third Codex review's P0-1, the orchestrator's ruling
    1): « Zélvine » corrected as « Zélvine Mézor » is never widened over
    « Mèzor », which is neither a Q3 form of « Mézor » nor « Mézor » of it;
    that place keeps its old span, as on main."""
    source, output = _files(tmp_path, ".txt")
    source.write_text(T31_TEXT, encoding="utf-8")
    run = _process(
        source, output, db, [("Zélvine", PERSON)], [ENTER, "e"], ["Zélvine Mézor"]
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    assert "Mèzor a signé" in written, written  # expected failure
    p = _pseudonym(db, "Zélvine Mézor")
    assert written == f"{p} a signé. {p} Mèzor a signé."


def test_cli_widening_across_a_line_break(tmp_path: Path, db: Path) -> None:
    """T35 CLI (AC3; the third Codex review's P0-5, the orchestrator's ruling
    2): the added word written after a line break is covered too."""
    source, output = _files(tmp_path, ".txt")
    source.write_text(T35_TEXT, encoding="utf-8")
    run = _process(
        source,
        output,
        db,
        [("ZÉLVINE", PERSON), ("Zélvine", PERSON)],
        [ENTER, "e"],
        ["ZELVINE TESTAR"],
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    assert "TESTAR" not in written.upper(), written  # expected failure
    p = _pseudonym(db, "ZELVINE TESTAR")
    assert written == f"Signé : {p}. {p} a signé."


def test_cli_repeated_add_recognises_held_capitals_place(
    tmp_path: Path, db: Path
) -> None:
    """T36 CLI (AC2; the third Codex review's P1, the orchestrator's ruling
    6): [A] « Vorlène » holds « VORLENE » too; [A] « VORLENE » then adds
    nothing and shows the approved line, and one row is saved."""
    source, output = _files(tmp_path, ".txt")
    source.write_text(T36_TEXT, encoding="utf-8")
    run = _process(
        source,
        output,
        db,
        [(VM, PERSON)],
        [ENTER, "a", "a", SPACE],
        ["Vorlène", PERSON, "VORLENE", PERSON],
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    line = "'VORLENE' is already in the review as PERSON at every place: nothing added"
    assert line in run.flat, run.shown  # expected failure
    assert _names(db) == sorted([(VM, PERSON), ("Vorlène", PERSON)])


def test_cli_repeated_add_reads_the_text_written_at_a_held_place(
    tmp_path: Path, db: Path
) -> None:
    """T36 CLI, case 2 (AC2; ADD-1019-01 through the text written): [A]
    « Belora » holds « BÊLORA », a capitals form of it; [A] « Bélora »,
    another name, adds « Bélora » only, since « BÊLORA » written there is
    already in the review."""
    source, output = _files(tmp_path, ".txt")
    source.write_text(T36B_TEXT, encoding="utf-8")
    run = _process(
        source,
        output,
        db,
        [(VM, PERSON)],
        [ENTER, "a", "a", SPACE],
        ["Belora", PERSON, "Bélora", PERSON],
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    assert "Processing 3 entities." in run.flat, run.shown  # expected failure
    vm, pb, pe = (
        _pseudonym(db, VM),
        _pseudonym(db, "Belora"),
        _pseudonym(db, "Bélora"),
    )
    assert _written(output) == f"{vm} valide. Signé : {pe}. Signé : {pb}."
    assert _names(db) == sorted([(VM, PERSON), ("Belora", PERSON), ("Bélora", PERSON)])


@pytest.mark.parametrize("quotes", list(QUOTES))
def test_cli_widening_inside_quotation_marks(
    tmp_path: Path, db: Path, quotes: str
) -> None:
    """T37 CLI (AC3; the fourth Codex review's P0-1): « Zélvine Testar »
    between quotation marks (straight, curly, curly double), « Zélvine »
    corrected as « ZELVINE TESTAR »: the widening covers « Testar », the
    closing mark is not part of the word, and the marks stay."""
    opening, closing = QUOTES[quotes]
    source, output = _files(tmp_path, ".txt")
    source.write_text(
        f"Signé : ZÉLVINE TESTAR. {opening}Zélvine Testar{closing} a signé.",
        encoding="utf-8",
    )
    run = _process(
        source,
        output,
        db,
        [("ZÉLVINE", PERSON), ("Zélvine", PERSON)],
        [ENTER, "e"],
        ["ZELVINE TESTAR"],
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    assert "TESTAR" not in written.upper(), written  # expected failure
    p = _pseudonym(db, "ZELVINE TESTAR")
    assert written == f"Signé : {p}. {opening}{p}{closing} a signé."


def test_cli_widening_keeps_the_elision_inside_quotation_marks(
    tmp_path: Path, db: Path
) -> None:
    """T37 CLI, the elision (AC3; the fourth Codex review's P0-1): the
    apostrophe inside « D'Émerine » stays part of the word; only the closing
    quotation mark is taken off."""
    source, output = _files(tmp_path, ".txt")
    source.write_text(
        "Signé : ZÉLVINE D'ÉMERINE. 'Zélvine D'Émerine' a signé.", encoding="utf-8"
    )
    run = _process(
        source,
        output,
        db,
        [("ZÉLVINE", PERSON), ("Zélvine", PERSON)],
        [ENTER, "e"],
        ["ZELVINE D'EMERINE"],
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    assert "EMERINE" not in fold(written).upper(), written  # expected failure
    p = _pseudonym(db, "ZELVINE D'EMERINE")
    assert written == f"Signé : {p}. '{p}' a signé."


@pytest.mark.parametrize("mark", list(PUNCTUATION))
def test_cli_widening_before_any_punctuation(
    tmp_path: Path, db: Path, mark: str
) -> None:
    """T40 CLI (AC3; the fifth Codex review's P0-1, QA's LEAK-1019-03): a word
    is read by what it is, so an ellipsis, a dash, a slash, an asterisk or a
    footnote digit right after « Testar » ends it; « Zélvine » corrected as
    « ZELVINE TESTAR » covers « Testar », and the punctuation stays."""
    punctuation = PUNCTUATION[mark]
    source, output = _files(tmp_path, ".txt")
    source.write_text(
        f"Signé : ZÉLVINE TESTAR. Zélvine Testar{punctuation} a signé.",
        encoding="utf-8",
    )
    run = _process(
        source,
        output,
        db,
        [("ZÉLVINE", PERSON), ("Zélvine", PERSON)],
        [ENTER, "e"],
        ["ZELVINE TESTAR"],
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    assert "TESTAR" not in written.upper(), written  # expected failure
    p = _pseudonym(db, "ZELVINE TESTAR")
    assert written == f"Signé : {p}. {p}{punctuation} a signé."


@pytest.mark.parametrize("case", list(T42_CASES))
def test_cli_widening_reads_the_correction_by_the_same_word_rule(
    tmp_path: Path, db: Path, case: str
) -> None:
    """T42 CLI (AC3; the sixth Codex review's P0-1, QA's LEAK-1019-04): the
    typed correction is read with the document's word rule, and what lies
    between two words is the same on both sides (an initial's full stop
    belongs to the initial): a middle initial (« J. », « J.-P. »,
    « J.-M. »), a comma, an initial typed without its full stops; nothing of
    the name is left, and the pseudonym is written twice."""
    text, found, typed = T42_CASES[case]
    source, output = _files(tmp_path, ".txt")
    source.write_text(text, encoding="utf-8")
    run = _process(
        source,
        output,
        db,
        [(found[0], PERSON), (found[1], PERSON)],
        [ENTER, "e"],
        [typed],
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    left = [word for word in ("Testar", "TESTAR", "Zélvine") if word in written]
    assert not left, written  # expected failure
    # « M. » in « J.-M. » is read as a title when the row is saved, as on
    # main, so the one row is read by its type, not by the text typed
    people = [row.pseudonym_full for row in _entities(db) if row.entity_type == PERSON]
    _pre(len(people) == 1, f"one row stored for {typed!r}: {_names(db)}")
    p = str(people[0])
    assert written == f"Signé : {p}. {p} a signé."


# ---------------------------------------------------------------------------
# AC1, AC5: accents in capitals
# ---------------------------------------------------------------------------


def test_accent_dropped_in_capitals(tmp_path: Path, db: Path) -> None:
    """T10 (AC1, AC5, Q3): [A] « Vorlène » writes « VORLÈNE » and « VORLENE »,
    never « Vorlene »."""
    source, output = _files(tmp_path, ".txt")
    source.write_text(T10_TEXT, encoding="utf-8")
    run = _process(
        source, output, db, [(VM, PERSON)], [ENTER, "a", SPACE], ["Vorlène", PERSON]
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    assert "VORLÈNE" not in written, written  # expected failure
    assert "VORLENE" not in written
    assert "Vorlene a signé." in written
    vm, p = _pseudonym(db, VM), _pseudonym(db, "Vorlène")
    assert written == (
        f"{vm} valide. Le contrat avec {p}. {p} a signé. {p} a signé. "
        "Vorlene a signé."
    )


# ---------------------------------------------------------------------------
# AC4, AC8, AC15: delete-mapping
# ---------------------------------------------------------------------------


def test_delete_mapping_removes_capital_variants(tmp_path: Path, db: Path) -> None:
    """T9 (AC4): « Zorbal », « ZORBAL » and « zorbal » are listed on the screen
    and erased, each with its entry; « Zorbalix » stays."""
    rows = [
        ("Zorbal", PERSON),
        ("ZORBAL", PERSON),
        ("zorbal", PERSON),
        ("Zorbalix", PERSON),
    ]
    _store(db, tmp_path, rows)
    _pre(_names(db) == sorted(rows), "the four rows stored")
    ids = _ids(db)
    run = _delete(db, "Zorbal")
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    assert (  # expected failure
        f"{NAME_LINE}ZORBAL" in run.screen and f"{NAME_LINE}zorbal" in run.screen
    ), run.shown
    assert _names(db) == [("Zorbalix", PERSON)]
    assert _erased_ids(db) == sorted(ids[row] for row in rows[:3])
    assert all("zorbal" not in raw.lower() for raw in _erasures(db))


def test_screen_line_for_several(tmp_path: Path, db: Path) -> None:
    """T11 (AC8, Q2): two saved pseudonyms of the name in two capitals: Lionel's
    approved line, with the count."""
    _store(db, tmp_path, [("Zorbal", PERSON), ("ZORBAL", PERSON), ("Zorbalix", PERSON)])
    _pre(len(_names(db)) == 3, "three rows stored")
    run = _delete(db, "Zorbal")
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    assert SEVERAL_LINE in run.flat, run.shown  # expected failure


def test_force_names_every_deleted(tmp_path: Path, db: Path) -> None:
    """T12 (AC4, the PO's S1): with ``--force``, no screen, and the success lines
    name every saved pseudonym deleted."""
    _store(db, tmp_path, [("Zorbal", PERSON), ("ZORBAL", PERSON), ("Zorbalix", PERSON)])
    _pre(len(_names(db)) == 3, "three rows stored")
    run = _delete(db, "Zorbal", "--force", typed=())
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    _pre("WARNING: PERMANENT DELETION" not in run.shown, "no screen")
    assert "Entity: ZORBAL (PERSON)" in run.shown, run.shown  # expected failure
    assert "Entity: Zorbal (PERSON)" in run.shown
    assert _names(db) == [("Zorbalix", PERSON)]
    assert len(_erasures(db)) == 2


def test_delete_mapping_accent_dropped_in_capitals(tmp_path: Path, db: Path) -> None:
    """T13 (AC4, Q3): « VORLENE » is listed and erased with « Vorlène »;
    « Vorlene » and « Zorbalix » stay."""
    _store(
        db,
        tmp_path,
        [
            ("Vorlène", PERSON),
            ("VORLENE", PERSON),
            ("Vorlene", PERSON),
            ("Zorbalix", PERSON),
        ],
    )
    _pre(len(_names(db)) == 4, "four rows stored")
    run = _delete(db, "Vorlène")
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    assert f"{NAME_LINE}VORLENE" in run.screen, run.shown  # expected failure
    assert _names(db) == sorted([("Vorlene", PERSON), ("Zorbalix", PERSON)])
    assert len(_erasures(db)) == 2


@pytest.mark.parametrize("left", ["one", "two"])
def test_id_line_says_mapping(tmp_path: Path, db: Path, left: str) -> None:
    """T18 (AC15, Lionel's Q5b): after ``delete-mapping --id``, the line about
    the others left says « mapping », byte for byte."""
    types = [LOCATION, ORG] if left == "one" else [LOCATION, ORG, PERSON]
    _store(db, tmp_path, [("Zorbal", entity_type) for entity_type in types])
    _pre(
        _names(db) == sorted(("Zorbal", entity_type) for entity_type in types),
        "one saved pseudonym per type",
    )
    deleted = LOCATION if left == "one" else PERSON
    run = _delete(db, "--id", _ids(db)[("Zorbal", deleted)], "--force", typed=())
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    _pre(len(_erasures(db)) == 1, "one ERASURE entry")
    expected = ID_LINE_ONE if left == "one" else ID_LINE_TWO
    assert expected in run.after, run.shown  # expected failure


def test_delete_mapping_reaches_old_damaged_row(tmp_path: Path, db: Path) -> None:
    """T20 (AC4, Q1 with Lionel's 10.20 Qa): the row v2.2.0 saved from
    « QUOVÔME ZORBAL » read with separate marks is listed and erased with
    today's row of « Quovôme Zorbal »."""
    circumflex = unicodedata.lookup("COMBINING CIRCUMFLEX ACCENT")
    old = _key_steps(nfd("QUOVÔME ZORBAL"), PERSON)
    _pre(old == "QUOVO" + circumflex + "ZORBAL", "the old saved text")
    _save(db, [(QZ, PERSON, VM), (old, PERSON, "Zorbalix")])
    _pre(_names(db) == sorted([(QZ, PERSON), (old, PERSON)]), "both rows saved")
    ids = _ids(db)
    run = _delete(db, QZ)
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    assert f"{NAME_LINE}{old}" in run.screen, run.shown  # expected failure
    assert f"{NAME_LINE}{QZ}" in run.screen
    assert _names(db) == []
    assert _erased_ids(db) == sorted([ids[(QZ, PERSON)], ids[(old, PERSON)]])


def test_delete_mapping_reaches_old_row_in_capitals_without_accent(
    tmp_path: Path, db: Path
) -> None:
    """T23 (AC4, Codex P0-2): « BELORA MÉLORA », a Q3 form of « Bélora
    Mélora », read with separate marks by v2.2.0, was saved « BELORA » + a
    lone mark + « LORA »; deleting « Bélora Mélora » erases it with today's
    row, and leaves « Vorlan Mélora »."""
    acute = unicodedata.lookup("COMBINING ACUTE ACCENT")
    old = _key_steps(nfd("BELORA MÉLORA"), PERSON)
    _pre(old == "BELORA " + acute + "LORA", "the old saved text")
    _save(
        db,
        [
            ("Bélora Mélora", PERSON, VM),
            (old, PERSON, "Zorbalix"),
            ("Vorlan Mélora", PERSON, "Quovane Testar"),
        ],
    )
    _pre(len(_names(db)) == 3, "three rows saved")
    ids = _ids(db)
    run = _delete(db, "Bélora Mélora")
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    assert f"{NAME_LINE}{old}" in run.screen, run.shown  # expected failure
    assert f"{NAME_LINE}Bélora Mélora" in run.screen
    assert _names(db) == [("Vorlan Mélora", PERSON)]
    assert _erased_ids(db) == sorted(
        [ids[("Bélora Mélora", PERSON)], ids[(old, PERSON)]]
    )


def test_delete_mapping_reaches_old_row_with_another_accent(
    tmp_path: Path, db: Path
) -> None:
    """T28 (AC4, the second Codex review's third P0): « VORLAN MÈLORA », a Q3
    form of « Vorlan Mélora », read with separate marks by v2.2.0, was saved
    « VORLAN » + a lone grave accent + « LORA »; deleting « Vorlan Mélora »
    erases it with today's row, and leaves « Bélora Quentrel »."""
    grave = unicodedata.lookup("COMBINING GRAVE ACCENT")
    old = _key_steps(nfd("VORLAN MÈLORA"), PERSON)
    _pre(old == "VORLAN " + grave + "LORA", "the old saved text")
    _save(
        db,
        [
            ("Vorlan Mélora", PERSON, VM),
            (old, PERSON, "Zorbalix"),
            ("Bélora Quentrel", PERSON, "Quovane Testar"),
        ],
    )
    _pre(len(_names(db)) == 3, "three rows saved")
    ids = _ids(db)
    run = _delete(db, "Vorlan Mélora")
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    assert f"{NAME_LINE}{old}" in run.screen, run.shown  # expected failure
    assert f"{NAME_LINE}Vorlan Mélora" in run.screen
    assert _names(db) == [("Bélora Quentrel", PERSON)]
    assert _erased_ids(db) == sorted(
        [ids[("Vorlan Mélora", PERSON)], ids[(old, PERSON)]]
    )


def test_delete_mapping_typed_without_accents_reaches_old_capitals_row(
    tmp_path: Path, db: Path
) -> None:
    """T28, case 2 (AC4, QA's DEL-1019-01 second shape): typed without its
    accents, « Vorlan Melora » erases « VORLAN MÉLORA » (a form once accents
    are set aside) and the old row v2.2.0 saved of it, « VORLAN » + a lone
    mark + « LORA »; « Vorlan Mélora » in title case is another name, and its
    rows stay, today's and the old one."""
    acute = unicodedata.lookup("COMBINING ACUTE ACCENT")
    old_caps = _key_steps(nfd("VORLAN MÉLORA"), PERSON)
    old_title = _key_steps(nfd("Vorlan Mélora"), PERSON)
    _pre(
        old_caps == "VORLAN " + acute + "LORA"
        and old_title == "Vorlan " + acute + "lora",
        "the two old saved texts",
    )
    _save(
        db,
        [
            ("Vorlan Melora", PERSON, VM),
            ("VORLAN MÉLORA", PERSON, "Zorbalix"),
            (old_caps, PERSON, "Quovane Testar"),
            ("Vorlan Mélora", PERSON, "Brimel Testar"),
            (old_title, PERSON, "Zorbal Quentrix"),
        ],
    )
    _pre(len(_names(db)) == 5, "five rows saved")
    ids = _ids(db)
    run = _delete(db, "Vorlan Melora")
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    assert f"{NAME_LINE}{old_caps}" in run.screen, run.shown  # expected failure
    assert _names(db) == sorted([("Vorlan Mélora", PERSON), (old_title, PERSON)])
    assert _erased_ids(db) == sorted(
        [
            ids[("Vorlan Melora", PERSON)],
            ids[("VORLAN MÉLORA", PERSON)],
            ids[(old_caps, PERSON)],
        ]
    )


def test_delete_mapping_keeps_title_case_old_row_of_another_accent(
    tmp_path: Path, db: Path
) -> None:
    """T33 (AC4; the third Codex review's P0-3, the orchestrator's ruling 4):
    « Vorlan Mèlora » in title case is another name than « Vorlan Mélora »
    (Q3 sets accents aside in capitals only), so the row v2.2.0 saved of it,
    « Vorlan » + a lone grave accent + « lora », stays."""
    grave = unicodedata.lookup("COMBINING GRAVE ACCENT")
    old = _key_steps(nfd("Vorlan Mèlora"), PERSON)
    _pre(old == "Vorlan " + grave + "lora", "the old saved text")
    _save(db, [("Vorlan Mélora", PERSON, VM), (old, PERSON, "Zorbalix")])
    _pre(len(_names(db)) == 2, "two rows saved")
    ids = _ids(db)
    run = _delete(db, "Vorlan Mélora")
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    assert _names(db) == [(old, PERSON)], run.shown  # expected failure
    assert _erased_ids(db) == [ids[("Vorlan Mélora", PERSON)]]


def test_delete_mapping_reaches_old_hyphen_row_typed_without_accents(
    tmp_path: Path, db: Path
) -> None:
    """T34 (AC4; the third Codex review's P0-4, the orchestrator's ruling 5):
    v2.2.0 cut « ME » after the hyphen of « BELORA-MÉLORA » and at the end of
    « QUOVÔME »; deleting « Belora-Melora Quovome Zorbal » erases that row
    with today's row in capitals, and leaves « Vorlan Mestrix »."""
    acute = unicodedata.lookup("COMBINING ACUTE ACCENT")
    circumflex = unicodedata.lookup("COMBINING CIRCUMFLEX ACCENT")
    caps = "BELORA-MÉLORA QUOVÔME ZORBAL"
    old = _key_steps(nfd(caps), PERSON)
    _pre(
        old == "BELORA-" + acute + "LORA QUOVO" + circumflex + "ZORBAL",
        "the old saved text",
    )
    _save(
        db,
        [
            (caps, PERSON, VM),
            (old, PERSON, "Zorbalix"),
            ("Vorlan Mestrix", PERSON, "Quovane Testar"),
        ],
    )
    _pre(len(_names(db)) == 3, "three rows saved")
    ids = _ids(db)
    run = _delete(db, "Belora-Melora Quovome Zorbal")
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    assert _names(db) == [("Vorlan Mestrix", PERSON)], run.shown  # expected failure
    assert _erased_ids(db) == sorted([ids[(caps, PERSON)], ids[(old, PERSON)]])


def test_delete_mapping_reaches_old_row_with_two_marks_side_by_side(
    tmp_path: Path, db: Path
) -> None:
    """T38 (AC4; the fourth Codex review's P0-2): v2.2.0 cut « ME » at the end
    of « QUOVÔME » and at the start of « MÉLORA », leaving two marks side by
    side; deleting « Quovome Melora » erases that row with today's."""
    circumflex = unicodedata.lookup("COMBINING CIRCUMFLEX ACCENT")
    acute = unicodedata.lookup("COMBINING ACUTE ACCENT")
    document = "QUOVÔME MÉLORA"
    old = _key_steps(nfd(document), PERSON)
    _pre(old == "QUOVO" + circumflex + acute + "LORA", "the old saved text")
    left, erased, ids, _ = _old_rows_erased(db, document, "Quovome Melora")
    assert left == [("Vorlan Mestrix", PERSON)], left  # expected failure
    assert erased == sorted([ids[(document, PERSON)], ids[(old, PERSON)]])


def test_delete_mapping_reaches_old_row_past_six_spots(
    tmp_path: Path, db: Path
) -> None:
    """T39 (AC4; the fourth Codex review's P0-3): the first three words hold
    six possible cut places; the row v2.2.0 saved, cut at the last word, is
    erased with « Melome Melame Melime Melume »."""
    acute = unicodedata.lookup("COMBINING ACUTE ACCENT")
    document = "MELOME MELAME MELIME MÉLUME"
    old = _key_steps(nfd(document), PERSON)
    _pre(old == "MELOME MELAME MELIME " + acute + "LUME", "the old saved text")
    left, erased, ids, _ = _old_rows_erased(db, document, "Melome Melame Melime Melume")
    assert left == [("Vorlan Mestrix", PERSON)], left  # expected failure
    assert erased == sorted([ids[(document, PERSON)], ids[(old, PERSON)]])


def test_delete_mapping_reaches_old_row_past_eight_places(
    tmp_path: Path, db: Path
) -> None:
    """T41 (AC4; the fifth Codex review's P0-2): the first words, one of them
    hyphenated, hold eight possible cut places before the surname's real
    cut; the row v2.2.0 saved is erased with « Melome Melame Melime-Melume
    Melora », whatever the number of places."""
    acute = unicodedata.lookup("COMBINING ACUTE ACCENT")
    document = "MELOME MELAME MELIME-MELUME MÉLORA"
    old = _key_steps(nfd(document), PERSON)
    _pre(old == "MELOME MELAME MELIME-MELUME " + acute + "LORA", "the old saved text")
    left, erased, ids, _ = _old_rows_erased(
        db, document, "Melome Melame Melime-Melume Melora"
    )
    assert left == [("Vorlan Mestrix", PERSON)], left  # expected failure
    assert erased == sorted([ids[(document, PERSON)], ids[(old, PERSON)]])


def test_delete_mapping_reaches_old_row_with_a_repeated_cut(
    tmp_path: Path, db: Path
) -> None:
    """T43 (AC4; the sixth Codex review's P0-2): v2.2.0 cut « MÉ » twice from
    « MÉMÉLIE TESTAR », leaving two marks + « LIE TESTAR »; the row is
    erased with « Memelie Testar », with today's row of that text."""
    acute = unicodedata.lookup("COMBINING ACUTE ACCENT")
    document = "MÉMÉLIE TESTAR"
    old = _key_steps(nfd(document), PERSON)
    _pre(old == acute + acute + "LIE TESTAR", "the old saved text")
    left, erased, ids, _ = _old_rows_erased(db, document, "Memelie Testar")
    assert left == [("Vorlan Mestrix", PERSON)], left  # expected failure
    assert erased == sorted([ids[(document, PERSON)], ids[(old, PERSON)]])


@pytest.mark.parametrize(
    "typed",
    [
        pytest.param("Clemelie Mezor", id="typed_without_accents"),
        pytest.param("Clémélie Mézor", id="typed_with_accents"),
    ],
)
def test_delete_mapping_reaches_old_row_cut_after_an_own_accent(
    tmp_path: Path, db: Path, typed: str
) -> None:
    """T44 (AC4; DEL-1019-07): v2.2.0 cut « MÉ » from « CLÉMÉLIE » right after
    the « É » of « CLÉ » (its own accent made the word's bound) and from
    « MÉZOR »: « CLE » + two marks + « LIE », a mark + « ZOR ». The row is
    erased typed without the accents (expected failure on 63da306) and with
    them (a guard: 10.20 already reaches it)."""
    acute = unicodedata.lookup("COMBINING ACUTE ACCENT")
    document = "CLÉMÉLIE MÉZOR"
    old = _key_steps(nfd(document), PERSON)
    _pre(old == "CLE" + acute + acute + "LIE " + acute + "ZOR", "the old saved text")
    left, erased, ids, _ = _old_rows_erased(db, document, typed)
    assert left == [("Vorlan Mestrix", PERSON)], left  # expected failure
    assert erased == sorted([ids[(document, PERSON)], ids[(old, PERSON)]])


@pytest.mark.parametrize(
    "typed",
    [
        pytest.param("Meme Testar", id="typed_without_accents"),
        pytest.param("Mème Testar", id="typed_with_accents"),
    ],
)
def test_delete_mapping_reaches_old_row_cut_on_both_sides_of_one_mark(
    tmp_path: Path, db: Path, typed: str
) -> None:
    """T45 (AC4; the seventh Codex review's P0, DEL-1019-08): v2.2.0 cut « ME »
    from « MÈME TESTAR » before the « È »'s mark, then the second « ME »
    after it, with the space, leaving the mark + « TESTAR ». The row is
    erased typed without the accents (expected failure on 46202a3, a
    regression from 63da306) and with them (a guard: 10.20 reaches it)."""
    grave = unicodedata.lookup("COMBINING GRAVE ACCENT")
    document = "MÈME TESTAR"
    old = _key_steps(nfd(document), PERSON)
    _pre(old == grave + "TESTAR", "the old saved text")
    left, erased, ids, _ = _old_rows_erased(db, document, typed)
    assert left == [("Vorlan Mestrix", PERSON)], left  # expected failure
    assert erased == sorted([ids[(document, PERSON)], ids[(old, PERSON)]])


def test_delete_mapping_never_reaches_another_names_old_row_through_a_consonant(
    tmp_path: Path, db: Path
) -> None:
    """T46 (AC4; found measuring v1.23): « Drlime » is not « MÉLIME », whose
    v2.2.0 row is a mark + « LIME »; deleting « Drlime » erases its own row
    only, since no French letter « R » carries an accent the old step could
    have bounded « DR » with."""
    acute = unicodedata.lookup("COMBINING ACUTE ACCENT")
    old = _key_steps(nfd("MÉLIME"), PERSON)
    _pre(old == acute + "LIME", "the old saved text")
    _save(
        db,
        [
            ("Drlime", PERSON, "Quovane Testar"),
            ("MÉLIME", PERSON, VM),
            (old, PERSON, "Zorbalix"),
        ],
    )
    _pre(len(_names(db)) == 3, "three rows saved")
    ids = _ids(db)
    run = _delete(db, "Drlime")
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    left = _names(db)
    assert left == sorted([("MÉLIME", PERSON), (old, PERSON)]), left  # expected failure
    assert _erased_ids(db) == [ids[("Drlime", PERSON)]]


def _inserted_rows_erased(
    db: Path, today: str, source: str, typed: str
) -> tuple[list[tuple[str, str]], list[str], dict[tuple[str, str], str], str]:
    """Saves today's row of ``today``, the row the old steps saved from
    ``source`` (the name with a word 10.20's replay inserts) read with
    separate marks, and ('Vorlan Mestrix', PERSON); deletes ``typed``.
    Returns the names left, the erased ids, the ids and the old text."""
    old = _key_steps(nfd(source), PERSON)
    _save(
        db,
        [
            (today, PERSON, VM),
            (old, PERSON, "Zorbalix"),
            ("Vorlan Mestrix", PERSON, "Quovane Testar"),
        ],
    )
    _pre(len(_names(db)) == 3, "three rows saved")
    ids = _ids(db)
    run = _delete(db, typed)
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    return _names(db), _erased_ids(db), ids, old


@pytest.mark.parametrize(
    "typed",
    [
        pytest.param("Quovome Testar", id="typed_without_accents"),
        pytest.param("Quovôme Testar", id="typed_with_accents"),
    ],
)
def test_delete_mapping_reaches_old_row_with_maitre_glued_after_a_cut(
    tmp_path: Path, db: Path, typed: str
) -> None:
    """T47 (AC4; the eighth Codex review's P0, DEL-1019-09): the old steps'
    row of « QUOVÔME MAÎTRE TESTAR » read with separate marks, « QUOVO » + a
    mark + « MAI » + a mark + « TRE TESTAR » (« ME » cut, the inserted
    « MAÎTRE » glued to the mark), is erased with today's « QUOVÔME
    TESTAR », typed without the accents (expected failure on c2379c4, a
    regression from 6d87d22) and with them (a guard)."""
    circumflex = unicodedata.lookup("COMBINING CIRCUMFLEX ACCENT")
    today = "QUOVÔME TESTAR"
    left, erased, ids, old = _inserted_rows_erased(
        db, today, "QUOVÔME MAÎTRE TESTAR", typed
    )
    _pre(
        old == "QUOVO" + circumflex + "MAI" + circumflex + "TRE TESTAR",
        "the old saved text",
    )
    assert left == [("Vorlan Mestrix", PERSON)], left  # expected failure
    assert erased == sorted([ids[(today, PERSON)], ids[(old, PERSON)]])


@pytest.mark.parametrize(
    "typed",
    [
        pytest.param("Melora Vorlan", id="typed_without_accents"),
        pytest.param("Mélora Vorlan", id="typed_with_accents"),
    ],
)
def test_delete_mapping_reaches_old_row_ending_with_maitre(
    tmp_path: Path, db: Path, typed: str
) -> None:
    """T48 (AC4; found measuring v1.25, a regression from 6d87d22): the old
    steps' row of « MÉLORA VORLAN MAÎTRE », a mark + « LORA VORLAN MAI » + a
    mark + « TRE » (the inserted word kept after the name's last word), is
    erased with today's « MÉLORA VORLAN », typed without the accents
    (expected failure on c2379c4) and with them (a guard)."""
    acute = unicodedata.lookup("COMBINING ACUTE ACCENT")
    circumflex = unicodedata.lookup("COMBINING CIRCUMFLEX ACCENT")
    today = "MÉLORA VORLAN"
    left, erased, ids, old = _inserted_rows_erased(
        db, today, "MÉLORA VORLAN MAÎTRE", typed
    )
    _pre(
        old == acute + "LORA VORLAN MAI" + circumflex + "TRE",
        "the old saved text",
    )
    assert left == [("Vorlan Mestrix", PERSON)], left  # expected failure
    assert erased == sorted([ids[(today, PERSON)], ids[(old, PERSON)]])


@pytest.mark.parametrize(
    "typed",
    [
        pytest.param("Éloria-Mèlora Mmé Vorlan", id="typed_with_accents"),
        pytest.param("Eloria-Melora Mme Vorlan", id="typed_without_accents"),
    ],
)
def test_delete_mapping_reaches_old_row_that_lost_a_whole_title_word(
    tmp_path: Path, db: Path, typed: str
) -> None:
    """T49 (AC4; found measuring v1.25, a regression from 6d87d22): the old
    steps' row of « ELORIA-MELORA MAÎTRE MME VORLAN », « ELORIA-MELORA MAI »
    + a mark + « TRE VORLAN » (« MME » lost whole, no letter of the name
    accented), is erased with today's « ÉLORIA-MÈLORA MMÉ VORLAN », typed
    with the accents (expected failure on c2379c4) and without them (a
    guard)."""
    circumflex = unicodedata.lookup("COMBINING CIRCUMFLEX ACCENT")
    today = "ÉLORIA-MÈLORA MMÉ VORLAN"
    left, erased, ids, old = _inserted_rows_erased(
        db, today, "ELORIA-MELORA MAÎTRE MME VORLAN", typed
    )
    _pre(
        old == "ELORIA-MELORA MAI" + circumflex + "TRE VORLAN",
        "the old saved text",
    )
    assert left == [("Vorlan Mestrix", PERSON)], left  # expected failure
    assert erased == sorted([ids[(today, PERSON)], ids[(old, PERSON)]])


# ---------------------------------------------------------------------------
# Guards (pass on the unfixed and the fixed code)
# ---------------------------------------------------------------------------


def test_hyphen_compound_not_a_form(tmp_path: Path, db: Path) -> None:
    """G1 (AC1): « SAINT-ZORBAL » is not a place of « Zorbal » (a hyphen joins
    words)."""
    source, output = _files(tmp_path, ".txt")
    source.write_text(G1_TEXT, encoding="utf-8")
    run = _process(
        source, output, db, [(VM, PERSON)], [ENTER, "a", SPACE], ["Zorbal", PERSON]
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    vm, z = _pseudonym(db, VM), _pseudonym(db, "Zorbal")
    assert _written(output) == f"{vm} valide. Le contrat avec {z} et SAINT-ZORBAL."


def test_delete_mapping_without_variant_unchanged(tmp_path: Path, db: Path) -> None:
    """G2 (AC12): a name with no other spelling: the one-row screen, that row
    erased with its entry."""
    _store(db, tmp_path, [(VM, PERSON), ("Zorbalix", PERSON)])
    _pre(len(_names(db)) == 2, "two rows stored")
    run = _delete(db, VM)
    assert run.exit_code == 0, run.shown
    assert ONE_ROW_LINE in run.flat
    assert _names(db) == [("Zorbalix", PERSON)]
    assert len(_erasures(db)) == 1


def test_cli_review_keeps_other_letters_apart(tmp_path: Path, db: Path) -> None:
    """G7 (AC1, AC3; Q3): « Vorlene » (not in capitals, without its accent) is
    another name: two names, and the correction of « Vorlène » leaves it."""
    source, output = _files(tmp_path, ".txt")
    source.write_text(G7_TEXT, encoding="utf-8")
    run = _process(
        source,
        output,
        db,
        [("Vorlène", PERSON), ("Vorlene", PERSON)],
        [ENTER, "e", SPACE],
        [VQ],
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    assert "(2 occurrences, 2 unique)" in run.flat, run.shown
    vq, other = _pseudonym(db, VQ), _pseudonym(db, "Vorlene")
    assert _written(output) == f"{vq} a signé. Signé : {other} Quentrix."
    assert _names(db) == sorted([(VQ, PERSON), ("Vorlene", PERSON)])


def test_cli_three_spellings_keep_other_letters_apart(tmp_path: Path, db: Path) -> None:
    """G8 CLI (AC1, AC3; the orchestrator's ruling on the PO's v1.2 M1): with
    « Vorlène », « VORLENE » and « Vorlene », « VORLENE » stays with « Vorlene »
    as today, and « Vorlene » never takes the corrected name's pseudonym."""
    source, output = _files(tmp_path, ".txt")
    source.write_text(G8_TEXT, encoding="utf-8")
    run = _process(
        source,
        output,
        db,
        [("Vorlène", PERSON), ("VORLENE", PERSON), ("Vorlene", PERSON)],
        [ENTER, "e", SPACE],
        [VQ],
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    vq = _pseudonym(db, VQ)
    assert f"avec {vq}" not in written, written
    assert "(3 occurrences, 2 unique)" in run.flat
    caps, other = _pseudonym(db, "VORLENE"), _pseudonym(db, "Vorlene")
    assert written == (
        f"{vq} a signé. Signé : {caps} QUENTRIX. Le contrat avec {other} Quentrix."
    )
    assert _names(db) == sorted(
        [(VQ, PERSON), ("VORLENE", PERSON), ("Vorlene", PERSON)]
    )


def test_delete_mapping_keeps_name_saved_today_equal_to_lost_letter_form(
    tmp_path: Path, db: Path
) -> None:
    """T23, case 2 (AC4, the sign rule; the orchestrator's addition): the old
    row « QUOVO » + a mark + « ZORBAL » is « QUOVÔZORBAL » once composed, but
    « QUOVÔZORBAL » saved today, with no mark, is another person: deleting
    « Quovôme Zorbal » erases the old row and today's row of the name, never
    that one."""
    circumflex = unicodedata.lookup("COMBINING CIRCUMFLEX ACCENT")
    old = _key_steps(nfd("QUOVÔME ZORBAL"), PERSON)
    _pre(
        old == "QUOVO" + circumflex + "ZORBAL"
        and unicodedata.normalize("NFC", old) == "QUOVÔZORBAL",
        "the old saved text and its composed form",
    )
    _save(
        db,
        [
            (QZ, PERSON, VM),
            (old, PERSON, "Zorbalix"),
            ("QUOVÔZORBAL", PERSON, "Quovane Testar"),
        ],
    )
    _pre(len(_names(db)) == 3, "three rows saved")
    ids = _ids(db)
    run = _delete(db, QZ)
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    left = [(e.full_name, e.entity_type, e.pseudonym_full) for e in _entities(db)]
    assert left == [("QUOVÔZORBAL", PERSON, "Quovane Testar")], run.shown
    assert f"{NAME_LINE}QUOVÔZORBAL" not in run.screen
    assert f"{NAME_LINE}{old}" in run.screen and f"{NAME_LINE}{QZ}" in run.screen
    assert _erased_ids(db) == sorted([ids[(QZ, PERSON)], ids[(old, PERSON)]])


def test_capitals_first_joined(tmp_path: Path, db: Path) -> None:
    """G11 (AC3, QA's TEST-1019-01): the capitals form first in the document;
    « VORLENE » and « Vorlène » are one name, and the correction covers both
    places whole."""
    source, output = _files(tmp_path, ".txt")
    source.write_text(G11_TEXT, encoding="utf-8")
    run = _process(
        source,
        output,
        db,
        [("VORLENE", PERSON), ("Vorlène", PERSON)],
        [ENTER, "e"],
        [VQ],
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    assert "Quentrix" not in written and "QUENTRIX" not in written, written
    assert "(2 occurrences, 1 unique)" in run.flat
    vq = _pseudonym(db, VQ)
    assert written == f"Signé : {vq}. {vq} a signé."


def test_delete_mapping_reaches_row_with_no_mark_that_keeps_every_word(
    tmp_path: Path, db: Path
) -> None:
    """G12 (AC4, QA's TEST-1019-02): the row v2.2.0 saved from « A
    ZORBAVILLE », the capitals form of « à Zorbaville » without its accent,
    carries no mark; it keeps every word of « Zorbaville », so deleting
    « Zorbaville » erases it with today's row, and leaves « Zorbavillette »."""
    old = _key_steps(nfd("A ZORBAVILLE"), LOCATION)
    _pre(
        old == "A ZORBAVILLE" and not any(unicodedata.combining(char) for char in old),
        "the old saved text, with no mark",
    )
    _save(
        db,
        [
            ("Zorbaville", LOCATION, "Xantorville"),
            (old, LOCATION, "Quentrelville"),
            ("Zorbavillette", LOCATION, "Vorlanville"),
        ],
    )
    _pre(len(_names(db)) == 3, "three rows saved")
    ids = _ids(db)
    run = _delete(db, "Zorbaville")
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    assert _names(db) == [("Zorbavillette", LOCATION)], run.shown
    assert _erased_ids(db) == sorted(
        [ids[("Zorbaville", LOCATION)], ids[(old, LOCATION)]]
    )


def test_cli_correction_reaches_capitals_form_before_found_place(
    tmp_path: Path, db: Path
) -> None:
    """G13 (AC3, QA's TEST-1019-03): « QUENTRIX » found; the correction
    « Zorbal Quentrix » reaches its capitals form « ZORBAL QUENTRIX », which
    starts before the found place."""
    source, output = _files(tmp_path, ".txt")
    source.write_text(G13_TEXT, encoding="utf-8")
    run = _process(source, output, db, [("QUENTRIX", PERSON)], [ENTER, "e"], [ZQ])
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    assert "ZORBAL" not in written, written
    p = _pseudonym(db, ZQ)
    assert written == f"Le contrat est validé. Signé : {p}."


def _old_rows_erased(
    db: Path, document: str, typed: str
) -> tuple[list[tuple[str, str]], list[str], dict[tuple[str, str], str], str]:
    """Saves today's row of ``document``, the row v2.2.0 saved from it read
    with separate marks, and ('Vorlan Mestrix', PERSON); deletes ``typed``.
    Returns the names left, the erased ids, the ids and the old text."""
    old = _key_steps(nfd(document), PERSON)
    _save(
        db,
        [
            (document, PERSON, VM),
            (old, PERSON, "Zorbalix"),
            ("Vorlan Mestrix", PERSON, "Quovane Testar"),
        ],
    )
    _pre(len(_names(db)) == 3, "three rows saved")
    ids = _ids(db)
    run = _delete(db, typed)
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    return _names(db), _erased_ids(db), ids, old


def test_delete_mapping_reaches_old_row_with_two_cuts(tmp_path: Path, db: Path) -> None:
    """G14 (AC4; QA's TEST-1019-04): v2.2.0 saved « MÉLORA MÉZOR » cut twice;
    deleting « Melora Mezor » erases it with today's row."""
    acute = unicodedata.lookup("COMBINING ACUTE ACCENT")
    left, erased, ids, old = _old_rows_erased(db, "MÉLORA MÉZOR", "Melora Mezor")
    _pre(old == acute + "LORA " + acute + "ZOR", "the old saved text")
    assert left == [("Vorlan Mestrix", PERSON)]
    assert erased == sorted([ids[("MÉLORA MÉZOR", PERSON)], ids[(old, PERSON)]])


def test_delete_mapping_reaches_old_row_cut_at_word_end(
    tmp_path: Path, db: Path
) -> None:
    """G15 (AC4; QA's TEST-1019-04): v2.2.0 cut « ME » at the end of
    « QUOVÔME »; deleting « Quovome Zorbal » erases that row with today's."""
    circumflex = unicodedata.lookup("COMBINING CIRCUMFLEX ACCENT")
    left, erased, ids, old = _old_rows_erased(db, "QUOVÔME ZORBAL", "Quovome Zorbal")
    _pre(old == "QUOVO" + circumflex + "ZORBAL", "the old saved text")
    assert left == [("Vorlan Mestrix", PERSON)]
    assert erased == sorted([ids[("QUOVÔME ZORBAL", PERSON)], ids[(old, PERSON)]])


def _widened(tmp_path: Path, db: Path, text: str, typed: str) -> tuple[str, str]:
    """« ZÉLVINE » and « Zélvine » found as one name, corrected as ``typed``:
    the written text and the corrected name's pseudonym."""
    source, output = _files(tmp_path, ".txt")
    source.write_text(text, encoding="utf-8")
    run = _process(
        source,
        output,
        db,
        [("ZÉLVINE", PERSON), ("Zélvine", PERSON)],
        [ENTER, "e"],
        [typed],
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    return _written(output), _pseudonym(db, typed)


def test_widening_never_over_a_word_in_small_letters(tmp_path: Path, db: Path) -> None:
    """G16 (AC1, AC3; QA's TEST-1019-04): « testar » in small letters is never
    covered, as on main."""
    written, p = _widened(
        tmp_path,
        db,
        "Signé : ZÉLVINE TESTAR. Zélvine testar a signé.",
        "ZELVINE TESTAR",
    )
    assert written == f"Signé : {p}. {p} testar a signé."


def test_widening_over_capitals_form_of_accented_word(tmp_path: Path, db: Path) -> None:
    """G17 (AC3; QA's TEST-1019-04, the orchestrator's ruling 1): the typed
    « TESTAR » is a Q3 form of the written « Tèstar », so it is covered."""
    written, p = _widened(
        tmp_path,
        db,
        "Signé : ZÉLVINE TÈSTAR. Zélvine Tèstar a signé.",
        "ZELVINE TESTAR",
    )
    assert written == f"Signé : {p}. {p} a signé."


def test_widening_over_every_added_word(tmp_path: Path, db: Path) -> None:
    """G18 (AC3; QA's TEST-1019-04): two added words are both covered."""
    written, p = _widened(
        tmp_path,
        db,
        "Signé : ZÉLVINE BRIMEL TESTAR. Zélvine Brimel Testar a signé.",
        "ZELVINE BRIMEL TESTAR",
    )
    assert written == f"Signé : {p}. {p} a signé."


def test_widening_needs_the_held_words_first(tmp_path: Path, db: Path) -> None:
    """G19 (AC3; QA's TEST-1019-05): « Zélvine » corrected as « Quovane
    Testar » does not start with the word the place holds, so nothing is
    widened: « TESTAR » and « Testar » stay as on main."""
    written, p = _widened(
        tmp_path,
        db,
        "Signé : ZÉLVINE TESTAR. Zélvine Testar a signé.",
        "Quovane Testar",
    )
    assert written == f"Signé : {p} TESTAR. {p} Testar a signé."


def test_widening_over_capitals_form_of_typed_word(tmp_path: Path, db: Path) -> None:
    """G20 (AC3; QA's TEST-1019-05): typed « Zélvine Testar » covers
    « Zélvine TÉSTAR », the written word a Q3 form of the typed one."""
    source, output = _files(tmp_path, ".txt")
    source.write_text("Signé : Zélvine TÉSTAR.", encoding="utf-8")
    run = _process(
        source, output, db, [("Zélvine", PERSON)], [ENTER, "e"], ["Zélvine Testar"]
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    p = _pseudonym(db, "Zélvine Testar")
    assert _written(output) == f"Signé : {p}."


def test_delete_mapping_reaches_old_row_cut_after_an_apostrophe(
    tmp_path: Path, db: Path
) -> None:
    """G21 (AC4; QA's TEST-1019-05): v2.2.0 cut « ME » after the apostrophe of
    « O'MÉLORA »; deleting « Quovome O'Melora » erases that row with today's."""
    document = "QUOVÔME O'MÉLORA"
    left, erased, ids, old = _old_rows_erased(db, document, "Quovome O'Melora")
    _pre(
        any(unicodedata.combining(char) for char in old) and "O'" in old,
        "the old saved text, cut after the apostrophe",
    )
    assert left == [("Vorlan Mestrix", PERSON)]
    assert erased == sorted([ids[(document, PERSON)], ids[(old, PERSON)]])


def test_delete_mapping_reaches_old_titled_row(tmp_path: Path, db: Path) -> None:
    """G22 (AC4; QA's TEST-1019-06): v2.2.0 saved « MAÎTRE VORLAN MÉLORA »
    with the title kept (read with a separate mark) and the surname cut;
    deleting « Vorlan Melora » erases it with today's row of that text."""
    circumflex = unicodedata.lookup("COMBINING CIRCUMFLEX ACCENT")
    acute = unicodedata.lookup("COMBINING ACUTE ACCENT")
    document = "MAÎTRE VORLAN MÉLORA"
    left, erased, ids, old = _old_rows_erased(db, document, "Vorlan Melora")
    _pre(
        old == "MAI" + circumflex + "TRE VORLAN " + acute + "LORA",
        "the old saved text",
    )
    assert left == [("Vorlan Mestrix", PERSON)]
    assert erased == sorted([ids[(document, PERSON)], ids[(old, PERSON)]])


def test_delete_mapping_sets_aside_the_accent_beside_a_cut(
    tmp_path: Path, db: Path
) -> None:
    """G23 (AC4; QA's TEST-1019-06): v2.2.0 saved « ZÈLVINE MÈZOR » with
    « È » keeping its own accent beside the cut; deleting « Zelvine Mezor »
    erases it with today's row of that text."""
    grave = unicodedata.lookup("COMBINING GRAVE ACCENT")
    document = "ZÈLVINE MÈZOR"
    left, erased, ids, old = _old_rows_erased(db, document, "Zelvine Mezor")
    _pre(old == "ZE" + grave + "LVINE " + grave + "ZOR", "the old saved text")
    assert left == [("Vorlan Mestrix", PERSON)]
    assert erased == sorted([ids[(document, PERSON)], ids[(old, PERSON)]])


@pytest.mark.parametrize(
    "text,typed",
    [
        (
            "Signé : ZÉLVINE D'ÉMERINE. Zélvine d'Émerine a signé.",
            "ZELVINE D'EMERINE",
        ),
        (
            "Signé : ZÉLVINE L'ÉMERINE. Zélvine l'Émerine a signé.",
            "ZELVINE L'EMERINE",
        ),
    ],
    ids=["d_elision", "l_elision"],
)
def test_widening_over_an_elided_word(
    tmp_path: Path, db: Path, text: str, typed: str
) -> None:
    """G24 (AC3; v1.21): « d'Émerine » and « l'Émerine » are one word each (an
    apostrophe between letters), on both sides: a hyphen after one letter
    ends a word, an apostrophe does not. Covered whole."""
    written, p = _widened(tmp_path, db, text, typed)
    assert written == f"Signé : {p}. {p} a signé."


def test_delete_mapping_reaches_old_place_row_with_its_a(
    tmp_path: Path, db: Path
) -> None:
    """G25 (AC4; QA's TEST-1019-07): v2.2.0 saved the place « À VORLAN
    MÈLORA » with its « À » read with a separate mark and the surname cut;
    deleting « Vorlan Melora » erases it with today's « VORLAN MÈLORA », as
    10.20's replay puts the « à » in front of a place."""
    grave = unicodedata.lookup("COMBINING GRAVE ACCENT")
    old = _key_steps(nfd("À VORLAN MÈLORA"), LOCATION)
    _pre(old == "A" + grave + " VORLAN " + grave + "LORA", "the old saved text")
    _save(
        db,
        [
            ("VORLAN MÈLORA", LOCATION, "Zorbaville"),
            (old, LOCATION, "Quentreville"),
            ("Vorlan Mestrix", PERSON, "Quovane Testar"),
        ],
    )
    _pre(len(_names(db)) == 3, "three rows saved")
    ids = _ids(db)
    run = _delete(db, "Vorlan Melora")
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    assert _names(db) == [("Vorlan Mestrix", PERSON)]
    assert _erased_ids(db) == sorted(
        [ids[("VORLAN MÈLORA", LOCATION)], ids[(old, LOCATION)]]
    )


def test_widening_stops_at_a_number(tmp_path: Path, db: Path) -> None:
    """G26 (AC3; v1.21): what lies between two words is compared with the
    correction, so a number and a comma there are never covered:
    « Zélvine 14, Testar 12 » corrected as « ZELVINE TESTAR » keeps
    « 14, Testar 12 »."""
    source, output = _files(tmp_path, ".txt")
    source.write_text("Note : Zélvine 14, Testar 12.", encoding="utf-8")
    run = _process(
        source,
        output,
        db,
        [("Zélvine", PERSON)],
        [ENTER, "e"],
        ["ZELVINE TESTAR"],
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    p = _pseudonym(db, "ZELVINE TESTAR")
    assert _written(output) == f"Note : {p} 14, Testar 12."


@pytest.mark.parametrize(
    "text,tail",
    [
        (
            "Signé : ZÉLVINE TESTAR. Zélvine, Testar a signé.",
            "{p}, Testar a signé.",
        ),
        (
            "Signé : ZÉLVINE TESTAR. Merci Zélvine. Testar a signé.",
            "Merci {p}. Testar a signé.",
        ),
    ],
    ids=["comma", "sentence_full_stop"],
)
def test_widening_never_over_a_comma_or_a_full_stop_the_correction_lacks(
    tmp_path: Path, db: Path, text: str, tail: str
) -> None:
    """G27 (AC3; v1.21, QA's r3 probe « added_comma »): what lies between two
    words is what the correction has there; a comma or a sentence's full
    stop the correction does not hold is never covered, so the text keeps
    it and the next word."""
    written, p = _widened(tmp_path, db, text, "ZELVINE TESTAR")
    assert written == f"Signé : {p}. " + tail.format(p=p)


@pytest.mark.parametrize(
    "typed",
    ["Melora Quovome Testar", "Mélora Quovômé Testar"],
    ids=["typed_without_accents", "typed_with_accents"],
)
def test_delete_mapping_reaches_old_row_cut_between_two_marks(
    tmp_path: Path, db: Path, typed: str
) -> None:
    """G28 (AC4; v1.23): v2.2.0 cut « MÉ » from « MÉLORA » and the « MÉ » of
    « QUOVÔMÉ » between the « Ô »'s mark and its own, both marks left before
    the space; the row is erased typed with or without the accents."""
    acute = unicodedata.lookup("COMBINING ACUTE ACCENT")
    circumflex = unicodedata.lookup("COMBINING CIRCUMFLEX ACCENT")
    document = "MÉLORA QUOVÔMÉ TESTAR"
    old = _key_steps(nfd(document), PERSON)
    _pre(
        old == acute + "LORA QUOVO" + circumflex + acute + " TESTAR",
        "the old saved text",
    )
    left, erased, ids, _ = _old_rows_erased(db, document, typed)
    assert left == [("Vorlan Mestrix", PERSON)], left
    assert erased == sorted([ids[(document, PERSON)], ids[(old, PERSON)]])


def test_delete_mapping_reaches_old_row_of_a_title_run(
    tmp_path: Path, db: Path
) -> None:
    """G29 (AC4; QA's note on 46202a3): v2.2.0 cut both « MÉ » of « MÉMÉ
    TESTAR », leaving two marks + « TESTAR »; deleting « Meme Testar »
    erases it with today's row of that text."""
    acute = unicodedata.lookup("COMBINING ACUTE ACCENT")
    document = "MÉMÉ TESTAR"
    old = _key_steps(nfd(document), PERSON)
    _pre(old == acute + acute + " TESTAR", "the old saved text")
    left, erased, ids, _ = _old_rows_erased(db, document, "Meme Testar")
    assert left == [("Vorlan Mestrix", PERSON)], left
    assert erased == sorted([ids[(document, PERSON)], ids[(old, PERSON)]])


def test_delete_mapping_reaches_old_row_with_a_cedilla(
    tmp_path: Path, db: Path
) -> None:
    """G30 (AC4; QA's TEST-1019-08): v2.2.0 saved « MÉLORA VORÇANE » cut at
    « MÉ », the « Ç » keeping its cedilla as a separate mark; deleting
    « Melora Vorcane » erases it with today's row of that text."""
    acute = unicodedata.lookup("COMBINING ACUTE ACCENT")
    cedilla = unicodedata.lookup("COMBINING CEDILLA")
    document = "MÉLORA VORÇANE"
    old = _key_steps(nfd(document), PERSON)
    _pre(old == acute + "LORA VORC" + cedilla + "ANE", "the old saved text")
    left, erased, ids, _ = _old_rows_erased(db, document, "Melora Vorcane")
    assert left == [("Vorlan Mestrix", PERSON)], left
    assert erased == sorted([ids[(document, PERSON)], ids[(old, PERSON)]])


@pytest.mark.parametrize(
    "typed",
    [
        pytest.param("Melorine (Melavie) Testar", id="typed_without_accents"),
        pytest.param("Mélorine (Mélavie) Testar", id="typed_with_accents"),
    ],
)
def test_delete_mapping_reaches_old_row_cut_after_a_bracket(
    tmp_path: Path, db: Path, typed: str
) -> None:
    """T50 (AC4; the ninth Codex review's P0, QA's DEL-1019-10): v2.2.0 cut
    « MÉ » from « MÉLORINE » and from « MÉLAVIE » after the « ( », the old
    title pattern's own bound; the row is erased typed without the accents
    (expected failure on 5d110a4) and with them (a guard)."""
    acute = unicodedata.lookup("COMBINING ACUTE ACCENT")
    document = "MÉLORINE (MÉLAVIE) TESTAR"
    old = _key_steps(nfd(document), PERSON)
    _pre(old == acute + "LORINE (" + acute + "LAVIE) TESTAR", "the old saved text")
    left, erased, ids, _ = _old_rows_erased(db, document, typed)
    assert left == [("Vorlan Mestrix", PERSON)], left  # expected failure
    assert erased == sorted([ids[(document, PERSON)], ids[(old, PERSON)]])


@pytest.mark.parametrize("case", list(T51_CASES))
def test_delete_mapping_reaches_old_row_cut_after_any_bound(
    tmp_path: Path, db: Path, case: str
) -> None:
    """T51 (AC4; QA's DEL-1019-10): v2.2.0 cut « MÉ » after whatever the old
    title pattern's bound allows before a title word, « ( », « [ » or « / »,
    leaving the mark there (10.20's sign, read with the same bound); the row
    is erased with today's row of the name."""
    document, typed = T51_CASES[case]
    old = _key_steps(nfd(document), PERSON)
    _pre(
        any(unicodedata.combining(char) for char in old) and old != nfd(document),
        "the old saved text, cut",
    )
    left, erased, ids, _ = _old_rows_erased(db, document, typed)
    assert left == [("Vorlan Mestrix", PERSON)], left  # expected failure
    assert erased == sorted([ids[(document, PERSON)], ids[(old, PERSON)]])
