"""Story 10.20: an accent stored as a separate mark, in the readers, the
command line and the saved names (T1-T10, T12-T18, T20-T23, T26-T30, G2-G7).

Written by the BMAD dev agent (Claude). The readers are called directly; the
command-line paths drive the real ``gdpr-pseudo`` commands (``process``,
``batch --workers 1``, ``delete-mapping``, ``list-entities``,
``list-mappings``) through ``CliRunner``, on a database made by the real
``init`` command, as Story 10.15's harness (``test_cli_corrections_hand_added.py``).
Detection is replaced below the reader by a stand-in that finds each name
given in its composed and its decomposed spelling, in the text it receives;
T3 uses the real regex layer and T4 the real detector. The keys, the typed
texts and the yes/no answers come from lists (no to « Add a name the
detector missed », yes otherwise). The assertions read what a reader
returns, the written file, the stored rows, the ERASURE entries and the
lines shown.

Every text is written composed in this file; a decomposed form is made at
run time with ``unicodedata.normalize("NFD", ...)``, and the one mark with no
composed letter with ``unicodedata.lookup``: no combining mark is written in
this file. Preconditions call ``pytest.fail("PRE: ...")``, never ``assert``;
the first ``assert`` after them is the expected failure on the unfixed code
(``# expected failure``). Invented names and main-corpus words only. No Qt;
no spaCy model except T4.
"""

from __future__ import annotations

import csv
import dataclasses
import re
import sqlite3
import unicodedata
from collections.abc import Callable, Sequence
from contextlib import ExitStack, closing
from pathlib import Path
from typing import Any
from unittest.mock import patch

import openpyxl
import pytest
import readchar
from rich.prompt import Confirm, Prompt
from typer.testing import CliRunner

from gdpr_pseudonymizer.cli.main import app
from gdpr_pseudonymizer.core.custom_pseudonym import check_custom_pseudonym
from gdpr_pseudonymizer.core.document_processor import DocumentProcessor
from gdpr_pseudonymizer.data.database import open_database
from gdpr_pseudonymizer.data.models import Entity
from gdpr_pseudonymizer.data.repositories.mapping_repository import (
    SQLiteMappingRepository,
)
from gdpr_pseudonymizer.nlp.entity_detector import DetectedEntity
from gdpr_pseudonymizer.nlp.regex_matcher import RegexMatcher
from gdpr_pseudonymizer.utils.file_handler import read_file
from gdpr_pseudonymizer.utils.tabular_reader import (
    read_csv_structured,
    read_excel_structured,
)
from gdpr_pseudonymizer.validation.workflow import ADD_MISSED_NAME_QUESTION

PASSPHRASE = "zorbalia_pass_A_123!"
ENV = {"GDPR_PSEUDO_PASSPHRASE": PASSPHRASE}
SHEET = "Zorbalia"
PERSON = "PERSON"
LOCATION = "LOCATION"

ENTER = "\r"
SPACE = " "
MAX_KEYS = 60

ZO = "Zórbalia"
ZQ = "Zórbalia Quentrel"
VM = "Vorlan Mestrix"
XV = "Xantor Vellion"
XQ = "Xantor Quovix"
XVE = "Xantor Véllion"

SENT = "Le contrat avec Zórbalia Quentrel a été signé."
A_FIRST = "Vorlan Mestrix valide. Le contrat avec Zórbalia."
A_SECOND = "Signé : Zórbalia."
A_TEXT = f"{A_FIRST} {A_SECOND}"
A_SENTENCES = ("Vorlan Mestrix valide.", "Le contrat avec Zórbalia.", A_SECOND)
E_TEXT = "Zórbalia Quentrel a signé. Signé : Zórbalia Quentrel."
R_TEXT = "Le contrat avec Zórbalia Quentrel a signé."
R_REUSED = "Le contrat avec Xantor Vellion a signé."
MME_TEXT = "Vorlan Mestrix valide. Signé : Mme Zórbalia."
G7_REST = " a signé avec quentrel quovan et Zorbal Quentrix."

NOT_FOUND = "not found in document"
STILL_ONE = "Still to review: 1 (left in plain text)"
CONFIRM_PROMPT = "Type 'yes'"
# Lionel's Qc line (2026-10-21, "A: Use it (Rec.)"), N = 2, byte for byte.
SEVERAL_LINE = (
    "You are about to delete the following 2 saved pseudonyms "
    "(all of this name, whatever its type):"
)

XFAIL = "Story 10.20: reproduces {} on the unfixed code; fixed in Task 2"

_ANSI = re.compile(r"\x1b\[[0-9;?]*[A-Za-z]")


def nfc(text: str) -> str:
    return unicodedata.normalize("NFC", text)


def nfd(text: str) -> str:
    return unicodedata.normalize("NFD", text)


FORMS: dict[str, Callable[[str], str]] = {"nfc": nfc, "nfd": nfd}

# « Zox » + U+0301 + « rbalia »: a mark with no composed letter (AC4)
ACUTE = unicodedata.lookup("COMBINING ACUTE ACCENT")
ZOX = "Zox" + ACUTE + "rbalia"


def _pre(condition: bool, what: str) -> None:
    """A precondition: a red failure on every leg, never an AssertionError."""
    if not condition:
        pytest.fail(f"PRE: {what}")


def _has_mark(text: str) -> bool:
    """Whether ``text`` holds a combining mark (AC4's rule)."""
    return any(unicodedata.combining(char) for char in text)


def _collapsed(text: str) -> str:
    """``text`` with every whitespace run made one space (Rich wraps lines)."""
    return " ".join(text.split())


@pytest.fixture(autouse=True)
def _settings(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    """No user or project settings file is read or written."""
    monkeypatch.setenv("HOME", str(tmp_path))
    monkeypatch.setenv("USERPROFILE", str(tmp_path))
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(
        "gdpr_pseudonymizer.gui.config._config_path",
        lambda: tmp_path / ".gdpr-pseudo.yaml",
    )


@pytest.fixture
def db(tmp_path: Path) -> Path:
    """A new database, made by the real ``init`` command."""
    path = tmp_path / "mappings.db"
    result = CliRunner().invoke(
        app, ["init", "--db", str(path)], env=ENV, catch_exceptions=False
    )
    _pre(result.exit_code == 0 and path.is_file(), "init made the database")
    return path


# ---------------------------------------------------------------------------
# Harness
# ---------------------------------------------------------------------------


class EitherForm:
    """Stand-in detector: every occurrence of each (name, type) given, in its
    composed and its decomposed spelling, in the text it receives; an
    occurrence inside a longer one found is dropped."""

    def __init__(self, names: Sequence[tuple[str, str]]) -> None:
        self.names = list(names)

    def detect_entities(self, text: str) -> list[DetectedEntity]:
        found: list[DetectedEntity] = []
        for name, entity_type in self.names:
            for form in dict.fromkeys((nfc(name), nfd(name))):
                start = text.find(form)
                while start != -1:
                    found.append(
                        DetectedEntity(
                            text=form,
                            entity_type=entity_type,
                            start_pos=start,
                            end_pos=start + len(form),
                            confidence=0.9,
                            source="spacy",
                        )
                    )
                    start = text.find(form, start + 1)
        return [
            entity
            for entity in found
            if not any(
                other.start_pos <= entity.start_pos
                and entity.end_pos <= other.end_pos
                and other.end_pos - other.start_pos > entity.end_pos - entity.start_pos
                for other in found
            )
        ]


class RegexLayer:
    """The real regex layer of the detector (the one that cuts « Mme Zo »)."""

    def __init__(self) -> None:
        self.matcher = RegexMatcher()
        self.matcher.load_patterns()

    def detect_entities(self, text: str) -> list[DetectedEntity]:
        return self.matcher.match_entities(text)


class Keys:
    """``readchar.readkey``: the keys in order, then Space."""

    def __init__(self, keys: Sequence[str]) -> None:
        self.keys = list(keys)
        self.read = 0

    def __call__(self) -> str:
        self.read += 1
        if self.read > MAX_KEYS:
            raise RuntimeError("more than 60 keys read")
        return self.keys.pop(0) if self.keys else SPACE


class Typed:
    """``Prompt.ask``: the typed texts in order. The prompt is printed, so the
    output shows where it was asked."""

    def __init__(self, texts: Sequence[str]) -> None:
        self.texts = list(texts)

    def __call__(self, prompt: Any, *args: Any, **kwargs: Any) -> str:
        print(f"[asked] {prompt}")
        if not self.texts:
            raise RuntimeError("no typed text left")
        return self.texts.pop(0)


class Answers:
    """``Confirm.ask``: no to Story 10.18's question, yes otherwise."""

    def __call__(self, prompt: Any, *args: Any, **kwargs: Any) -> bool:
        return str(prompt) != ADD_MISSED_NAME_QUESTION


@dataclasses.dataclass
class Run:
    exit_code: int
    shown: str

    @property
    def flat(self) -> str:
        return _collapsed(self.shown)


def _invoke(
    args: Sequence[str],
    names: Sequence[tuple[str, str]] = (),
    keys: Sequence[str] = (),
    typed: Sequence[str] = (),
    detect: str = "stand-in",
) -> Run:
    """One real command run. ``detect``: "stand-in" (``EitherForm``),
    "regex" (``RegexLayer``) or "real" (the detector itself)."""
    with ExitStack() as stack:
        if detect != "real":
            detector: Any = RegexLayer() if detect == "regex" else EitherForm(names)
            stack.enter_context(
                patch.object(
                    DocumentProcessor, "_get_detector", new=lambda self: detector
                )
            )
            stack.enter_context(
                patch.object(
                    DocumentProcessor,
                    "_detect_and_filter_entities",
                    new=lambda self, text, entity_type_filter=None: (
                        detector.detect_entities(text)
                    ),
                )
            )
        stack.enter_context(patch.object(readchar, "readkey", new=Keys(keys)))
        stack.enter_context(patch.object(Prompt, "ask", new=Typed(typed)))
        stack.enter_context(patch.object(Confirm, "ask", new=Answers()))
        result = CliRunner().invoke(app, list(args), env=ENV, catch_exceptions=False)
    return Run(result.exit_code, _ANSI.sub("", result.output))


def _process(
    source: Path,
    output: Path,
    db: Path,
    names: Sequence[tuple[str, str]] = (),
    keys: Sequence[str] = (ENTER,),
    typed: Sequence[str] = (),
    detect: str = "stand-in",
) -> Run:
    """One real ``gdpr-pseudo process`` run."""
    return _invoke(
        ["process", str(source), "--db", str(db), "-o", str(output)],
        names,
        keys,
        typed,
        detect,
    )


def _batch(folder: Path, out: Path, db: Path, names: Sequence[tuple[str, str]]) -> Run:
    """One real ``gdpr-pseudo batch --workers 1`` run."""
    return _invoke(
        [
            "batch",
            str(folder),
            "--db",
            str(db),
            "--output",
            str(out),
            "--workers",
            "1",
        ],
        names,
        (ENTER,),
    )


def _delete(db: Path, name: str) -> Run:
    """One real ``gdpr-pseudo delete-mapping <name>``, « yes » typed."""
    return _invoke(["delete-mapping", name, "--db", str(db)], typed=["yes"])


def _write_sheet(path: Path, rows: list[list[str]]) -> None:
    if path.suffix == ".xlsx":
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = SHEET
        for row in rows:
            sheet.append(row)
        workbook.save(path)
        workbook.close()
    else:
        with path.open("w", encoding="utf-8", newline="") as handle:
            csv.writer(handle).writerows(rows)


def _write_doc(path: Path, text: str) -> None:
    """A document holding ``text``; a sheet holds it in A2, « ok » in B2,
    under the header « Nom », « Note »."""
    if path.suffix in (".xlsx", ".csv"):
        _write_sheet(path, [["Nom", "Note"], [text, "ok"]])
    else:
        path.write_text(text, encoding="utf-8")


def _written(path: Path) -> str:
    """The written text; a sheet's cells, one per line."""
    _pre(path.is_file(), f"{path.name} is written")
    if path.suffix == ".xlsx":
        workbook = openpyxl.load_workbook(path)
        try:
            return "\n".join(
                str(value)
                for sheet in workbook.worksheets
                for row in sheet.iter_rows(values_only=True)
                for value in row
                if value is not None
            )
        finally:
            workbook.close()
    if path.suffix == ".csv":
        with path.open(encoding="utf-8", newline="") as handle:
            return "\n".join(cell for row in csv.reader(handle) for cell in row)
    return path.read_text(encoding="utf-8")


def _save(db: Path, rows: Sequence[tuple[str, str, str]]) -> None:
    """Rows saved in order, each with the exact name given, as an earlier
    version saved them: (full_name, type, pseudonym)."""
    with open_database(str(db), PASSPHRASE) as session:
        repo = SQLiteMappingRepository(session)
        for full_name, entity_type, pseudonym in rows:
            repo.save(
                Entity(
                    entity_type=entity_type,
                    full_name=full_name,
                    pseudonym_full=pseudonym,
                    theme="neutral",
                )
            )


def _rows(db: Path) -> list[tuple[str, str, str]]:
    """Every stored row: (full_name, entity_type, pseudonym_full)."""
    with open_database(str(db), PASSPHRASE) as session:
        return [
            (row.full_name, row.entity_type, row.pseudonym_full)
            for row in SQLiteMappingRepository(session).find_all()
        ]


def _names(db: Path) -> list[tuple[str, str]]:
    """Every stored (full_name, entity_type), sorted."""
    return sorted((name, entity_type) for name, entity_type, _ in _rows(db))


def _erasures(db: Path) -> int:
    """The ERASURE entries of the operations table."""
    with closing(sqlite3.connect(str(db))) as connection:
        (count,) = connection.execute(
            "SELECT COUNT(*) FROM operations WHERE operation_type = 'ERASURE'"
        ).fetchone()
    return int(count)


# A PDF whose accents are drawn apart (no font file): the standard Helvetica
# font, WinAnsi encoding, each accented letter drawn as its plain letter then
# a spacing accent, which the font's /ToUnicode map gives as a combining mark.
_SPACING_ACCENT = {
    unicodedata.lookup("COMBINING ACUTE ACCENT"): 0xB4,
    unicodedata.lookup("COMBINING GRAVE ACCENT"): 0x60,
    unicodedata.lookup("COMBINING CIRCUMFLEX ACCENT"): 0x88,
    unicodedata.lookup("COMBINING DIAERESIS"): 0xA8,
    unicodedata.lookup("COMBINING CEDILLA"): 0xB8,
}


def _stream(data: bytes) -> bytes:
    return f"<< /Length {len(data)} >>\nstream\n".encode() + data + b"\nendstream"


def _write_pdf(path: Path, text: str) -> None:
    codes = bytearray()
    for char in nfd(text):
        code = _SPACING_ACCENT.get(char, ord(char))
        _pre(code < 0x100, f"{char!r} drawn with WinAnsi")
        codes.append(code)
    content = b"BT /F1 12 Tf 40 700 Td <" + codes.hex().upper().encode() + b"> Tj ET"
    to_unicode = (
        b"/CIDInit /ProcSet findresource begin\n12 dict begin\nbegincmap\n"
        b"/CIDSystemInfo << /Registry (Adobe) /Ordering (UCS) /Supplement 0 >> def\n"
        b"/CMapName /Adobe-Identity-UCS def\n/CMapType 2 def\n"
        b"1 begincodespacerange\n<00> <FF>\nendcodespacerange\n"
        + f"{len(_SPACING_ACCENT)} beginbfchar\n".encode()
        + b"".join(
            f"<{code:02X}> <{ord(mark):04X}>\n".encode()
            for mark, code in _SPACING_ACCENT.items()
        )
        + b"endbfchar\nendcmap\nCMapName currentdict /CMap defineresource pop\n"
        b"end\nend"
    )
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica "
        b"/Encoding /WinAnsiEncoding /ToUnicode 6 0 R >>",
        _stream(content),
        _stream(to_unicode),
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets: list[int] = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref}\n%%EOF\n"
    ).encode()
    path.write_bytes(bytes(out))


def _read_back(tmp_path: Path, kind: str, text: str) -> str:
    """``text`` written as a file of ``kind``, then read by the product's
    reader (the structured readers' cells joined, one per line)."""
    ext = "." + kind.split("_")[0]
    path = tmp_path / f"doc{ext}"
    if kind == "docx":
        import docx

        document = docx.Document()
        document.add_paragraph(text)
        document.save(str(path))
    elif kind == "pdf":
        _write_pdf(path, text)
    else:
        _write_doc(path, text)
    if kind == "xlsx_cells":
        return "\n".join(cell.value for cell in read_excel_structured(str(path)).cells)
    if kind == "csv_cells":
        return "\n".join(cell.value for cell in read_csv_structured(str(path)).cells)
    return read_file(str(path))


# ---------------------------------------------------------------------------
# AC1, AC5: every reader returns composed text; written files are composed
# ---------------------------------------------------------------------------


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason=XFAIL.format("the reader returning the stored marks"),
)
@pytest.mark.parametrize(
    "kind", ["txt", "md", "docx", "pdf", "xlsx", "csv", "xlsx_cells", "csv_cells"]
)
def test_reader_returns_composed(tmp_path: Path, kind: str) -> None:
    """T1 (AC1): each reader returns the sentence composed, from a file that
    stores it decomposed (a PDF whose accents are drawn apart)."""
    text = _read_back(tmp_path, kind, nfd(SENT))
    _pre(SENT in nfc(text), f"the reader returned the sentence: {text!r}")
    assert not _has_mark(text), repr(text)  # expected failure
    assert SENT in text


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason=XFAIL.format("a written file keeping the stored marks"),
)
@pytest.mark.parametrize("ext", [".txt", ".xlsx", ".csv"])
def test_written_file_composed(tmp_path: Path, db: Path, ext: str) -> None:
    """T2 (AC5): a document stored decomposed, nothing detected, is written
    composed."""
    source, output = tmp_path / f"in{ext}", tmp_path / f"out{ext}"
    _write_doc(source, nfd(SENT))
    run = _process(source, output, db)
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    assert not _has_mark(written), repr(written)  # expected failure
    assert SENT in written


# ---------------------------------------------------------------------------
# Path 1: « Mme Zo »
# ---------------------------------------------------------------------------


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason=XFAIL.format(
        "« Mme Zo » (the regex layer cuts the title's name at the mark)"
    ),
)
def test_mme_zo_regex_layer(tmp_path: Path, db: Path) -> None:
    """T3 (path 1): with the real regex layer, a decomposed « Mme Zórbalia »
    is replaced whole: the end of the name is not left in the written file."""
    source, output = tmp_path / "in.txt", tmp_path / "out.txt"
    source.write_text(nfd(MME_TEXT), encoding="utf-8")
    run = _process(source, output, db, detect="regex")
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    _pre("Reviewing PERSON entities" in run.flat, "a PERSON reviewed")
    assert "rbalia" not in nfc(written), repr(written)  # expected failure
    assert not _has_mark(written)


@pytest.mark.spacy
@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason=XFAIL.format("« Mme Zo » with the real detector"),
)
def test_mme_zo_real_detector(tmp_path: Path, db: Path) -> None:
    """T4 (path 1): the same with the real detector."""
    source, output = tmp_path / "in.txt", tmp_path / "out.txt"
    source.write_text(nfd(MME_TEXT), encoding="utf-8")
    run = _process(source, output, db, detect="real")
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    _pre("Reviewing PERSON entities" in run.flat, "a PERSON reviewed")
    assert "rbalia" not in nfc(written), repr(written)  # expected failure
    assert not _has_mark(written)


# ---------------------------------------------------------------------------
# Path 2: [A]
# ---------------------------------------------------------------------------


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason=XFAIL.format("[A] « not found » for a name typed composed"),
)
@pytest.mark.parametrize("ext", [".txt", ".xlsx"])
def test_add_composed_in_decomposed_document(
    tmp_path: Path, db: Path, ext: str
) -> None:
    """T5 (path 2): « Zórbalia » typed composed at [A] in a decomposed
    document: no « not found », every place replaced, one row."""
    source, output = tmp_path / f"in{ext}", tmp_path / f"out{ext}"
    if ext == ".xlsx":
        _write_sheet(source, [[nfd(sentence)] for sentence in A_SENTENCES])
    else:
        source.write_text(nfd(A_TEXT), encoding="utf-8")
    run = _process(
        source, output, db, [(VM, PERSON)], [ENTER, "a", SPACE], [ZO, PERSON]
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    assert NOT_FOUND not in run.flat, run.shown  # expected failure
    assert "rbalia" not in nfc(written), repr(written)
    assert _names(db) == sorted([(VM, PERSON), (ZO, PERSON)])


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason=XFAIL.format("[A] « not found » for a name typed decomposed"),
)
def test_add_decomposed_in_composed_document(tmp_path: Path, db: Path) -> None:
    """T6 (path 2): « Zórbalia » typed decomposed at [A] in a composed
    document: no « not found », one row, composed."""
    source, output = tmp_path / "in.txt", tmp_path / "out.txt"
    source.write_text(A_TEXT, encoding="utf-8")
    run = _process(
        source, output, db, [(VM, PERSON)], [ENTER, "a", SPACE], [nfd(ZO), PERSON]
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    assert NOT_FOUND not in run.flat, run.shown  # expected failure
    assert _names(db) == sorted([(VM, PERSON), (ZO, PERSON)])


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason=XFAIL.format("[A] replacing the places of one form only"),
)
@pytest.mark.parametrize("typed_form", ["nfc", "nfd"])
def test_add_in_document_with_both_forms(
    tmp_path: Path, db: Path, typed_form: str
) -> None:
    """T7 (path 2): the first place composed, the second decomposed; [A]
    typed in either form replaces both; one row, composed."""
    source, output = tmp_path / "in.txt", tmp_path / "out.txt"
    source.write_text(f"{nfc(A_FIRST)} {nfd(A_SECOND)}", encoding="utf-8")
    typed = FORMS[typed_form](ZO)
    run = _process(
        source, output, db, [(VM, PERSON)], [ENTER, "a", SPACE], [typed, PERSON]
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    _pre("Added PERSON entity" in run.flat, f"added: {run.shown!r}")
    written = _written(output)
    assert "rbalia" not in nfc(written), repr(written)  # expected failure
    assert [row for row in _names(db) if row[0] != VM] == [(ZO, PERSON)]


# ---------------------------------------------------------------------------
# Path 3: [E]
# ---------------------------------------------------------------------------


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason=XFAIL.format("a correction typed in the other form leaving « Quentrel »"),
)
@pytest.mark.parametrize(
    "doc_form, typed_form", [("nfd", "nfc"), ("nfc", "nfd")], ids=["nfd-nfc", "nfc-nfd"]
)
def test_correction_across_forms(
    tmp_path: Path, db: Path, doc_form: str, typed_form: str
) -> None:
    """T8 (path 3): « Zórbalia » found twice, corrected at [E] into
    « Zórbalia Quentrel » typed in the other form: no part of it in clear,
    one row, composed."""
    source, output = tmp_path / "in.txt", tmp_path / "out.txt"
    source.write_text(FORMS[doc_form](E_TEXT), encoding="utf-8")
    run = _process(
        source, output, db, [(ZO, PERSON)], [ENTER, "e"], [FORMS[typed_form](ZQ)]
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    _pre("2 occurrences, 1 unique" in run.flat, f"one group of 2: {run.shown!r}")
    written = _written(output)
    assert "Quentrel" not in written, repr(written)  # expected failure
    assert _names(db) == [(ZQ, PERSON)]


# ---------------------------------------------------------------------------
# Path 4: the review's grouping and 10.13's count
# ---------------------------------------------------------------------------


def _mixed() -> str:
    """« Zórbalia Quentrel a signé. » composed, then « Signé : Zórbalia
    Quentrel. » decomposed."""
    return f"{nfc('Zórbalia Quentrel a signé.')} {nfd('Signé : Zórbalia Quentrel.')}"


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason=XFAIL.format("two names in the review for one name in two forms"),
)
def test_both_forms_are_one_name(tmp_path: Path, db: Path) -> None:
    """T9 (path 4): the name composed at one place and decomposed at another
    is one name in the review, one pseudonym in the file, one row."""
    source, output = tmp_path / "in.txt", tmp_path / "out.txt"
    source.write_text(_mixed(), encoding="utf-8")
    run = _process(source, output, db, [(ZQ, PERSON)], [ENTER, SPACE])
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    assert "2 occurrences, 1 unique" in run.flat, run.shown  # expected failure
    rows = _rows(db)
    assert [row[:2] for row in rows] == [(ZQ, PERSON)]
    pseudonym = rows[0][2]
    assert _written(output) == f"{pseudonym} a signé. Signé : {pseudonym}."


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason=XFAIL.format("« Still to review: 2 » for one name in two forms"),
)
def test_still_to_review_counts_it_once(tmp_path: Path, db: Path) -> None:
    """T10 (path 4, 10.13's count): the name in two forms, left to review,
    is counted once."""
    source, output = tmp_path / "in.txt", tmp_path / "out.txt"
    source.write_text(f"{_mixed()} Vorlan Mestrix valide.", encoding="utf-8")
    run = _process(
        source, output, db, [(ZQ, PERSON), (VM, PERSON)], [ENTER, "n", "n", SPACE]
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    _pre("Still to review" in run.flat, f"names left to review: {run.shown!r}")
    assert STILL_ONE in run.flat, run.shown  # expected failure


# ---------------------------------------------------------------------------
# Path 7: a saved name reused in a later document
# ---------------------------------------------------------------------------

_ORDERS = pytest.mark.parametrize(
    "first, second", [("nfc", "nfd"), ("nfd", "nfc")], ids=["nfc-nfd", "nfd-nfc"]
)


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason=XFAIL.format("a second pseudonym for the name in the other form"),
)
@_ORDERS
def test_reuse_process_twice(tmp_path: Path, db: Path, first: str, second: str) -> None:
    """T12 (path 7): the name saved from one form and met in the other in
    the next ``process``: the same pseudonym, the same written text, one
    row."""
    source_a, output_a = tmp_path / "a.txt", tmp_path / "a_out.txt"
    source_b, output_b = tmp_path / "b.txt", tmp_path / "b_out.txt"
    source_a.write_text(FORMS[first](R_TEXT), encoding="utf-8")
    source_b.write_text(FORMS[second](R_TEXT), encoding="utf-8")
    run_a = _process(source_a, output_a, db, [(ZQ, PERSON)], [ENTER, SPACE])
    _pre(run_a.exit_code == 0 and len(_rows(db)) == 1, "the first run stored one row")
    run_b = _process(source_b, output_b, db, [(ZQ, PERSON)], [ENTER, SPACE])
    _pre(run_b.exit_code == 0, f"exit code 0: {run_b.shown!r}")
    written_b = _written(output_b)
    assert written_b == _written(output_a), repr(written_b)  # expected failure
    assert _names(db) == [(ZQ, PERSON)]


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason=XFAIL.format("a second pseudonym in batch for the name in the other form"),
)
@_ORDERS
def test_reuse_batch_one_at_a_time(
    tmp_path: Path, db: Path, first: str, second: str
) -> None:
    """T13 (path 7): the same with ``batch --workers 1``."""
    folder, out = tmp_path / "in", tmp_path / "out"
    folder.mkdir()
    (folder / "a.txt").write_text(FORMS[first](R_TEXT), encoding="utf-8")
    (folder / "b.txt").write_text(FORMS[second](R_TEXT), encoding="utf-8")
    run = _batch(folder, out, db, [(ZQ, PERSON)])
    output_a, output_b = out / "a_pseudonymized.txt", out / "b_pseudonymized.txt"
    _pre(
        run.exit_code == 0 and output_a.is_file() and output_b.is_file(),
        f"two files written: {run.shown!r}",
    )
    written_b = _written(output_b)
    assert written_b == _written(output_a), repr(written_b)  # expected failure
    assert _names(db) == [(ZQ, PERSON)]


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason=XFAIL.format("a row saved decomposed not reused"),
)
def test_old_decomposed_row_reused(tmp_path: Path, db: Path) -> None:
    """T14 (path 7, AC3): a row saved decomposed, as the unfixed code saves
    it, gives its pseudonym to the composed name."""
    _save(db, [(nfd(ZQ), PERSON, XV)])
    _pre(_rows(db) == [(nfd(ZQ), PERSON, XV)], "the decomposed row is stored")
    source, output = tmp_path / "in.txt", tmp_path / "out.txt"
    source.write_text(R_TEXT, encoding="utf-8")
    run = _process(source, output, db, [(ZQ, PERSON)], [ENTER, SPACE])
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    assert written == R_REUSED, repr(written)  # expected failure
    assert len(_rows(db)) == 1


# ---------------------------------------------------------------------------
# Path 8: delete-mapping, list-entities, list-mappings
# ---------------------------------------------------------------------------


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason=XFAIL.format("« Entity Not Found » for a row saved decomposed"),
)
def test_delete_decomposed_row_by_composed_name(tmp_path: Path, db: Path) -> None:
    """T15 (path 8): ``delete-mapping`` with the composed name erases the row
    saved decomposed, with its erasure entry."""
    _save(db, [(nfd(ZQ), PERSON, XV)])
    _pre(len(_rows(db)) == 1, "one row stored")
    run = _delete(db, ZQ)
    assert run.exit_code == 0, run.shown  # expected failure
    assert _rows(db) == []
    assert _erasures(db) == 1


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason=XFAIL.format("only the composed row listed and erased"),
)
def test_delete_both_forms(tmp_path: Path, db: Path) -> None:
    """T16 (path 8, Open Point (b)): a name saved both ways: both rows listed
    before the confirmation, both erased, each with its entry."""
    _save(db, [(ZQ, PERSON, XV), (nfd(ZQ), PERSON, XQ)])
    _pre(len(_rows(db)) == 2, "two rows stored")
    run = _delete(db, ZQ)
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    _pre(CONFIRM_PROMPT in run.shown, "the confirmation is asked")
    before = _collapsed(run.shown.split(CONFIRM_PROMPT, 1)[0])
    assert XQ in before, before  # expected failure
    assert _rows(db) == []
    assert _erasures(db) == 2


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason=XFAIL.format("« No entities found » for the name in the other form"),
)
@pytest.mark.parametrize(
    "saved, searched", [("nfd", "nfc"), ("nfc", "nfd")], ids=["nfd-nfc", "nfc-nfd"]
)
def test_list_entities_search(
    tmp_path: Path, db: Path, saved: str, searched: str
) -> None:
    """T17 (path 8, AC2): ``list-entities --search`` in the other form lists
    the row."""
    _save(db, [(FORMS[saved](ZQ), PERSON, XV)])
    run = _invoke(["list-entities", "--db", str(db), "--search", FORMS[searched](ZO)])
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    assert XV in run.flat, run.shown  # expected failure


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason=XFAIL.format("« No mappings found » for a row saved decomposed"),
)
def test_list_mappings_search(tmp_path: Path, db: Path) -> None:
    """T18 (AC2): ``list-mappings --search`` composed lists the row saved
    decomposed."""
    _save(db, [(nfd(ZQ), PERSON, XV)])
    run = _invoke(["list-mappings", "--db", str(db), "--search", ZO])
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    assert XV in run.flat, run.shown  # expected failure


# ---------------------------------------------------------------------------
# AC4: a word keeps its marks, also a mark with no composed letter
# ---------------------------------------------------------------------------


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason=XFAIL.format("the custom pseudonym check cutting a word at its mark"),
)
def test_custom_check_keeps_the_mark_inside_the_word() -> None:
    """T20 (AC4): « Zox » + U+0301 + « rbalia » is one word for the custom
    pseudonym check, so « Xantor Zoxrbalia » is refused."""
    document = f"Quentrel {ZOX} a signé le contrat."
    refusal = check_custom_pseudonym(
        [], f"Quentrel {ZOX}", PERSON, "Xantor Zoxrbalia", document, {}
    )
    assert refusal is not None  # expected failure


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason=XFAIL.format("[A] placing « Zox » inside a word with a mark"),
)
def test_add_does_not_place_inside_a_word_with_a_mark(tmp_path: Path, db: Path) -> None:
    """T21 (AC4): [A] « Zox » is not placed inside « Zox » + U+0301 +
    « rbalia »."""
    source, output = tmp_path / "in.txt", tmp_path / "out.txt"
    source.write_text(f"Vorlan Mestrix valide. {ZOX} a signé. Zox valide.", "utf-8")
    run = _process(
        source, output, db, [(VM, PERSON)], [ENTER, "a", SPACE], ["Zox", PERSON]
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    _pre("Zox valide" not in written, f"« Zox valide » replaced: {written!r}")
    assert ZOX in written, repr(written)  # expected failure


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason=XFAIL.format("a name with a mark not redacted from the error text"),
)
def test_log_redaction_keeps_the_mark_inside_the_name() -> None:
    """T22 (AC4): the error text's name-shaped run takes the mark, so the
    whole name is redacted."""
    message = DocumentProcessor._sanitize_error_message(
        ValueError(f"failed near {ZOX} Quentrel today")
    )
    _pre(message.startswith("ValueError: "), message)
    assert "rbalia" not in message, message  # expected failure
    assert "Quentrel" not in message, message


# ---------------------------------------------------------------------------
# AC2: a custom pseudonym typed decomposed
# ---------------------------------------------------------------------------


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason=XFAIL.format("a custom pseudonym typed decomposed stored decomposed"),
)
def test_cli_custom_pseudonym_typed_decomposed(tmp_path: Path, db: Path) -> None:
    """T23 (AC2): [C] « Xantor Véllion » typed decomposed is stored and
    written composed."""
    source, output = tmp_path / "in.txt", tmp_path / "out.txt"
    source.write_text("Vorlan Mestrix valide.", encoding="utf-8")
    run = _process(source, output, db, [(VM, PERSON)], [ENTER, "c"], [nfd(XVE)])
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    _pre(nfc(written) == f"{XVE} valide.", f"the custom pseudonym written: {written!r}")
    rows = _rows(db)
    _pre(len(rows) == 1, f"one row: {rows!r}")
    assert rows[0][2] == XVE, repr(rows[0][2])  # expected failure
    assert written == nfc(written)


# ---------------------------------------------------------------------------
# The PO's S2 and M1, Lionel's Qa, Qb, Qc
# ---------------------------------------------------------------------------


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason=XFAIL.format("« Me<pseudonym> » for a decomposed « Mélora »"),
)
def test_fresh_decomposed_title_like_name(tmp_path: Path, db: Path) -> None:
    """T26 (the PO's S2): a decomposed « Mélora Quentrel » on a fresh
    database is replaced whole and saved composed."""
    source, output = tmp_path / "in.txt", tmp_path / "out.txt"
    source.write_text(nfd("Le contrat avec Mélora Quentrel a signé."), "utf-8")
    run = _process(source, output, db, [("Mélora Quentrel", PERSON)], [ENTER, SPACE])
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    rows = _rows(db)
    _pre(len(rows) == 1, f"one row: {rows!r}")
    expected = f"Le contrat avec {rows[0][2]} a signé."
    assert nfc(written) == expected, repr(written)  # expected failure
    assert rows[0][:2] == ("Mélora Quentrel", PERSON)


@dataclasses.dataclass(frozen=True)
class OldRow:
    """A row an earlier version saved from a decomposed document, and the
    composed document that later meets the name."""

    stored: str
    entity_type: str
    pseudonym: str
    document: str
    detected: str
    expected: str
    bare: str


def _old_titled(case: str) -> OldRow:
    """The PO's M1: « Maître » or « à » kept in front, decomposed."""
    if case == "maitre_vorlan":
        return OldRow(
            nfd("Maître Vorlan Mestrix"),
            PERSON,
            XV,
            "Le contrat avec Maître Vorlan Mestrix a signé.",
            "Maître Vorlan Mestrix",
            "Le contrat avec Maître Xantor Vellion a signé.",
            VM,
        )
    if case == "maitre_eloria":
        return OldRow(
            nfd("Maître Éloria Quentrel"),
            PERSON,
            XV,
            "Le contrat avec Maître Éloria Quentrel a signé.",
            "Maître Éloria Quentrel",
            "Le contrat avec Maître Xantor Vellion a signé.",
            "Éloria Quentrel",
        )
    return OldRow(
        nfd("à Zórbaville"),
        LOCATION,
        "Xantorville",
        "Le contrat a été signé à Zórbaville.",
        "Zórbaville",
        "Le contrat a été signé à Xantorville.",
        "Zórbaville",
    )


def _old_cut(case: str) -> OldRow:
    """Lionel's Qa: a row that lost letters to the title pattern."""
    name = "Mélora Quentrel" if case == "melora" else "Zorbôme Quentrel"
    stored = ACUTE + "lora Quentrel" if case == "melora" else nfd("Zorbô") + "Quentrel"
    return OldRow(
        stored,
        PERSON,
        XV,
        f"Le contrat avec {name} a signé.",
        name,
        R_REUSED,
        name,
    )


TITLED = ["maitre_vorlan", "maitre_eloria", "a_zorbaville"]
CUT = ["melora", "zorbome"]


def _reuse(tmp_path: Path, db: Path, old: OldRow) -> str:
    """The old row saved, then the composed document processed: the written
    text (after the preconditions)."""
    _save(db, [(old.stored, old.entity_type, old.pseudonym)])
    _pre(
        _rows(db) == [(old.stored, old.entity_type, old.pseudonym)],
        "the old row is stored",
    )
    source, output = tmp_path / "in.txt", tmp_path / "out.txt"
    source.write_text(old.document, encoding="utf-8")
    run = _process(
        source, output, db, [(old.detected, old.entity_type)], [ENTER, SPACE]
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    return _written(output)


def _erase(db: Path, old: OldRow) -> Run:
    """The old row saved, then ``delete-mapping`` with the bare name."""
    _save(db, [(old.stored, old.entity_type, old.pseudonym)])
    _pre(len(_rows(db)) == 1, "the old row is stored")
    return _delete(db, old.bare)


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason=XFAIL.format("a new pseudonym beside a row saved with « Maître » or « à »"),
)
@pytest.mark.parametrize("case", TITLED)
def test_old_titled_row_reused(tmp_path: Path, db: Path, case: str) -> None:
    """T27 (the PO's M1, AC3): a row saved with a decomposed « Maître » or
    « à » in its key gives its pseudonym to the composed name."""
    old = _old_titled(case)
    written = _reuse(tmp_path, db, old)
    assert written == old.expected, repr(written)  # expected failure
    assert len(_rows(db)) == 1


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason=XFAIL.format(
        "« Entity Not Found » for a row saved with « Maître » or « à »"
    ),
)
@pytest.mark.parametrize("case", TITLED)
def test_old_titled_row_erased_by_name(tmp_path: Path, db: Path, case: str) -> None:
    """T27 (the PO's M1, AC3): ``delete-mapping`` with the bare composed name
    erases that row, with its entry."""
    run = _erase(db, _old_titled(case))
    assert run.exit_code == 0, run.shown  # expected failure
    assert _rows(db) == []
    assert _erasures(db) == 1


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason=XFAIL.format("a new pseudonym beside a row that lost letters"),
)
@pytest.mark.parametrize("case", CUT)
def test_row_that_lost_letters_reused(tmp_path: Path, db: Path, case: str) -> None:
    """T28 (Lionel's Qa, AC3): a row whose key lost letters to the title
    pattern gives its pseudonym to the composed name."""
    old = _old_cut(case)
    written = _reuse(tmp_path, db, old)
    assert written == old.expected, repr(written)  # expected failure
    assert len(_rows(db)) == 1


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason=XFAIL.format("« Entity Not Found » for a row that lost letters"),
)
@pytest.mark.parametrize("case", CUT)
def test_row_that_lost_letters_erased_by_name(
    tmp_path: Path, db: Path, case: str
) -> None:
    """T28 (Lionel's Qa, AC3): ``delete-mapping`` with the composed name
    erases that row, with its entry."""
    run = _erase(db, _old_cut(case))
    assert run.exit_code == 0, run.shown  # expected failure
    assert _rows(db) == []
    assert _erasures(db) == 1


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason=XFAIL.format("the decomposed row's pseudonym for a decomposed document"),
)
def test_rule_a_the_usual_way_pseudonym(tmp_path: Path, db: Path) -> None:
    """T29 (Lionel's Qb, rule A): a name saved both ways (decomposed first):
    documents in either form get the composed row's pseudonym; both rows
    stay."""
    _save(db, [(nfd(ZQ), PERSON, XQ), (ZQ, PERSON, XV)])
    _pre(len(_rows(db)) == 2, "two rows stored")
    source_a, output_a = tmp_path / "a.txt", tmp_path / "a_out.txt"
    source_b, output_b = tmp_path / "b.txt", tmp_path / "b_out.txt"
    source_a.write_text(nfd(R_TEXT), encoding="utf-8")
    source_b.write_text(R_TEXT, encoding="utf-8")
    run_a = _process(source_a, output_a, db, [(ZQ, PERSON)], [ENTER, SPACE])
    run_b = _process(source_b, output_b, db, [(ZQ, PERSON)], [ENTER, SPACE])
    _pre(run_a.exit_code == 0 and run_b.exit_code == 0, "both written")
    written_a, written_b = _written(output_a), _written(output_b)
    assert nfc(written_a) == R_REUSED, repr(written_a)  # expected failure
    assert nfc(written_b) == R_REUSED, repr(written_b)
    assert len(_rows(db)) == 2


@pytest.mark.xfail(
    strict=True,
    raises=AssertionError,
    reason=XFAIL.format("the one-row screen or « one per type of this name »"),
)
@pytest.mark.parametrize("other", ["decomposed_same_type", "composed_other_type"])
def test_delete_mapping_several_rows_line(tmp_path: Path, db: Path, other: str) -> None:
    """T30 (Lionel's Qc, AC9): the line for several saved pseudonyms of one
    name, two forms of one type or two types."""
    second = (
        (nfd(ZQ), PERSON, XQ)
        if other == "decomposed_same_type"
        else (ZQ, LOCATION, "Xantorville")
    )
    _save(db, [(ZQ, PERSON, XV), second])
    _pre(len(_rows(db)) == 2, "two rows stored")
    run = _delete(db, ZQ)
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    assert SEVERAL_LINE in run.flat, run.shown  # expected failure
    assert "one per type of this name" not in run.flat


# ---------------------------------------------------------------------------
# Guards (pass on the unfixed and the fixed code)
# ---------------------------------------------------------------------------


def test_guard_add_reaches_the_whole_word_with_a_mark(tmp_path: Path, db: Path) -> None:
    """G2 (AC4): [A] « Zox » + U+0301 + « rbalia » typed whole replaces the
    word; no mark is left."""
    source, output = tmp_path / "in.txt", tmp_path / "out.txt"
    source.write_text(f"Vorlan Mestrix valide. {ZOX} a signé. Zox valide.", "utf-8")
    run = _process(
        source, output, db, [(VM, PERSON)], [ENTER, "a", SPACE], [ZOX, PERSON]
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    assert "rbalia" not in written, repr(written)
    assert not _has_mark(written)


def test_text_without_accents_unchanged(tmp_path: Path, db: Path) -> None:
    """G3 (AC13): a document without accents, nothing detected, is written
    byte for byte."""
    source, output = tmp_path / "in.txt", tmp_path / "out.txt"
    source.write_text(
        "Vorlan Mestrix valide.\nLe contrat avec Zorbal Quentrix.", encoding="utf-8"
    )
    run = _process(source, output, db)
    _pre(run.exit_code == 0 and output.is_file(), f"written: {run.shown!r}")
    assert output.read_bytes() == source.read_bytes()


def test_compatibility_forms_unchanged(tmp_path: Path) -> None:
    """G4 (AC1, AC13): the reader keeps a compatibility form (the ligature
    U+FB01)."""
    text = "Le contrat avec Zorbaﬁx a signé."
    path = tmp_path / "doc.txt"
    path.write_text(text, encoding="utf-8")
    assert read_file(str(path)) == text


def test_composed_everywhere_unchanged(tmp_path: Path, db: Path) -> None:
    """G5 (AC13): T5 composed and typed composed, as today."""
    source, output = tmp_path / "in.txt", tmp_path / "out.txt"
    source.write_text(A_TEXT, encoding="utf-8")
    run = _process(
        source, output, db, [(VM, PERSON)], [ENTER, "a", SPACE], [ZO, PERSON]
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    assert NOT_FOUND not in run.flat, run.shown
    assert "rbalia" not in written, repr(written)
    assert _names(db) == sorted([(VM, PERSON), (ZO, PERSON)])


def test_no_saved_row_rewritten(tmp_path: Path, db: Path) -> None:
    """G6 (Lionel's Qa, AC13): a row saved decomposed keeps its stored name
    after a composed document met the name."""
    _save(db, [(nfd(ZQ), PERSON, XV)])
    source, output = tmp_path / "in.txt", tmp_path / "out.txt"
    source.write_text(R_TEXT, encoding="utf-8")
    run = _process(source, output, db, [(ZQ, PERSON)], [ENTER, SPACE])
    _pre(run.exit_code == 0 and output.is_file(), f"written: {run.shown!r}")
    assert (nfd(ZQ), PERSON, XV) in _rows(db)


@pytest.mark.parametrize("typed", ["Xantor Zorbal", "Xantor Zorbal Quentrix"])
def test_capitalised_words_after_an_accent_stored_apart(
    tmp_path: Path, db: Path, typed: str
) -> None:
    """G7 (QA's TEST-004 on 10.16, AC4): after « Éloria » written decomposed,
    « Zorbal » and « Zorbal Quentrix » are capitalised words of the
    document, so [C] refuses a pseudonym holding them."""
    source, output = tmp_path / "in.txt", tmp_path / "out.txt"
    source.write_text(nfd("Éloria") + G7_REST, encoding="utf-8")
    run = _process(
        source, output, db, [("quentrel quovan", PERSON)], [ENTER, "c", SPACE], [typed]
    )
    _pre(run.exit_code == 0, f"exit code 0: {run.shown!r}")
    written = _written(output)
    assert "Xantor" not in written, repr(written)
    composed = nfc(written)
    assert composed.startswith("Éloria a signé avec"), repr(written)
    assert composed.endswith("et Zorbal Quentrix."), repr(written)
